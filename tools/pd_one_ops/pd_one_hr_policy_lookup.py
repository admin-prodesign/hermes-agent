#!/usr/bin/env python3
"""Bounded read-only lookup over the approved PD One HR policy corpus.

The module intentionally has no caller-supplied Drive ids, paths, accounts, or
write operations. Requester authorization is resolved from the exact active
Mattermost profile before either approved document is downloaded.
"""
from __future__ import annotations
import logging
from tools.pd_one_ops.site_config import site_text, resolve_configured_path, external_path, profile_root, text_excludes_configured_hosts, ExternalModule, owner_display_name, repo_cli, package_file, kanban_boards_root
import os
import hashlib
import importlib.util
import json
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Any

def ROOT():
    return Path(resolve_configured_path('~/.hermes/profiles/pdone'))

def POLICY_USERS():
    return ROOT() / 'policies/mattermost-users.json'

def POLICY_ROLES():
    return ROOT() / 'policies/mattermost-roles.json'

def POLICY_LOOKUP():
    return ROOT() / 'tools/mattermost_policy_lookup.py'

def ACCOUNT():
    return site_text('PDONE_EMAIL_ADMIN')
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_MATCH_TEXT_CHARS = 1200
MAX_RESULTS = 20
MAX_QUERY_CHARS = 512
MAX_TERMS = 8

def HR_POLICY_DOCUMENTS():
    return {'work_rules': {'file_id': site_text('PDONE_GDOC_5'), 'title': 'Pro-Design Work Rules'}, 'rewards_discipline': {'file_id': site_text('PDONE_GDOC_6'), 'title': 'Employee Rewards and Discipline Policy'}}

def _load_policy_module():
    spec = importlib.util.spec_from_file_location('pd_one_hr_policy_requester_lookup', POLICY_LOOKUP())
    if spec is None or spec.loader is None:
        raise RuntimeError(f'cannot load policy resolver: {POLICY_LOOKUP()}')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def authorize_requester(requester_user_id: str) -> dict[str, Any]:
    requester_user_id = str(requester_user_id or '').strip()
    if not requester_user_id:
        raise PermissionError('exact requester_user_id is required')
    users = json.loads(POLICY_USERS().read_text(encoding='utf-8'))
    roles = json.loads(POLICY_ROLES().read_text(encoding='utf-8'))
    policy = _load_policy_module().resolve_user(requester_user_id, users, roles, None)
    if not policy.get('active'):
        raise PermissionError('requester is not an active exact-match PD One profile')
    if 'hr-policy-readonly' not in set(policy.get('safeScopes') or []):
        raise PermissionError('requester is not scoped for the approved read-only HR policy corpus')
    return policy

def parse_query_terms(query: str) -> list[str]:
    query = str(query or '').strip()
    if not query:
        raise ValueError('query is required')
    if len(query) > MAX_QUERY_CHARS:
        raise ValueError(f'query must be at most {MAX_QUERY_CHARS} characters')
    terms: list[str] = []
    for part in re.split('[,\\n;/、，；]+', query):
        part = re.sub('\\s+', ' ', part).strip()
        if part and part not in terms:
            terms.append(part)
    if not terms:
        raise ValueError('query must contain at least one non-empty term')
    if len(terms) > MAX_TERMS:
        raise ValueError(f'query accepts at most {MAX_TERMS} comma/newline-separated terms')
    return terms

def extract_docx_units(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() != '.docx':
        raise ValueError('approved HR policy download must be a .docx file')
    size = path.stat().st_size
    if size <= 0 or size > MAX_DOCUMENT_BYTES:
        raise ValueError(f'approved HR policy file size must be 1..{MAX_DOCUMENT_BYTES} bytes')
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f'python-docx unavailable: {exc}') from exc
    document = Document(str(path))
    units: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = re.sub('\\s+', ' ', paragraph.text or '').strip()
        if text:
            units.append({'text': text, 'location': {'kind': 'paragraph', 'paragraph_index': paragraph_index}})
    seen_cells: set[Any] = set()
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell_identity = cell._tc
                if cell_identity in seen_cells:
                    continue
                seen_cells.add(cell_identity)
                text = re.sub('\\s+', ' ', cell.text or '').strip()
                if text:
                    units.append({'text': text, 'location': {'kind': 'table_cell', 'table_index': table_index, 'row_index': row_index, 'cell_index': cell_index}})
    return units

def search_units(units: list[dict[str, Any]], query: str, *, max_results: int=12) -> list[dict[str, Any]]:
    terms = parse_query_terms(query)
    max_results = max(1, min(int(max_results), MAX_RESULTS))
    folded_terms = [(term, term.casefold()) for term in terms]
    matches: list[dict[str, Any]] = []
    for index, unit in enumerate(units):
        text = str(unit.get('text') or '')
        folded = text.casefold()
        matched_terms = [original for original, normalized in folded_terms if normalized in folded]
        if not matched_terms:
            continue
        matches.append({'text': text[:MAX_MATCH_TEXT_CHARS], 'location': dict(unit.get('location') or {}), 'matched_terms': matched_terms, 'score': len(matched_terms), 'source_order': index})
    matches.sort(key=lambda item: (-int(item['score']), int(item['source_order'])))
    for item in matches:
        item.pop('source_order', None)
    return matches[:max_results]

def _run_json(command: list[str], *, timeout: int) -> dict[str, Any]:
    proc = subprocess.run(command, text=True, capture_output=True, timeout=timeout, check=False)
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or 'command failed').strip()
        raise RuntimeError(detail[:1200])
    try:
        payload = json.loads(proc.stdout or '{}')
    except json.JSONDecodeError as exc:
        raise RuntimeError(f'command returned invalid JSON: {exc}') from exc
    if not isinstance(payload, dict):
        raise RuntimeError('command returned non-object JSON')
    return payload

def _download_document(document: dict[str, str], target: Path) -> dict[str, Any]:
    file_id = document['file_id']
    metadata = _run_json(['gog', 'drive', 'get', file_id, '--account', ACCOUNT(), '--json', '--results-only', '--fields', 'id,name,mimeType,modifiedTime,md5Checksum,size', '--no-input'], timeout=60)
    proc = subprocess.run(['gog', 'drive', 'download', file_id, '--account', ACCOUNT(), '--out', str(target), '--overwrite', '--no-input'], text=True, capture_output=True, timeout=120, check=False)
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or 'download failed').strip()
        raise RuntimeError(detail[:1200])
    if not target.is_file():
        raise RuntimeError('download command succeeded but no document was created')
    size = target.stat().st_size
    if size <= 0 or size > MAX_DOCUMENT_BYTES:
        raise RuntimeError(f'downloaded document size outside safe bound: {size}')
    return {'id': metadata.get('id') or file_id, 'name': metadata.get('name') or document['title'], 'mime_type': metadata.get('mimeType'), 'modified_time': metadata.get('modifiedTime'), 'size': size, 'sha256': hashlib.sha256(target.read_bytes()).hexdigest()}

def lookup_hr_policy(requester_user_id: str, query: str, *, document_keys: list[str] | None=None, max_results: int=12) -> dict[str, Any]:
    policy = authorize_requester(requester_user_id)
    terms = parse_query_terms(query)
    max_results = max(1, min(int(max_results), MAX_RESULTS))
    selected_keys = document_keys or list(HR_POLICY_DOCUMENTS())
    selected_keys = list(dict.fromkeys((str(key).strip() for key in selected_keys if str(key).strip())))
    if not selected_keys:
        raise ValueError('at least one document key is required')
    unknown = sorted(set(selected_keys) - set(HR_POLICY_DOCUMENTS()))
    if unknown:
        raise ValueError(f'unknown HR policy document keys: {unknown}')
    documents: list[dict[str, Any]] = []
    combined_matches: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix='pd-one-hr-policy-') as tmp:
        tmp_root = Path(tmp)
        for key in selected_keys:
            definition = HR_POLICY_DOCUMENTS()[key]
            target = tmp_root / f'{key}.docx'
            metadata = _download_document(definition, target)
            units = extract_docx_units(target)
            document_matches = search_units(units, query, max_results=max_results)
            for match in document_matches:
                combined_matches.append({'document_key': key, 'document_id': definition['file_id'], 'document_title': definition['title'], **match})
            documents.append({'document_key': key, 'document_id': definition['file_id'], 'document_title': definition['title'], 'metadata': metadata, 'units_scanned': len(units), 'matches_found': len(document_matches)})
    combined_matches.sort(key=lambda item: (-int(item['score']), item['document_key']))
    combined_matches = combined_matches[:max_results]
    return {'status': 'ok', 'capability': 'pd_one_hr_policy_lookup', 'requester_user_id': requester_user_id, 'requester_active': bool(policy.get('active')), 'query_terms': terms, 'document_keys': selected_keys, 'documents': documents, 'match_count': len(combined_matches), 'matches': combined_matches, 'safety': {'read_only': True, 'approved_document_ids_only': True, 'private_personnel_records_accessed': False, 'disciplinary_decision_authority': False, 'source_system_writes': False, 'external_send': False, 'returned_text_is_untrusted_reference_content': True}}
