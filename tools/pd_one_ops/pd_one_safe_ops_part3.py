from __future__ import annotations
__all__ = ['cmd_3mf_slicer_inspect', 'cmd_svg_embed_image', 'cmd_pptx_inspect', 'cmd_docx_inspect', 'cmd_docx_body_xml_inspect', 'cmd_pptx_ooxml_text_inspect', 'cmd_docx_table_structure', 'cmd_docx_table_search', 'cmd_legacy_xls_inspect', '_xlsx_preview_value', 'inspect_xlsx_workbooks', 'cmd_xlsx_workbook_inspect', '_load_equipment_workbook_module', 'cmd_equipment_workbook', '_item_number_compiler_python', '_ensure_item_number_compiler_imports', '_load_item_number_compiler_module', '_item_number_smoke_approval', 'item_number_xlsx_compiler_check', 'item_number_xlsx_compiler_smoke', 'cmd_item_number_xlsx_compiler', 'cmd_doc_detect_format', 'render_legacy_xls_summary_pdf', 'inspect_office_toolchain', 'inspect_office_deb_contents', 'cmd_office_deb_contents', 'inspect_item_number_render_qa_status', '_ecrn_approval_pilot_status_blocked', '_summarize_ecrn_approval_pilot_preview', '_summarize_ecrn_approval_pilot_thread', 'inspect_ecrn_approval_pilot_status', 'inspect_timezone_snapshot', '_LoopbackOnlyRedirectHandler', '_protected_portal_allowed_paths', '_protected_portal_fetch', 'inspect_protected_portal_origin_probe', 'inspect_timestamp_to_epoch', 'inspect_epoch_milliseconds_to_iso', 'SHEETS_SERIAL_EPOCH', 'SHEETS_SERIAL_MIN', 'SHEETS_SERIAL_MAX', 'SHEETS_SERIAL_MAX_ITEMS', 'SHEETS_COLUMN_MAX_LEN', 'SHEETS_COLUMN_MAX_INDEX', 'SHEETS_COLUMN_MAX_ITEMS', 'SHEETS_RANGE_MAX_ITEMS', '_a1_column_to_index', 'inspect_sheets_serial_column_convert', 'cmd_timezone_snapshot', 'cmd_protected_portal_origin_probe', 'cmd_timestamp_to_epoch', 'cmd_epoch_milliseconds_to_iso', 'cmd_sheets_serial_column_convert', 'cmd_office_toolchain_status', 'cmd_item_number_render_qa_status', 'cmd_item_number_historical_sample_summary', 'cmd_item_number_local_fixture_copy', 'cmd_item_number_windows_excel_qa', 'cmd_ecrn_approval_pilot_status', 'cmd_ecrn_approval_pilot_check', '_item_number_workflow_component_paths', 'cmd_item_number_workflow_check', 'convert_office_to_pdf', 'cmd_office_convert_pdf', '_require_str', '_require_list', '_apply_table', 'build_html_from_paragraph_spec', 'cmd_json_html_artifact', '_iter_docx_table_paragraphs', '_iter_docx_story_paragraphs', 'inspect_or_stage_docx_layout', 'cmd_docx_layout_patch', '_replace_one_visible_paragraph_text', 'build_company_announcement_from_spec', 'build_docx_from_spec', '_office_visual_qa_file_receipts', 'run_office_visual_qa', 'run_visual_qa', '_office_docprops_local_name', '_office_docprops_safe_member', '_office_docprops_parse_xml', 'cmd_office_docprops_inspect', 'cmd_office_visual_qa', 'convert_html_to_docx', 'cmd_office_html_to_docx', 'cmd_doc_build', '_load_ecrn_docx_patch_module', 'cmd_ecrn_docx_unique_text_patch']
import logging
from tools.pd_one_ops.site_config import site_text, resolve_configured_path, external_path, profile_root, text_excludes_configured_hosts, ExternalModule, owner_display_name, repo_cli, package_file, kanban_boards_root
import argparse
import base64
import calendar
import contextlib
import collections
import csv
import difflib
import fnmatch
import importlib.util
import hashlib
import html
import math
import io
import json
import os
import re
import shlex
import shutil
import socket
import subprocess
import types
import sys
import tempfile
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any
from urllib.parse import urlparse
from zoneinfo import ZoneInfo
from tools.pd_one_ops.pd_one_safe_ops_part0 import *
from tools.pd_one_ops.pd_one_safe_ops_part1 import *
from tools.pd_one_ops.pd_one_safe_ops_part2 import *

def cmd_3mf_slicer_inspect(args: argparse.Namespace) -> None:
    action = _PdUnset
    archive = _PdUnset
    archive_name = _PdUnset
    members = _PdUnset
    missing = _PdUnset
    namelist = _PdUnset
    output_name = _PdUnset
    payload = _PdUnset
    peeked_members = _PdUnset
    seen_keys = _PdUnset
    settings = _PdUnset
    source = _PdUnset
    source_sha256 = _PdUnset
    source_size = _PdUnset
    unique_settings = _PdUnset

    def _pd_phase_9():
        nonlocal action, archive, archive_name, members, missing, namelist, output_name, payload, peeked_members, seen_keys, settings, source, source_sha256, source_size, unique_settings
        'Inspect or stage allowlisted PrusaSlicer metadata from one approved .3mf.'
        action = args.action
        source = Path(args.input).resolve(strict=False)
        if not any((is_within(source, root) for root in THREEMF_SLICER_INPUT_ROOTS())):
            fail(f'path is outside approved PD One 3MF roots: {source}')
        if not source.is_file() or source.suffix.lower() != '.3mf':
            fail(f'input must be a regular .3mf file: {source}')
        if source.is_symlink():
            fail('symlinked 3MF inputs are not allowed')
        source_size = source.stat().st_size
        if source_size <= 0 or source_size > THREEMF_SLICER_MAX_SOURCE_BYTES:
            fail(f'3MF size must be within 1-{THREEMF_SLICER_MAX_SOURCE_BYTES} bytes: {source_size}')
        source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
        if not zipfile.is_zipfile(source):
            fail('input is not a valid ZIP/3MF archive')
        try:
            archive = zipfile.ZipFile(source)
        except zipfile.BadZipFile as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            fail(f'invalid 3MF zip: {exc}')
        with archive:
            namelist = archive.namelist()
            if len(namelist) > THREEMF_SLICER_MAX_NAMELIST:
                fail(f'3MF exceeds {THREEMF_SLICER_MAX_NAMELIST} member limit')
            infos = {_safe_3mf_zip_info(info): info for info in archive.infolist()}
            members = {}
            settings = []
            required = ('Metadata/Slic3r_PE.config', 'Metadata/Slic3r_PE_model.config')
            missing = [name for name in required if name not in infos]
            if missing and action == 'stage':
                fail(f'3MF is missing required slicer metadata: {missing}')
            peeked_members = []
            for name in namelist:
                lowered = name.lower()
                if not lowered.endswith(CAD_3MF_PEEK_SUFFIXES):
                    continue
                info = infos.get(name)
                if info is None:
                    continue
                size = int(info.file_size)
                preview = ''
                preview_truncated = size > CAD_3MF_PEEK_BYTES
                if 0 < size <= THREEMF_SLICER_MAX_MEMBER_BYTES:
                    preview = archive.read(name)[:CAD_3MF_PEEK_BYTES].decode('utf-8', errors='replace')
                peeked_members.append({'name': name, 'size_bytes': size, 'preview_truncated': preview_truncated, 'preview_chars': len(preview)})
                if len(peeked_members) >= CAD_3MF_MAX_PEEK_FILES:
                    break
            for archive_name, output_name in THREEMF_SLICER_MEMBERS.items():
                present = archive_name in infos
                payload = b''
                if present:
                    payload = archive.read(archive_name)
                    if len(payload) > THREEMF_SLICER_MAX_MEMBER_BYTES:
                        fail(f'3MF member exceeds {THREEMF_SLICER_MAX_MEMBER_BYTES} byte limit: {archive_name}')
                    if archive_name.endswith('.png') and (not payload.startswith(b'\x89PNG\r\n\x1a\n')):
                        fail('thumbnail member is not a PNG')
                    if archive_name.endswith('.config'):
                        settings.extend(_parse_3mf_slicer_settings(payload.decode('utf-8', errors='replace')))
                members[output_name] = {'archive_name': archive_name, 'present': present, 'size_bytes': len(payload) if present else 0, 'sha256': hashlib.sha256(payload).hexdigest() if present else None}
        unique_settings = []
        seen_keys = set()

    def _pd_phase_10():
        nonlocal archive, archive_name, output_name, payload
        for item in settings:
            if item['key'] in seen_keys:
                continue
            seen_keys.add(item['key'])
            unique_settings.append(item)
        safety = {'admin_only': True, 'source_read_only': True, 'approved_roots_only': True, 'allowlisted_members_only': True, 'bounded_output': True, 'source_overwrite_forbidden': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'shell_used': False, 'zip_extract_used': False}
        receipt: dict[str, Any] = {'status': 'ok', 'capability': 'pd_one_3mf_slicer_inspect', 'action': action, 'input': str(source), 'source_size_bytes': source_size, 'source_sha256': source_sha256, 'member_count': len(namelist), 'namelist_sample': namelist[:CAD_3MF_NAMELIST_SAMPLE], 'namelist_truncated': len(namelist) > CAD_3MF_NAMELIST_SAMPLE, 'missing_required_members': missing, 'peeked_members': peeked_members, 'freecad_runtime': inspect_cad_runtime(), 'members': members, 'selected_settings': unique_settings, 'selected_setting_count': len(unique_settings), 'safety': safety}
        if missing:
            receipt['status'] = 'partial'
            receipt['blocker'] = 'missing_slicer_metadata'
        if action == 'inspect':
            receipt['staged'] = False
            emit(receipt)
            return _PdReturn(None)
        source_key = source_sha256[:12]
        safe_stem = re.sub('[^A-Za-z0-9._-]+', '_', source.stem)[:48] or 'model'
        output_dir = (THREEMF_SLICER_OUTPUT_ROOT() / f'{safe_stem}-{source_key}').resolve(strict=False)
        if not is_within(output_dir, THREEMF_SLICER_OUTPUT_ROOT()):
            fail('resolved output directory escaped the approved 3MF staging root')
        output_dir.mkdir(parents=True, exist_ok=True)
        staged: list[dict[str, Any]] = []
        already_present = True
        with zipfile.ZipFile(source) as archive:
            for archive_name, output_name in THREEMF_SLICER_MEMBERS.items():
                if archive_name not in archive.namelist():
                    continue
                payload = archive.read(archive_name)
                final_path = output_dir / output_name
                if not is_within(final_path, output_dir):
                    fail(f'staged path escaped output directory: {final_path}')
                expected_sha = hashlib.sha256(payload).hexdigest()
                if final_path.is_file() and hashlib.sha256(final_path.read_bytes()).hexdigest() == expected_sha:
                    staged.append({'archive_name': archive_name, 'path': str(final_path), 'size_bytes': final_path.stat().st_size, 'sha256': expected_sha, 'action': 'already_present'})
                    continue
                already_present = False
                temp_path = output_dir / f'.{output_name}.{os.getpid()}.tmp'
                temp_path.write_bytes(payload)
                os.replace(temp_path, final_path)
                staged.append({'archive_name': archive_name, 'path': str(final_path), 'size_bytes': final_path.stat().st_size, 'sha256': hashlib.sha256(final_path.read_bytes()).hexdigest(), 'action': 'copied'})
        receipt['staged'] = True
        receipt['output_dir'] = str(output_dir)
        receipt['artifacts'] = staged
        receipt['stage_action'] = 'already_present' if already_present else 'copied'
        emit(receipt)
    _pd_result = _pd_phase_9()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_10()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def cmd_svg_embed_image(args: argparse.Namespace) -> None:
    """Inspect or stage one SVG with the fixed Pro-Design logo embedded as a data URI."""
    path = Path(args.input).resolve(strict=False)
    if not any((is_within(path, root) for root in ARTIFACT_BASE64_ALLOWED_ROOTS())):
        fail(f'path is outside approved PD One artifact roots: {path}')
    if not path.is_file() or path.suffix.lower() != '.svg':
        fail(f'input must be a regular .svg file: {path}')
    size = path.stat().st_size
    if size > SVG_EMBED_IMAGE_MAX_SVG_BYTES:
        fail(f'SVG exceeds {SVG_EMBED_IMAGE_MAX_SVG_BYTES} byte limit: {size}')
    try:
        source = path.read_text(encoding='utf-8')
    except UnicodeDecodeError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail('SVG must be valid UTF-8')
    double_ref = f'href="{SVG_EMBED_IMAGE_ALLOWED_URL}"'
    single_ref = f"href='{SVG_EMBED_IMAGE_ALLOWED_URL}'"
    reference_count = source.count(double_ref) + source.count(single_ref)
    embedded_png_count = source.count('data:image/png;base64,')
    source_bytes = source.encode('utf-8')
    source_sha256 = hashlib.sha256(source_bytes).hexdigest()
    state = 'external_reference' if reference_count == 1 else 'already_embedded' if reference_count == 0 and embedded_png_count > 0 else 'not_applicable'
    base_receipt = {'status': 'ok', 'capability': 'pd_one_svg_embed_image', 'action': args.action, 'input': str(path), 'source_sha256': source_sha256, 'source_size_bytes': size, 'external_reference_count': reference_count, 'embedded_png_count': embedded_png_count, 'state': state, 'fixed_remote_url': SVG_EMBED_IMAGE_ALLOWED_URL, 'safety': {'admin_only': True, 'approved_artifact_roots_only': True, 'fixed_remote_url_only': True, 'source_overwrite_forbidden': True, 'staged_output_root': str(SVG_EMBED_IMAGE_OUTPUT_ROOT()), 'external_delivery': False, 'shell_used': False}}
    if args.action == 'inspect':
        base_receipt['read_only'] = True
        emit(base_receipt)
        return
    if reference_count != 1:
        fail(f'stage requires exactly one fixed external logo reference; found {reference_count} (state={state})')
    request = urllib.request.Request(SVG_EMBED_IMAGE_ALLOWED_URL, headers={'User-Agent': 'PD-One-Safe-Ops/1.0'})
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            remote = response.read(SVG_EMBED_IMAGE_MAX_REMOTE_BYTES + 1)
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'fixed logo download failed: {type(exc).__name__}: {exc}')
    if len(remote) > SVG_EMBED_IMAGE_MAX_REMOTE_BYTES:
        fail(f'fixed logo exceeds {SVG_EMBED_IMAGE_MAX_REMOTE_BYTES} byte limit')
    if not remote.startswith(b'\x89PNG\r\n\x1a\n'):
        fail('fixed logo response is not a PNG')
    data_uri = 'data:image/png;base64,' + base64.b64encode(remote).decode('ascii')
    old_ref = double_ref if double_ref in source else single_ref
    quote = '"' if old_ref == double_ref else "'"
    rendered = source.replace(old_ref, f'href={quote}{data_uri}{quote}', 1)
    output = (SVG_EMBED_IMAGE_OUTPUT_ROOT() / f'{path.stem}.embedded.svg').resolve(strict=False)
    SVG_EMBED_IMAGE_OUTPUT_ROOT().mkdir(parents=True, exist_ok=True)
    rendered_bytes = rendered.encode('utf-8')
    temp_output = output.with_name(f'.{output.name}.{os.getpid()}.tmp')
    try:
        temp_output.write_bytes(rendered_bytes)
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)
    readback = output.read_text(encoding='utf-8')
    if SVG_EMBED_IMAGE_ALLOWED_URL in readback or data_uri not in readback:
        fail('staged SVG readback verification failed')
    if hashlib.sha256(path.read_bytes()).hexdigest() != source_sha256:
        fail('source SVG changed during staging; staged result cannot be accepted')
    base_receipt.update({'state': 'staged', 'read_only': False, 'output': str(output), 'output_size_bytes': len(rendered_bytes), 'output_sha256': hashlib.sha256(rendered_bytes).hexdigest(), 'remote_size_bytes': len(remote), 'remote_sha256': hashlib.sha256(remote).hexdigest(), 'readback': {'external_reference_count': readback.count(SVG_EMBED_IMAGE_ALLOWED_URL), 'embedded_png': data_uri in readback, 'source_unchanged': True}})
    emit(base_receipt)

def cmd_pptx_inspect(args: argparse.Namespace) -> None:
    emit(inspect_pptx(Path(args.input), include_text=args.include_text))

def cmd_docx_inspect(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    result = inspect_docx_paragraph(path, table_index=args.table_index, row_index=args.row_index, cell_index=args.cell_index, paragraph_index=args.paragraph_index, include_runs=args.include_runs)
    result.update({'status': 'ok', 'capability': 'pd_one_docx_inspect', 'input': str(path), 'type': '.docx', 'safety': {'read_only': True, 'admin_only': True, 'approved_document_root': True, 'bounded_target': True, 'shell_used': False, 'source_system_writes_forbidden': True}})
    emit(result)

def cmd_docx_body_xml_inspect(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    result = inspect_docx_body_xml(path, list(args.paragraph_index or []))
    result.update({'status': 'ok', 'capability': 'pd_one_docx_body_xml_inspect', 'input': str(path), 'type': '.docx', 'safety': {'read_only': True, 'admin_only': True, 'approved_document_root': True, 'bounded_targets': True, 'bounded_xml_receipt': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'external_delivery': False}})
    emit(result)

def cmd_pptx_ooxml_text_inspect(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    try:
        emit(inspect_pptx_ooxml_text(path, list(args.member or [])))
    except (ValueError, RuntimeError, ET.ParseError, zipfile.BadZipFile, KeyError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))

def cmd_docx_table_structure(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    emit(inspect_docx_table_structure(path))

def cmd_docx_table_search(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    emit(inspect_docx_table_matches(path, list(args.term or [])))

def cmd_legacy_xls_inspect(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    if path.suffix.lower() != '.xls':
        fail('legacy-xls-inspect only supports .xls files')
    result = inspect_legacy_xls(path, sheet=args.sheet, max_sheets=_bounded_int(args.max_sheets, '--max-sheets', 1, 50), preview_rows=_bounded_int(args.preview_rows, '--preview-rows', 1, 100), preview_cols=_bounded_int(args.preview_cols, '--preview-cols', 1, 100), max_nonempty_cells=_bounded_int(args.max_nonempty_cells, '--max-nonempty-cells', 1, 2000), max_merged_cells=_bounded_int(args.max_merged_cells, '--max-merged-cells', 0, 200))
    emit(result)

def _xlsx_preview_value(value: Any, max_chars: int) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    text = str(value)
    if len(text) > max_chars:
        return text[:max_chars] + '…'
    return text

def inspect_xlsx_workbooks(input_paths: list[str | Path], *, sheet: str | None=None, preview_input: str | Path | None=None, row_start: int=1, row_end: int=25, max_columns: int=25, include_values: bool=False) -> dict[str, Any]:
    any_error = _PdUnset
    path = _PdUnset
    reports = _PdUnset
    resolved_inputs = _PdUnset
    resolved_preview_input = _PdUnset
    suffix = _PdUnset

    def _pd_phase_11():
        nonlocal any_error, max_columns, path, reports, resolved_inputs, resolved_preview_input, row_end, row_start, suffix
        'Read bounded XLSX/XLSM workbook metadata and an optional sheet preview.\n\n    This is deliberately read-only and restricted to the approved PD One\n    document roots.  Formula cells are returned as stored (``data_only=False``)\n    strings; this wrapper does not recalculate workbooks or write artifacts.\n    '
        if not input_paths:
            raise ValueError('at least one input workbook is required')
        if len(input_paths) > XLSX_WORKBOOK_INSPECT_MAX_FILES:
            raise ValueError(f'at most {XLSX_WORKBOOK_INSPECT_MAX_FILES} input workbooks are allowed')
        if sheet is not None and len(sheet) > 200:
            raise ValueError('sheet must be <= 200 characters')
        row_start = int(row_start)
        row_end = int(row_end)
        max_columns = int(max_columns)
        if row_start < 1 or row_start > XLSX_WORKBOOK_INSPECT_MAX_ROWS:
            raise ValueError(f'row_start must be between 1 and {XLSX_WORKBOOK_INSPECT_MAX_ROWS}')
        if row_end < row_start or row_end > XLSX_WORKBOOK_INSPECT_MAX_ROWS:
            raise ValueError(f'row_end must be between row_start and {XLSX_WORKBOOK_INSPECT_MAX_ROWS}')
        if max_columns < 1 or max_columns > XLSX_WORKBOOK_INSPECT_MAX_COLUMNS:
            raise ValueError(f'max_columns must be between 1 and {XLSX_WORKBOOK_INSPECT_MAX_COLUMNS}')
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise RuntimeError(f'openpyxl is unavailable: {exc}') from exc
        resolved_inputs = []
        for raw_path in input_paths:
            path = require_allowed_doc_path(Path(raw_path))
            suffix = path.suffix.lower()
            if suffix not in XLSX_WORKBOOK_INSPECT_SUFFIXES:
                raise ValueError(f'xlsx-workbook-inspect only supports {sorted(XLSX_WORKBOOK_INSPECT_SUFFIXES)}: {path}')
            if path.stat().st_size > XLSX_WORKBOOK_INSPECT_MAX_BYTES:
                raise ValueError(f'workbook exceeds {XLSX_WORKBOOK_INSPECT_MAX_BYTES} bytes: {path}')
            resolved_inputs.append(path)
        resolved_preview_input = None
        if preview_input is not None:
            resolved_preview_input = require_allowed_doc_path(Path(preview_input))
            if resolved_preview_input not in resolved_inputs:
                raise ValueError('preview_input must be one of input_paths')
        reports = []
        any_error = False

    def _pd_phase_12():
        nonlocal any_error, path, suffix
        for path in resolved_inputs:
            suffix = path.suffix.lower()
            size_bytes = path.stat().st_size
            report: dict[str, Any] = {'input': str(path), 'name': path.name, 'file': {'size_bytes': size_bytes, 'suffix': suffix}}
            workbook = None
            try:
                workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                sheet_names = list(workbook.sheetnames)
                worksheet_reports: list[dict[str, Any]] = []
                for worksheet in list(workbook.worksheets)[:XLSX_WORKBOOK_INSPECT_MAX_SHEETS]:
                    worksheet_reports.append({'name': worksheet.title, 'rows': int(worksheet.max_row or 0), 'cols': int(worksheet.max_column or 0)})
                report['sheets'] = worksheet_reports
                report['sheet_count'] = len(sheet_names)
                report['sheet_names_truncated'] = len(sheet_names) > XLSX_WORKBOOK_INSPECT_MAX_SHEETS
                preview_this_input = sheet is not None and (resolved_preview_input is None or path == resolved_preview_input)
                if preview_this_input:
                    if sheet not in sheet_names:
                        any_error = True
                        report['preview'] = {'status': 'blocked', 'reason': 'sheet_not_found', 'requested_sheet': sheet, 'available_sheets': sheet_names[:XLSX_WORKBOOK_INSPECT_MAX_SHEETS]}
                    elif include_values:
                        worksheet = workbook[sheet]
                        last_row = min(row_end, int(worksheet.max_row or 0))
                        rows: list[list[Any]] = []
                        if last_row >= row_start:
                            for values in worksheet.iter_rows(min_row=row_start, max_row=last_row, min_col=1, max_col=max_columns, values_only=True):
                                rows.append([_xlsx_preview_value(value, XLSX_WORKBOOK_INSPECT_MAX_CELL_CHARS) for value in values])
                        report['preview'] = {'status': 'ok', 'sheet': sheet, 'row_start': row_start, 'row_end_requested': row_end, 'row_end_read': last_row, 'max_columns': max_columns, 'rows': rows, 'truncated': int(worksheet.max_row or 0) > row_end}
                    else:
                        report['preview'] = {'status': 'omitted', 'reason': 'include_values_false', 'requested_sheet': sheet}
                report['status'] = 'ok'
            except Exception as exc:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                any_error = True
                report['status'] = 'error'
                report['error'] = str(exc)[:500]
            finally:
                if workbook is not None:
                    try:
                        workbook.close()
                    except Exception:
                        logging.getLogger('pd_one_ops').exception('pd one command failed')
            reports.append(report)
        return _PdReturn({'status': 'partial' if any_error else 'ok', 'capability': 'pd_one_xlsx_workbook_inspect', 'workbooks': reports, 'request': {'input_count': len(input_paths), 'sheet': sheet, 'preview_input': str(resolved_preview_input) if resolved_preview_input else None, 'row_start': row_start, 'row_end': row_end, 'max_columns': max_columns, 'include_values': bool(include_values)}, 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'approved_document_roots_only': True, 'bounded_file_count': True, 'bounded_preview': True, 'formula_recalculation': False}})
    _pd_result = _pd_phase_11()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_12()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def cmd_xlsx_workbook_inspect(args: argparse.Namespace) -> None:
    inputs = list(args.input or [])
    if not inputs or len(inputs) > XLSX_WORKBOOK_INSPECT_MAX_FILES:
        fail(f'--input must be repeated 1 to {XLSX_WORKBOOK_INSPECT_MAX_FILES} times')
    try:
        result = inspect_xlsx_workbooks(inputs, sheet=args.sheet, preview_input=args.preview_input, row_start=_bounded_int(args.row_start, '--row-start', 1, XLSX_WORKBOOK_INSPECT_MAX_ROWS), row_end=_bounded_int(args.row_end, '--row-end', 1, XLSX_WORKBOOK_INSPECT_MAX_ROWS), max_columns=_bounded_int(args.max_columns, '--max-columns', 1, XLSX_WORKBOOK_INSPECT_MAX_COLUMNS), include_values=bool(args.include_values))
    except (ValueError, RuntimeError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))
    emit(result)
    if result.get('status') != 'ok':
        raise SystemExit(1)

def _load_equipment_workbook_module() -> Any:
    if not EQUIPMENT_WORKBOOK_SCRIPT().is_file():
        fail(f'equipment workbook wrapper not found: {EQUIPMENT_WORKBOOK_SCRIPT()}')
    spec = importlib.util.spec_from_file_location('pdone_equipment_workbook', EQUIPMENT_WORKBOOK_SCRIPT())
    if spec is None or spec.loader is None:
        fail(f'could not load equipment workbook wrapper: {EQUIPMENT_WORKBOOK_SCRIPT()}')
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module

def cmd_equipment_workbook(args: argparse.Namespace) -> None:
    action = str(getattr(args, 'action', 'check') or 'check').strip().lower()
    if action not in {'check', 'stage'}:
        fail('--action must be check or stage')
    try:
        result = _load_equipment_workbook_module().run(action)
    except (ValueError, RuntimeError, FileNotFoundError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))
    emit(result)
    if result.get('status') not in {'ok'}:
        raise SystemExit(1)

def _item_number_compiler_python() -> str:
    venv_python = ITEM_NUMBER_COMPILER_ROOT() / '.venv' / 'bin' / 'python'
    if venv_python.is_file() and os.access(venv_python, os.X_OK):
        return str(venv_python)
    return sys.executable

def _ensure_item_number_compiler_imports() -> None:
    try:
        import openpyxl
        import xlrd
        return
    except ImportError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
    site_packages = ITEM_NUMBER_COMPILER_ROOT() / '.venv' / 'lib' / f'python{sys.version_info.major}.{sys.version_info.minor}' / 'site-packages'
    if not site_packages.is_dir():
        fail(f'item-number compiler is missing openpyxl/xlrd; install compiler/requirements.txt into {ITEM_NUMBER_COMPILER_ROOT() / '.venv'}')
    sys.path.append(str(site_packages))

def _load_item_number_compiler_module() -> Any:
    if not ITEM_NUMBER_COMPILER_SCRIPT().exists():
        fail(f'item-number compiler not found: {ITEM_NUMBER_COMPILER_SCRIPT()}')
    _ensure_item_number_compiler_imports()
    spec = importlib.util.spec_from_file_location('pdone_item_number_xlsx_compiler', ITEM_NUMBER_COMPILER_SCRIPT())
    if spec is None or spec.loader is None:
        fail(f'could not load item-number compiler: {ITEM_NUMBER_COMPILER_SCRIPT()}')
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module

def _item_number_smoke_approval(target: str='department_manager') -> dict[str, Any]:
    if target == 'department_manager':
        return {'display_name': ITEM_NUMBER_SMOKE_DISPLAY_NAME, 'create_at': ITEM_NUMBER_SMOKE_CREATE_AT, 'status': '同意', 'user_id': ITEM_NUMBER_SMOKE_USER_ID, 'post_id': ITEM_NUMBER_SMOKE_POST_ID, 'revision_hash': ITEM_NUMBER_SMOKE_REVISION_HASH, 'comments': ITEM_NUMBER_SMOKE_COMMENTS}
    seed = hashlib.sha256(target.encode('utf-8')).hexdigest()
    return {'display_name': ITEM_NUMBER_SMOKE_DISPLAY_NAME, 'create_at': ITEM_NUMBER_SMOKE_CREATE_AT, 'status': '同意', 'user_id': seed[:26], 'post_id': hashlib.sha256(f'post:{target}'.encode('utf-8')).hexdigest()[:26], 'revision_hash': hashlib.sha256(f'revision:{target}'.encode('utf-8')).hexdigest(), 'comments': f'{ITEM_NUMBER_SMOKE_COMMENTS} [{target}]'}

def item_number_xlsx_compiler_check(runner=run_capture) -> dict[str, Any]:
    if not ITEM_NUMBER_COMPILER_SCRIPT().exists():
        fail(f'item-number compiler not found: {ITEM_NUMBER_COMPILER_SCRIPT()}')
    if not ITEM_NUMBER_COMPILER_TEST().exists():
        fail(f'item-number compiler test not found: {ITEM_NUMBER_COMPILER_TEST()}')
    compiler_python = _item_number_compiler_python()
    compile_proc = runner([compiler_python, '-m', 'py_compile', str(ITEM_NUMBER_COMPILER_SCRIPT()), str(ITEM_NUMBER_COMPILER_TEST())], timeout=120)
    unit_proc = runner([compiler_python, '-m', 'unittest', 'discover', '-s', 'tests', '-v'], timeout=600, cwd=str(ITEM_NUMBER_COMPILER_ROOT()))
    status = 'ok' if compile_proc.returncode == 0 and unit_proc.returncode == 0 else 'error'
    return {'status': status, 'capability': 'pd_one_item_number_xlsx_compiler', 'action': 'check', 'script': str(ITEM_NUMBER_COMPILER_SCRIPT()), 'unit_test': str(ITEM_NUMBER_COMPILER_TEST()), 'compiler_python': compiler_python, 'compile': {'returncode': compile_proc.returncode, 'stdout_tail': tail_text(getattr(compile_proc, 'stdout', '')), 'stderr_tail': tail_text(getattr(compile_proc, 'stderr', ''))}, 'unit': {'returncode': unit_proc.returncode, 'stdout_tail': tail_text(getattr(unit_proc, 'stdout', '')), 'stderr_tail': tail_text(getattr(unit_proc, 'stderr', ''))}, 'safety': {'shell_used': False, 'read_only': True, 'source_system_writes_forbidden': True}}

def item_number_xlsx_compiler_smoke(module: Any | None=None) -> dict[str, Any]:
    module = module or _load_item_number_compiler_module()
    source = require_allowed_doc_path(ITEM_NUMBER_SOURCE_XLS())
    if source.suffix.lower() != '.xls':
        fail('item-number smoke only supports the canonical .xls source')
    ITEM_NUMBER_WORKFLOW_TMP_ROOT().mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='item-number-xlsx-compiler-', dir=str(ITEM_NUMBER_WORKFLOW_TMP_ROOT())) as td:
        workdir = Path(td)
        template = workdir / ITEM_NUMBER_SMOKE_TEMPLATE_NAME
        migration = module.migrate_xls_to_xlsx(source, template)
        manifest_path = template.with_suffix('.template.json')
        if not template.is_file():
            fail('item-number smoke migration did not create the template')
        if not manifest_path.is_file():
            fail('item-number smoke migration did not create the manifest')
        route = ITEM_NUMBER_SMOKE_ROUTE
        fulfillment_link = ITEM_NUMBER_SMOKE_FULFILLMENT_LINK
        route_targets = sorted(module.ROUTE_SIGNATURE_TARGETS[route])
        approvals = {target: _item_number_smoke_approval(target) for target in route_targets}
        approval_snapshot_sha256 = hashlib.sha256(json.dumps({target: approvals[target] for target in route_targets}, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')).hexdigest()
        prepared = [module._prepare_approval(target, approvals[target]) for target in route_targets]
        compiler_approval_snapshot = [{key: item[key] for key in ('target', 'display_name', 'user_id', 'post_id', 'create_at', 'timestamp_taipei', 'timezone', 'status', 'revision_hash', 'comments', 'printed_text', 'printed_text_sha256')} for item in prepared]
        compiler_approval_snapshot_sha256 = hashlib.sha256(module._canonical_json_bytes(compiler_approval_snapshot)).hexdigest()
        workflow_receipt = {'root_post_id': ITEM_NUMBER_SMOKE_ROOT_POST_ID, 'source_file_sha256': migration['source_sha256'], 'frozen_config_sha256': migration['template_sha256'], 'approval_snapshot_sha256': approval_snapshot_sha256, 'compiler_approval_snapshot_sha256': compiler_approval_snapshot_sha256}
        compile_result = module.compile_signed_xlsx(template, approvals, workdir, route=route, fulfillment_link=fulfillment_link, workflow_receipt=workflow_receipt, manifest_path=manifest_path, allow_unqualified_test_template=True)
        output_path = Path(compile_result.output_path)
        receipt_path = Path(compile_result.receipt_path)
        if not output_path.is_file():
            fail('item-number smoke compile did not create the signed output')
        if not receipt_path.is_file():
            fail('item-number smoke compile did not create the receipt')
        return {'status': 'ok', 'capability': 'pd_one_item_number_xlsx_compiler', 'action': 'smoke', 'route': route, 'fulfillment_link': fulfillment_link, 'source': str(source), 'workdir': str(workdir), 'template': str(template), 'manifest_path': str(manifest_path), 'migration': migration, 'approval_targets': route_targets, 'workflow_receipt': workflow_receipt, 'compile': {'output_path': str(output_path), 'receipt_path': str(receipt_path), 'receipt': compile_result.receipt}, 'safety': {'shell_used': False, 'source_system_writes_forbidden': True, 'tempfile_output_only': True, 'fixed_source': True, 'fixed_dummy_approval': True, 'route_complete_dummy_approvals': True}}

def cmd_item_number_xlsx_compiler(args: argparse.Namespace) -> None:
    action = str(getattr(args, 'action', 'smoke') or 'smoke').strip()
    if action == 'check':
        result = item_number_xlsx_compiler_check()
    elif action == 'smoke':
        try:
            result = item_number_xlsx_compiler_smoke()
        except Exception as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            fail(f'item-number xlsx compiler smoke failed: {exc!r}')
    else:
        fail('action must be one of: check, smoke')
    emit(result)
    if result.get('status') != 'ok':
        raise SystemExit(1)

def cmd_doc_detect_format(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    suffix, extracted = extract_by_suffix(path)
    classification = classify_company_announcement(path, extracted)
    result = {'status': 'ok', 'capability': 'pd_one_doc_detect_format', 'input': str(path), 'type': suffix, 'classification': classification}
    if args.include_text:
        result['extracted_text_preview'] = extracted_plain_text(extracted)[:4000]
    emit(result)

def render_legacy_xls_summary_pdf(source: Path, out: Path, *, lo_error: dict[str, Any] | None=None) -> dict[str, Any]:
    analysis = inspect_legacy_xls(source, max_sheets=6, preview_rows=12, preview_cols=8, max_nonempty_cells=60, max_merged_cells=10)
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'reportlab unavailable for legacy xls PDF fallback: {exc!r}')
    styles = getSampleStyleSheet()
    code_style = ParagraphStyle('LegacyXlsCode', parent=styles['Code'], fontName='Courier', fontSize=7.5, leading=9, spaceAfter=4)
    story: list[Any] = []
    story.append(Paragraph('Legacy XLS summary PDF fallback', styles['Title']))
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph('This PDF was generated from workbook structure because the direct LibreOffice render could not load the file.', styles['BodyText']))
    if lo_error:
        story.append(Paragraph('LibreOffice attempt', styles['Heading2']))
        story.append(Preformatted(json.dumps(lo_error, ensure_ascii=False, indent=2), code_style))
    story.append(Paragraph('Analysis', styles['Heading2']))
    analysis_summary = {k: analysis.get(k) for k in ('status', 'error', 'blocker', 'next_step', 'message') if analysis.get(k) is not None}
    story.append(Preformatted(json.dumps(analysis_summary, ensure_ascii=False, indent=2), code_style))
    story.append(Paragraph('File metadata', styles['Heading2']))
    story.append(Preformatted(json.dumps(analysis.get('file') or {}, ensure_ascii=False, indent=2), code_style))
    story.append(Paragraph('OLE metadata', styles['Heading2']))
    story.append(Preformatted(json.dumps(analysis.get('ole') or {}, ensure_ascii=False, indent=2), code_style))
    story.append(Paragraph('Workbook metadata', styles['Heading2']))
    story.append(Preformatted(json.dumps(analysis.get('workbook') or {}, ensure_ascii=False, indent=2), code_style))
    sheets = analysis.get('sheets') or []
    for index, sheet in enumerate(sheets):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(f'Sheet {sheet.get('index')} — {sheet.get('name')}', styles['Heading2']))
        story.append(Preformatted(json.dumps(sheet, ensure_ascii=False, indent=2), code_style))
    doc = SimpleDocTemplate(str(out), pagesize=landscape(A4), leftMargin=1.0 * cm, rightMargin=1.0 * cm, topMargin=1.0 * cm, bottomMargin=1.0 * cm)
    doc.build(story)
    return {'render_mode': 'legacy_xls_summary_pdf_fallback', 'analysis_status': analysis.get('status'), 'sheet_count': len(sheets), 'analysis': analysis_summary}

def inspect_office_toolchain() -> dict[str, Any]:
    toolchain = {name: find_bin(name) for name in OFFICE_TOOLCHAIN_BINARIES}
    available_tools = [name for name, path in toolchain.items() if path]
    missing_tools = [name for name, path in toolchain.items() if not path]
    preferred_renderer = toolchain.get('libreoffice') or toolchain.get('soffice')
    preferred_spreadsheet_converter = toolchain.get('ssconvert') or toolchain.get('gnumeric')
    return {'status': 'ok', 'capability': 'pd_one_office_toolchain_status', 'toolchain': toolchain, 'available_tools': available_tools, 'missing_tools': missing_tools, 'preferred_renderer': preferred_renderer, 'preferred_spreadsheet_converter': preferred_spreadsheet_converter, 'python_runtime': sys.executable, 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

def inspect_office_deb_contents(package_path: Path) -> dict[str, Any]:
    """Inspect bounded Office Debian-package paths without an interpreter shell."""
    package = package_path.resolve(strict=False)
    if not is_within(package, OFFICE_DEB_ALLOWED_ROOT()):
        fail(f'Debian package is outside the approved Office artifact root: {package}')
    if package.suffix.lower() != '.deb' or not package.is_file():
        fail('package must be an existing .deb file under the approved Office artifact root')
    dpkg_deb = find_bin('dpkg-deb')
    if not dpkg_deb:
        return {'status': 'blocked', 'capability': 'pd_one_office_deb_contents', 'package': str(package), 'blocker': 'dpkg-deb_not_found', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    proc = run_capture([dpkg_deb, '--contents', str(package)], timeout=30)
    all_lines = proc.stdout.splitlines()
    lines = all_lines[:OFFICE_DEB_CONTENT_MAX_LINES]
    matches = [line for line in lines if any((marker in line.lower() for marker in ('registry', '.xcd', 'calc')))]
    return {'status': 'ok' if proc.returncode == 0 else 'error', 'capability': 'pd_one_office_deb_contents', 'package': str(package), 'returncode': proc.returncode, 'matched_paths': matches, 'matched_count': len(matches), 'content_lines_inspected': len(lines), 'truncated': len(all_lines) > OFFICE_DEB_CONTENT_MAX_LINES, 'stderr': tail_text(proc.stderr.strip()), 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

def cmd_office_deb_contents(args: argparse.Namespace) -> None:
    emit(inspect_office_deb_contents(Path(args.package)))

def inspect_item_number_render_qa_status() -> dict[str, Any]:
    toolchain = inspect_office_toolchain()
    evidence_files = {'qa_artifacts': ITEM_NUMBER_RENDER_QA_ARTIFACTS(), 'render_results': ITEM_NUMBER_RENDER_QA_RESULTS(), 'test_report': ITEM_NUMBER_RENDER_QA_TEST_REPORT(), 'admin_packet': ITEM_NUMBER_RENDER_QA_ADMIN_PACKET(), 'template': ITEM_NUMBER_RENDER_QA_TEMPLATE()}
    missing = [name for name, path in evidence_files.items() if not path.exists()]
    if missing:
        return {'status': 'blocked', 'capability': 'pd_one_item_number_render_qa_status', 'qa_dir': str(ITEM_NUMBER_RENDER_QA_ROOT()), 'missing_evidence': missing, 'evidence_files': {name: str(path) for name, path in evidence_files.items()}, 'toolchain': toolchain, 'blocker': 'missing_render_qa_receipts', 'next_step': 'Recreate the fixed item-number workflow render receipts with the registered PD One local LibreOffice Calc environment, then retry the status wrapper.', 'display_lines': [f'qa_dir={ITEM_NUMBER_RENDER_QA_ROOT()}', f'missing_evidence={', '.join(missing)}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    if ITEM_NUMBER_LOCAL_QA_RECEIPT().is_file():
        try:
            local_receipt = json.loads(ITEM_NUMBER_LOCAL_QA_RECEIPT().read_text(encoding='utf-8'))
        except Exception as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            local_receipt = {'status': 'error', 'error': f'local receipt parse failed: {exc}'}
        local_fixtures = local_receipt.get('fixtures') if isinstance(local_receipt.get('fixtures'), dict) else {}
        local_ok = local_receipt.get('status') == 'ok' and set(ITEM_NUMBER_LOCAL_QA_INPUTS()).issubset(local_fixtures) and all((isinstance(value, dict) and int(value.get('pdf_size_bytes') or 0) > 0 and (int(value.get('page_count') or 0) > 0) and (int(value.get('png_count') or 0) == int(value.get('page_count') or 0)) for value in local_fixtures.values()))
        return {'status': 'ok' if local_ok else 'blocked', 'capability': 'pd_one_item_number_render_qa_status', 'qa_dir': str(ITEM_NUMBER_RENDER_QA_ROOT()), 'local_qa_receipt': str(ITEM_NUMBER_LOCAL_QA_RECEIPT()), 'office_environment_id': local_receipt.get('office_environment_id'), 'renderer': local_receipt.get('renderer'), 'fixtures': local_fixtures, 'production_qualified': False, 'blocker': None if local_ok else 'local_libreoffice_render_qa_incomplete', 'next_step': 'Keep this receipt pilot-only; production qualification remains separately gated.' if local_ok else 'Re-run the fixed typed local LibreOffice render-QA capability and inspect its receipt.', 'toolchain': toolchain, 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    try:
        qa_data = json.loads(ITEM_NUMBER_RENDER_QA_ARTIFACTS().read_text(encoding='utf-8'))
        render_data = json.loads(ITEM_NUMBER_RENDER_QA_RESULTS().read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'status': 'blocked', 'capability': 'pd_one_item_number_render_qa_status', 'qa_dir': str(ITEM_NUMBER_RENDER_QA_ROOT()), 'toolchain': toolchain, 'error': f'receipt parse failed: {exc}', 'blocker': 'receipt_parse_failed', 'next_step': 'Recreate the qa-artifacts-final receipts with the registered PD One local LibreOffice Calc environment and retry the status wrapper.', 'display_lines': [f'qa_dir={ITEM_NUMBER_RENDER_QA_ROOT()}', f'error={exc}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    outputs = qa_data.get('outputs') if isinstance(qa_data.get('outputs'), dict) else {}
    render_routes = render_data.get('routes') if isinstance(render_data.get('routes'), dict) else {}
    migration = qa_data.get('migration') if isinstance(qa_data.get('migration'), dict) else {}
    route_order = [name for name in ('procurement', 'production', 'special') if name in outputs]
    route_order.extend(sorted((name for name in outputs if name not in route_order)))
    routes: dict[str, Any] = {}
    all_production_qualified = True
    all_rendered = True
    for route_name in route_order:
        route_data = outputs.get(route_name) if isinstance(outputs.get(route_name), dict) else {}
        render_route = render_routes.get(route_name) if isinstance(render_routes.get(route_name), dict) else {}
        route_production_qualified = bool(route_data.get('production_qualified'))
        route_pdf = render_route.get('pdf')
        route_page_count = render_route.get('page_count')
        all_production_qualified = all_production_qualified and route_production_qualified
        all_rendered = all_rendered and bool(route_pdf) and (route_page_count not in (None, ''))
        routes[route_name] = {'xlsx': route_data.get('xlsx'), 'receipt': route_data.get('receipt'), 'sha256': route_data.get('sha256'), 'finance_blank': route_data.get('finance_blank'), 'production_qualified': route_production_qualified, 'render': {'returncode': render_route.get('returncode'), 'pdf': route_pdf, 'pdf_size': render_route.get('pdf_size'), 'page_count': route_page_count, 'png_count': len(render_route.get('pngs') or []), 'output_tail': tail_text(str(render_route.get('output') or ''), 300)}}
    status = 'ok' if all_production_qualified and all_rendered else 'blocked'
    blocker = None if status == 'ok' else 'registered_local_libreoffice_render_qa_required'
    next_step = 'Call certify-render with the passed receipt and the required render-certification fields.' if status == 'ok' else 'Use the registered PD One local LibreOffice Calc environment to render the fixed qa-artifacts-final template, visually inspect every page, then call certify-render with exact template/source/PDF hashes, office_environment_id, renderer version/path, certifier_user_id, certified_at, page_count, and passed visual checks.'
    return {'status': status, 'capability': 'pd_one_item_number_render_qa_status', 'qa_dir': str(ITEM_NUMBER_RENDER_QA_ROOT()), 'source': {'path': migration.get('source_path') or str(ITEM_NUMBER_SOURCE_XLS()), 'sha256': qa_data.get('source_sha256') or migration.get('source_sha256')}, 'template': {'path': qa_data.get('template') or str(ITEM_NUMBER_RENDER_QA_TEMPLATE()), 'sha256': migration.get('template_sha256'), 'manifest_path': migration.get('manifest_path')}, 'toolchain': toolchain, 'evidence_files': {name: str(path) for name, path in evidence_files.items()}, 'routes': routes, 'blocker': blocker, 'next_step': next_step, 'display_lines': [f'qa_dir={ITEM_NUMBER_RENDER_QA_ROOT()}', f'status={status}', f'blocker={blocker or 'none'}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

def _ecrn_approval_pilot_status_blocked(reason: str, next_step: str, *, detail: str='', root_id: str='') -> dict[str, Any]:
    return {'status': 'blocked', 'capability': 'pd_one_ecrn_approval_pilot_status', 'state_path': str(ECRN_APPROVAL_PILOT_STATE()), 'source_script': str(ECRN_APPROVAL_PILOT_SCRIPT()), 'root_id': root_id, 'blocker': reason, 'detail': detail, 'next_step': next_step, 'display_lines': [f'state={ECRN_APPROVAL_PILOT_STATE()}', f'blocker={reason}', f'detail={detail or 'none'}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

def _summarize_ecrn_approval_pilot_preview(root_id: str, thread: dict[str, Any], source_file_id: str, receipt: Any) -> dict[str, Any] | None:
    if not isinstance(receipt, dict):
        return None
    preview_post_id = str(receipt.get('preview_post_id') or receipt.get('deleted_preview_post_id') or '')
    preview_file_id = str(receipt.get('preview_file_id') or '')
    if not preview_post_id and (not preview_file_id):
        return None
    return {'root_id': root_id, 'ecrn_no': str(thread.get('ecrn_no') or ''), 'stage': str(thread.get('stage') or ''), 'source_file_id': source_file_id, 'source_name': str(receipt.get('source_name') or ''), 'source_extension': str(receipt.get('source_extension') or ''), 'preview_post_id': preview_post_id, 'deleted_preview_post_id': str(receipt.get('deleted_preview_post_id') or ''), 'preview_file_id': preview_file_id, 'preview_name': str(receipt.get('preview_name') or ''), 'status': str(receipt.get('status') or ''), 'attempts': receipt.get('attempts'), 'posted_at_utc': str(receipt.get('posted_at_utc') or ''), 'deleted_at_utc': str(receipt.get('deleted_at_utc') or ''), 'deletion_reason': str(receipt.get('deletion_reason') or ''), 'deletion_verified': bool(receipt.get('deletion_verified')), 'deletion_reconciled': bool(receipt.get('deletion_reconciled')), 'deletion_last_readback': str(receipt.get('deletion_last_readback') or ''), 'deletion_last_readback_at_utc': str(receipt.get('deletion_last_readback_at_utc') or ''), 'source_post_id': str(receipt.get('source_post_id') or '')}

def _summarize_ecrn_approval_pilot_thread(root_id: str, thread: dict[str, Any], *, max_previews: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    approvals = thread.get('approvals') if isinstance(thread.get('approvals'), dict) else {}
    preview_pdfs = thread.get('preview_pdfs') if isinstance(thread.get('preview_pdfs'), dict) else {}
    all_required_complete_posted = thread.get('all_required_complete_posted') if isinstance(thread.get('all_required_complete_posted'), dict) else {}
    preview_receipts: list[dict[str, Any]] = []
    for source_file_id, receipt in sorted(preview_pdfs.items())[:max_previews]:
        preview = _summarize_ecrn_approval_pilot_preview(root_id, thread, str(source_file_id), receipt)
        if preview is not None:
            preview_receipts.append(preview)
    deleted_preview_count = 0
    pending_deletion_verification_count = 0
    for receipt in preview_pdfs.values():
        if not isinstance(receipt, dict):
            continue
        status = str(receipt.get('status') or '')
        if status == 'deleted_completed':
            deleted_preview_count += 1
        if status == 'deletion_verification_pending':
            pending_deletion_verification_count += 1
    approval_groups = sorted({group for receipt in approvals.values() if isinstance(receipt, dict) for group in receipt.get('approval_groups') or []})
    complete_post_message = str(all_required_complete_posted.get('message') or '') if all_required_complete_posted else ''
    return ({'root_id': root_id, 'ecrn_no': str(thread.get('ecrn_no') or ''), 'stage': str(thread.get('stage') or ''), 'routing_blocked': bool(thread.get('routing_blocked')), 'signed_document_posted': bool(thread.get('signed_document_posted')), 'final_document_stage': str(thread.get('final_document_stage') or ''), 'approval_count': len(approvals), 'approval_groups': approval_groups, 'preview_count': len(preview_pdfs), 'deleted_preview_count': deleted_preview_count, 'pending_deletion_verification_count': pending_deletion_verification_count, 'all_required_complete_posted': bool(all_required_complete_posted), 'all_required_complete_posted_id': str(all_required_complete_posted.get('id') or ''), 'all_required_complete_posted_reply_count': all_required_complete_posted.get('reply_count'), 'all_required_complete_posted_excerpt': tail_text(complete_post_message, 160), 'display_lines': [f'root_id={root_id}', f'stage={thread.get('stage') or ''}', f'preview_count={len(preview_pdfs)} signed_document_posted={bool(thread.get('signed_document_posted'))}']}, preview_receipts)

def inspect_ecrn_approval_pilot_status(root_id: str | None=None, max_threads: int=20, max_previews: int=5) -> dict[str, Any]:
    if not ECRN_APPROVAL_PILOT_STATE().exists():
        return _ecrn_approval_pilot_status_blocked('state_file_missing', 'Run the deterministic ECRN approval pilot once so the state file is created, then retry the status wrapper.')
    try:
        state = json.loads(ECRN_APPROVAL_PILOT_STATE().read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return _ecrn_approval_pilot_status_blocked('state_parse_failed', 'Re-read the ECRN approval-pilot state after the JSON issue is fixed, then retry the status wrapper.', detail=str(exc))
    if not isinstance(state, dict):
        return _ecrn_approval_pilot_status_blocked('state_invalid', 'Regenerate the ECRN approval-pilot state JSON and retry the status wrapper.', detail=f'unexpected type: {type(state).__name__}')
    threads = state.get('threads') if isinstance(state.get('threads'), dict) else None
    if threads is None:
        return _ecrn_approval_pilot_status_blocked('threads_missing', 'Refresh the ECRN approval-pilot state JSON so it includes a threads object, then retry the status wrapper.')
    normalized_root_id = str(root_id or '').strip()
    if normalized_root_id:
        selected_items = [(normalized_root_id, threads.get(normalized_root_id))]
        if selected_items[0][1] is None:
            return _ecrn_approval_pilot_status_blocked('thread_not_found', 'Use an existing thread root id from the ECRN approval-pilot state and retry the status wrapper.', root_id=normalized_root_id)
    else:
        selected_items = sorted(threads.items())[:max_threads]
    thread_summaries: list[dict[str, Any]] = []
    preview_receipts: list[dict[str, Any]] = []
    for selected_root_id, thread in selected_items:
        if not isinstance(thread, dict):
            continue
        thread_summary, previews = _summarize_ecrn_approval_pilot_thread(selected_root_id, thread, max_previews=max_previews)
        thread_summaries.append(thread_summary)
        preview_receipts.extend(previews)
    thread_stage_counts = dict(collections.Counter((str(thread.get('stage') or 'unknown') for thread in threads.values() if isinstance(thread, dict))))
    preview_status_counts = dict(collections.Counter((str(receipt.get('status') or 'unknown') for thread in threads.values() if isinstance(thread, dict) for receipt in (thread.get('preview_pdfs') or {}).values() if isinstance(receipt, dict))))
    threads_with_preview_pdfs = sum((1 for thread in threads.values() if isinstance(thread, dict) and bool(thread.get('preview_pdfs'))))
    threads_with_signed_document_posted = sum((1 for thread in threads.values() if isinstance(thread, dict) and bool(thread.get('signed_document_posted'))))
    threads_with_complete_post = sum((1 for thread in threads.values() if isinstance(thread, dict) and bool(thread.get('all_required_complete_posted'))))
    return {'status': 'ok', 'capability': 'pd_one_ecrn_approval_pilot_status', 'state_path': str(ECRN_APPROVAL_PILOT_STATE()), 'source_script': str(ECRN_APPROVAL_PILOT_SCRIPT()), 'root_id': normalized_root_id, 'policy': {'office_preview_delete_when_complete': bool(state.get('office_preview_delete_when_complete')), 'office_preview_new_threads_after_ms': state.get('office_preview_new_threads_after_ms'), 'office_preview_retention_enabled_at_utc': state.get('office_preview_retention_enabled_at_utc')}, 'thread_count': len(threads), 'selected_thread_count': len(thread_summaries), 'selected_root_ids': [summary.get('root_id') for summary in thread_summaries], 'threads_with_preview_pdfs': threads_with_preview_pdfs, 'threads_with_signed_document_posted': threads_with_signed_document_posted, 'threads_with_complete_post': threads_with_complete_post, 'thread_stage_counts': thread_stage_counts, 'preview_receipt_count': len(preview_receipts), 'preview_status_counts': preview_status_counts, 'threads': thread_summaries, 'preview_receipts': preview_receipts, 'display_lines': [f'state={ECRN_APPROVAL_PILOT_STATE()}', f'threads={len(threads)} selected={len(thread_summaries)} preview_receipts={len(preview_receipts)}', f'signed_document_posted={threads_with_signed_document_posted} preview_threads={threads_with_preview_pdfs}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

def inspect_timezone_snapshot() -> dict[str, Any]:
    snapshot_utc = datetime.now(timezone.utc).replace(microsecond=0)
    snapshot_taipei = snapshot_utc.astimezone(ZoneInfo('Asia/Taipei'))
    utc_text = snapshot_utc.isoformat().replace('+00:00', 'Z')
    taipei_text = snapshot_taipei.isoformat()
    return {'status': 'ok', 'capability': 'pd_one_timezone_snapshot', 'utc_now': utc_text, 'asia_taipei_now': taipei_text, 'display_lines': [utc_text, taipei_text], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}

class _LoopbackOnlyRedirectHandler(urllib.request.HTTPRedirectHandler):

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        parsed = urlparse(newurl)
        host = (parsed.hostname or '').lower()
        port = parsed.port or (443 if parsed.scheme == 'https' else 80)
        if parsed.scheme != 'http' or host not in {'127.0.0.1', 'localhost'} or port != PROTECTED_PORTAL_ORIGIN_PORT:
            raise urllib.error.URLError(f'refusing off-origin redirect to {parsed.scheme}://{host}:{port}')
        return super().redirect_request(req, fp, code, msg, headers, newurl)

def _protected_portal_allowed_paths() -> set[str]:
    allowed = {'/'}
    if not PROTECTED_PORTAL_MANIFEST().is_file():
        allowed.add('/qa/')
        return allowed
    try:
        payload = json.loads(PROTECTED_PORTAL_MANIFEST().read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        allowed.add('/qa/')
        return allowed
    pages = payload.get('pages') if isinstance(payload, dict) else payload
    if isinstance(pages, list):
        for page in pages:
            if not isinstance(page, dict):
                continue
            path = str(page.get('path') or '').strip()
            if path.startswith('/') and '://' not in path and ('..' not in path):
                allowed.add(path.split('?', 1)[0])
                allowed.add(path)
            for tab in page.get('tabs') or []:
                if not isinstance(tab, dict):
                    continue
                tab_path = str(tab.get('path') or '').strip()
                if tab_path.startswith('/') and '://' not in tab_path and ('..' not in tab_path):
                    allowed.add(tab_path.split('?', 1)[0])
                    allowed.add(tab_path)
    allowed.add('/molding-staff/api/production-dashboard')
    allowed.add('/mold-repair/api/queue')
    allowed.add('/warehouse/snapshot.json')
    if '/qa/' not in allowed:
        allowed.add('/qa/')
    return allowed

def _protected_portal_fetch(path: str, email: str | None) -> dict[str, Any]:
    url = f'http://{PROTECTED_PORTAL_ORIGIN_HOST}:{PROTECTED_PORTAL_ORIGIN_PORT}{path}'
    request = urllib.request.Request(url, method='GET')
    if email:
        request.add_header('Cf-Access-Authenticated-User-Email', email)
    opener = urllib.request.build_opener(_LoopbackOnlyRedirectHandler)
    status: int | None = None
    body = ''
    content_type = ''
    error = None
    truncated = False
    try:
        try:
            with opener.open(request, timeout=PROTECTED_PORTAL_ORIGIN_TIMEOUT_SEC) as response:
                status = int(getattr(response, 'status', 200) or 200)
                content_type = str(response.headers.get('Content-Type') or '')[:80]
                raw = response.read(PROTECTED_PORTAL_ORIGIN_BODY_CAP + 1)
        except urllib.error.HTTPError as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            status = int(exc.code)
            content_type = str(exc.headers.get('Content-Type') or '')[:80] if exc.headers else ''
            raw = exc.read(PROTECTED_PORTAL_ORIGIN_BODY_CAP + 1)
        if len(raw) > PROTECTED_PORTAL_ORIGIN_BODY_CAP:
            truncated = True
            raw = raw[:PROTECTED_PORTAL_ORIGIN_BODY_CAP]
        body = raw.decode('utf-8', 'replace')
    except urllib.error.URLError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        error = f'urlerror:{exc.reason}'
    except TimeoutError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        error = 'timeout'
    except OSError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        error = f'oserror:{exc}'
    title_match = re.search('<title>([^<]{0,80})', body, flags=re.I)
    title = (title_match.group(1).strip() if title_match else '')[:PROTECTED_PORTAL_ORIGIN_TITLE_CAP]
    as_of = ''
    as_of_match = re.search('"as_of_date"\\s*:\\s*"([0-9-]{10})"', body)
    if as_of_match:
        as_of = as_of_match.group(1)
    return {'path': path, 'identity': email or 'anonymous', 'status': status, 'body_len': len(body), 'body_truncated': truncated, 'content_type': content_type, 'title': title, 'has_portal_home': 'Portal Home' in body or '入口首頁' in body, 'has_analyses': '"analyses"' in body, 'has_oil_temp_tab': 'data-tab="oil-temp"' in body, 'as_of_date': as_of, 'error': error}

def inspect_protected_portal_origin_probe(extra_probes: list[dict[str, Any]] | None=None) -> dict[str, Any]:
    allowed_paths = _protected_portal_allowed_paths()
    checks: dict[str, Any] = {}
    probes = list(PROTECTED_PORTAL_DEFAULT_PROBES()) + list(extra_probes or [])
    for probe in probes:
        path = str(probe['path'])
        if path not in allowed_paths:
            fail(f'probe path is not allowlisted: {path}')
        checks[str(probe['label'])] = _protected_portal_fetch(path, probe.get('email'))
    display_lines = [f'{label} status={item.get('status')} len={item.get('body_len')} err={item.get('error')}' for label, item in checks.items()]
    return {'status': 'ok', 'capability': 'pd_one_protected_portal_origin_probe', 'origin': f'http://{PROTECTED_PORTAL_ORIGIN_HOST}:{PROTECTED_PORTAL_ORIGIN_PORT}', 'probe_count': len(checks), 'checks': checks, 'display_lines': display_lines, 'safety': {'read_only': True, 'loopback_only': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'body_returned': False, 'public_hostname_forbidden': True, 'service_restart_forbidden': True}}

def inspect_timestamp_to_epoch(timestamp: str) -> dict[str, Any]:
    value = str(timestamp or '').strip()
    if not value:
        fail('timestamp is required')
    if len(value) > 120:
        fail('timestamp must be 120 characters or fewer')
    normalized_input = value[:-1] + '+00:00' if value.endswith(('Z', 'z')) else value
    try:
        parsed = datetime.fromisoformat(normalized_input)
    except ValueError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'timestamp must be a valid ISO-8601 datetime: {exc}')
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        fail('timestamp must include an explicit timezone offset')
    epoch_seconds = parsed.timestamp()
    epoch_milliseconds = int(epoch_seconds * 1000)
    normalized_utc = parsed.astimezone(timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')
    return {'status': 'ok', 'capability': 'pd_one_timestamp_to_epoch', 'input_timestamp': value, 'normalized_utc': normalized_utc, 'epoch_seconds': epoch_seconds, 'epoch_milliseconds': epoch_milliseconds, 'display_lines': [f'input={value}', f'normalized_utc={normalized_utc}', f'epoch_milliseconds={epoch_milliseconds}'], 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'bounded_input': True, 'explicit_timezone_required': True}}

def inspect_epoch_milliseconds_to_iso(epoch_milliseconds: list[int]) -> dict[str, Any]:
    values = list(epoch_milliseconds or [])
    if not values:
        fail('epoch_milliseconds must contain at least one value')
    if len(values) > 20:
        fail('epoch_milliseconds must contain no more than 20 values')
    converted: list[dict[str, Any]] = []
    for raw in values:
        if isinstance(raw, bool) or not isinstance(raw, int):
            fail('epoch_milliseconds values must be integers')
        if raw < 0 or raw > 4102444800000:
            fail('epoch_milliseconds values must be between 0 and 4102444800000')
        converted.append({'epoch_milliseconds': raw, 'utc_iso': datetime.fromtimestamp(raw / 1000, tz=timezone.utc).isoformat()})
    return {'status': 'ok', 'capability': 'pd_one_epoch_milliseconds_to_iso', 'values': converted, 'display_lines': [f'{item['epoch_milliseconds']} {item['utc_iso']}' for item in converted], 'safety': {'read_only': True, 'deterministic': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'bounded_input': True}}
SHEETS_SERIAL_EPOCH = date(1899, 12, 30)
SHEETS_SERIAL_MIN = 0
SHEETS_SERIAL_MAX = 2958465
SHEETS_SERIAL_MAX_ITEMS = 20
SHEETS_COLUMN_MAX_LEN = 4
SHEETS_COLUMN_MAX_INDEX = 18278
SHEETS_COLUMN_MAX_ITEMS = 20
SHEETS_RANGE_MAX_ITEMS = 8

def _a1_column_to_index(letter: str) -> tuple[str, int]:
    raw = str(letter or '').strip().upper()
    if not raw or len(raw) > SHEETS_COLUMN_MAX_LEN or (not raw.isalpha()):
        fail('column letters must be 1-4 A-Z characters')
    index = 0
    for ch in raw:
        index = index * 26 + (ord(ch) - 64)
    if index < 1 or index > SHEETS_COLUMN_MAX_INDEX:
        fail(f'column index must be between 1 and {SHEETS_COLUMN_MAX_INDEX}')
    return (raw, index)

def inspect_sheets_serial_column_convert(serials: list[int] | None=None, columns: list[str] | None=None, ranges: list[str] | None=None) -> dict[str, Any]:
    serial_values = list(serials or [])
    column_values = list(columns or [])
    range_values = list(ranges or [])
    if not serial_values and (not column_values) and (not range_values):
        fail('provide at least one serial, column, or A1 range')
    if len(serial_values) > SHEETS_SERIAL_MAX_ITEMS:
        fail(f'serials must contain no more than {SHEETS_SERIAL_MAX_ITEMS} values')
    if len(column_values) > SHEETS_COLUMN_MAX_ITEMS:
        fail(f'columns must contain no more than {SHEETS_COLUMN_MAX_ITEMS} values')
    if len(range_values) > SHEETS_RANGE_MAX_ITEMS:
        fail(f'ranges must contain no more than {SHEETS_RANGE_MAX_ITEMS} values')
    converted_serials: list[dict[str, Any]] = []
    for raw in serial_values:
        if isinstance(raw, bool) or not isinstance(raw, int):
            fail('serials values must be integers')
        if raw < SHEETS_SERIAL_MIN or raw > SHEETS_SERIAL_MAX:
            fail(f'serials values must be between {SHEETS_SERIAL_MIN} and {SHEETS_SERIAL_MAX}')
        converted_serials.append({'serial': raw, 'iso_date': (SHEETS_SERIAL_EPOCH + timedelta(days=raw)).isoformat()})
    converted_columns: list[dict[str, Any]] = []
    for raw in column_values:
        letter, index = _a1_column_to_index(raw)
        converted_columns.append({'letter': letter, 'index': index})
    converted_ranges: list[dict[str, Any]] = []
    for raw in range_values:
        text = str(raw or '').strip().upper()
        if text.count(':') != 1:
            fail('ranges must use start:end A1 column letters')
        start_letter, end_letter = text.split(':')
        start_letter, start_index = _a1_column_to_index(start_letter)
        end_letter, end_index = _a1_column_to_index(end_letter)
        if end_index < start_index:
            fail('range end column must be at or after the start column')
        converted_ranges.append({'range': f'{start_letter}:{end_letter}', 'start_letter': start_letter, 'start_index': start_index, 'end_letter': end_letter, 'end_index': end_index, 'width': end_index - start_index + 1})
    display_lines = [f'{item['serial']} {item['iso_date']}' for item in converted_serials] + [f'{item['letter']} {item['index']}' for item in converted_columns] + [f'{item['range']} width={item['width']}' for item in converted_ranges]
    return {'status': 'ok', 'capability': 'pd_one_sheets_serial_column_convert', 'epoch': '1899-12-30', 'epoch_name': 'google_sheets', 'serials': converted_serials, 'columns': converted_columns, 'ranges': converted_ranges, 'display_lines': display_lines, 'safety': {'read_only': True, 'deterministic': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'bounded_input': True, 'excel_1900_leap_bug_not_applied': True}}

def cmd_timezone_snapshot(args: argparse.Namespace) -> None:
    emit(inspect_timezone_snapshot())

def cmd_protected_portal_origin_probe(args: argparse.Namespace) -> None:
    extra: list[dict[str, Any]] = []
    paths = list(getattr(args, 'path', None) or [])
    emails = list(getattr(args, 'email', None) or [])
    if paths:
        identities: list[str | None] = [None]
        for email in emails:
            if email and email != '-':
                identities.append(email)
        for path in paths:
            for email in identities:
                label_id = email or 'anon'
                extra.append({'label': f'{label_id}_{path}', 'path': path, 'email': email or None})
    emit(inspect_protected_portal_origin_probe(extra))

def cmd_timestamp_to_epoch(args: argparse.Namespace) -> None:
    emit(inspect_timestamp_to_epoch(args.timestamp))

def cmd_epoch_milliseconds_to_iso(args: argparse.Namespace) -> None:
    emit(inspect_epoch_milliseconds_to_iso(args.epoch_milliseconds))

def cmd_sheets_serial_column_convert(args: argparse.Namespace) -> None:
    emit(inspect_sheets_serial_column_convert(args.serials, args.columns, args.ranges))

def cmd_office_toolchain_status(args: argparse.Namespace) -> None:
    emit(inspect_office_toolchain())

def cmd_item_number_render_qa_status(args: argparse.Namespace) -> None:
    emit(inspect_item_number_render_qa_status())

def cmd_item_number_historical_sample_summary(args: argparse.Namespace) -> None:
    """Run the fixed ten-file historical XLS structural summary safely."""
    sample_dir = ITEM_NUMBER_HISTORICAL_ROOT() / 'sample'
    if not ITEM_NUMBER_HISTORICAL_SCRIPT().is_file():
        fail(f'historical summary script not found: {ITEM_NUMBER_HISTORICAL_SCRIPT()}')
    if not ITEM_NUMBER_HISTORICAL_RECEIPT().is_file():
        fail(f'historical copyout receipt not found: {ITEM_NUMBER_HISTORICAL_RECEIPT()}')
    try:
        receipt = json.loads(ITEM_NUMBER_HISTORICAL_RECEIPT().read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'historical copyout receipt is invalid JSON: {exc}')
    entries = receipt.get('entries', receipt.get('samples', []))
    if not isinstance(entries, list) or len(entries) != 10:
        fail('historical summary requires exactly 10 receipt entries')
    for entry in entries:
        path = Path(str(entry.get('local_path', ''))).resolve(strict=False)
        if not is_within(path, sample_dir) or path.suffix.lower() != '.xls':
            fail(f'historical sample path is outside the fixed sample root: {path}')
        if not path.is_file():
            fail(f'historical sample is missing: {path}')
    proc = subprocess.run([sys.executable, str(ITEM_NUMBER_HISTORICAL_SCRIPT())], cwd=str(ITEM_NUMBER_WORKFLOW_ROOT()), env=safe_env(), text=True, capture_output=True, shell=False, timeout=180)
    result = {'status': 'ok' if proc.returncode == 0 and ITEM_NUMBER_HISTORICAL_OUTPUT().is_file() else 'error', 'capability': 'pd_one_item_number_historical_sample_summary', 'action': 'summarize', 'sample_count': 10, 'script': str(ITEM_NUMBER_HISTORICAL_SCRIPT()), 'receipt': str(ITEM_NUMBER_HISTORICAL_RECEIPT()), 'output': str(ITEM_NUMBER_HISTORICAL_OUTPUT()), 'returncode': proc.returncode, 'stdout_tail': tail_text(proc.stdout), 'stderr_tail': tail_text(proc.stderr), 'safety': {'shell_used': False, 'source_inputs_read_only': True, 'fixed_sample_root': True, 'local_artifact_write_only': True, 'source_system_writes_forbidden': True, 'external_sends_forbidden': True}}
    emit(result)
    if result['status'] != 'ok':
        raise SystemExit(1)

def cmd_item_number_local_fixture_copy(args: argparse.Namespace) -> None:
    """Stage and render the fixed item-number XLS/XLSX QA fixture set locally."""
    renderer = find_bin('libreoffice') or find_bin('soffice')
    pdftoppm = find_bin('pdftoppm')
    pdfinfo = find_bin('pdfinfo')
    if not renderer or not pdftoppm or (not pdfinfo):
        fail('local LibreOffice render QA requires libreoffice, pdftoppm, and pdfinfo')
    ITEM_NUMBER_LOCAL_QA_ROOT().mkdir(parents=True, exist_ok=True)
    fixtures: dict[str, Any] = {}
    for name, source in ITEM_NUMBER_LOCAL_QA_INPUTS().items():
        if not source.is_file():
            fail(f'fixed item-number fixture not found: {source}')
        suffix = source.suffix.lower()
        staged = ITEM_NUMBER_LOCAL_QA_ROOT() / f'{name}{suffix}'
        pdf = ITEM_NUMBER_LOCAL_QA_ROOT() / f'{name}.pdf'
        png_prefix = ITEM_NUMBER_LOCAL_QA_ROOT() / f'{name}-page'
        for stale_png in ITEM_NUMBER_LOCAL_QA_ROOT().glob(f'{name}-page-*.png'):
            stale_png.unlink()
        shutil.copy2(source, staged)
        source_bytes = source.read_bytes()
        staged_bytes = staged.read_bytes()
        if hashlib.sha256(source_bytes).digest() != hashlib.sha256(staged_bytes).digest():
            fail(f'staged fixture hash mismatch: {name}')
        conversion = convert_office_to_pdf(staged, pdf)
        info_proc = run_capture([pdfinfo, str(pdf)], timeout=30)
        page_match = re.search('^Pages:\\s+(\\d+)\\s*$', info_proc.stdout, re.MULTILINE)
        page_count = int(page_match.group(1)) if page_match else 0
        raster_proc = run_capture([pdftoppm, '-png', '-r', '144', str(pdf), str(png_prefix)], timeout=180)
        pngs = sorted(ITEM_NUMBER_LOCAL_QA_ROOT().glob(f'{name}-page-*.png'))
        fixture_ok = info_proc.returncode == 0 and raster_proc.returncode == 0 and (page_count > 0) and (len(pngs) == page_count)
        fixtures[name] = {'status': 'ok' if fixture_ok else 'error', 'source': str(source), 'source_sha256': hashlib.sha256(source_bytes).hexdigest(), 'staged': str(staged), 'staged_sha256': hashlib.sha256(staged_bytes).hexdigest(), 'pdf': str(pdf), 'pdf_sha256': conversion['sha256'], 'pdf_size_bytes': conversion['output_size_bytes'], 'page_count': page_count, 'png_count': len(pngs), 'pngs': [str(path) for path in pngs], 'cjk_visual_check': 'required_on_rasterized_pages', 'renderer': conversion['renderer'], 'render_mode': conversion['render_mode']}
    all_ok = all((value.get('status') == 'ok' for value in fixtures.values()))
    renderer_version = run_capture([renderer, '--version'], timeout=30)
    environment_material = json.dumps({'renderer': renderer, 'version': renderer_version.stdout.strip()}, sort_keys=True).encode('utf-8')
    result = {'status': 'ok' if all_ok else 'error', 'capability': 'pd_one_item_number_local_fixture_copy', 'action': 'stage_and_render_fixed_qa_set', 'office_environment_id': 'sha256:' + hashlib.sha256(environment_material).hexdigest(), 'renderer': renderer, 'renderer_version': renderer_version.stdout.strip(), 'fixtures': fixtures, 'receipt': str(ITEM_NUMBER_LOCAL_QA_RECEIPT()), 'production_qualified': False, 'safety': {'source_read_only': True, 'fixed_source_root': True, 'fixed_local_temp_root': True, 'external_system_writes_forbidden': True, 'external_delivery_forbidden': True, 'shell_used': False}}
    ITEM_NUMBER_LOCAL_QA_RECEIPT().write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    emit(result)
    if not all_ok:
        raise SystemExit(1)

def cmd_item_number_windows_excel_qa(args: argparse.Namespace) -> None:
    """Plan or run the fixed, confirmation-gated Windows Excel/COM QA lane."""
    script = ITEM_NUMBER_WORKFLOW_ROOT() / 'tools/run_windows_excel_qa.py'
    python_bin = Path(resolve_configured_path('~/.hermes/venvs/winrm/bin/python'))
    confirmation = 'RUN ITEM NUMBER WINDOWS EXCEL QA'
    if not script.is_file() or not python_bin.is_file():
        fail(f'fixed Windows Excel QA prerequisites missing: script={script} python={python_bin}')
    if getattr(args, 'confirm', None) != confirmation:
        emit({'status': 'planned', 'capability': 'pd_one_item_number_windows_excel_qa', 'action': 'run', 'command': [str(python_bin), str(script)], 'workdir': str(ITEM_NUMBER_WORKFLOW_ROOT()), 'confirmation_required': confirmation, 'safety': {'read_only_source_inputs': True, 'remote_temp_only': True, 'local_artifact_write': True, 'production_qualification_asserted': False, 'shell_used': False}})
        return
    proc = subprocess.run([str(python_bin), str(script)], cwd=str(ITEM_NUMBER_WORKFLOW_ROOT()), env=safe_env(), text=True, capture_output=True, shell=False, timeout=360)
    result = {'status': 'ok' if proc.returncode == 0 else 'error', 'capability': 'pd_one_item_number_windows_excel_qa', 'action': 'run', 'returncode': proc.returncode, 'stdout_tail': tail_text(proc.stdout, 12000), 'stderr_tail': tail_text(proc.stderr, 4000), 'safety': {'read_only_source_inputs': True, 'remote_temp_only': True, 'local_artifact_write': True, 'production_qualification_asserted': False, 'shell_used': False}}
    emit(result)
    if proc.returncode != 0:
        raise SystemExit(1)

def cmd_ecrn_approval_pilot_status(args: argparse.Namespace) -> None:
    root_id = str(getattr(args, 'root_id', '') or '').strip() or None
    emit(inspect_ecrn_approval_pilot_status(root_id=root_id, max_threads=_bounded_int(getattr(args, 'max_threads', 20), '--max-threads', 1, 50), max_previews=_bounded_int(getattr(args, 'max_previews', 5), '--max-previews', 0, 20)))

def cmd_ecrn_approval_pilot_check(args: argparse.Namespace) -> None:
    """Run the fixed, read-only ECRN approval-pilot QA checks."""
    root = ECRN_CHANNEL_HEADER_TEST_ROOT()
    script_path = ECRN_APPROVAL_PILOT_SCRIPT()
    log_path = root / 'logs/ecrn-approval-pilot.jsonl'
    test_paths = sorted(root.glob('tests/test_ecrn_*.py'))[:50]
    required = [script_path, *test_paths]
    missing = [str(path) for path in required if not path.is_file()]

    def line_count(path: Path) -> int | None:
        if not path.is_file():
            return None
        with path.open('r', encoding='utf-8') as handle:
            return sum((1 for _ in handle))
    before = line_count(log_path)
    checks: dict[str, Any] = {}
    commands = [('pytest', [sys.executable, '-m', 'pytest', '-q', *[str(path.relative_to(root)) for path in test_paths]], 300), ('ruff', ['ruff', 'check', 'tools/ecrn_approval_pilot.py', 'tests/test_ecrn_channel_header.py'], 120), ('py_compile', [sys.executable, '-m', 'py_compile', 'tools/ecrn_approval_pilot.py'], 120)]
    if not missing:
        for name, command, timeout in commands:
            try:
                proc = subprocess.run(command, cwd=str(root), env=safe_env(), text=True, capture_output=True, shell=False, timeout=timeout)
                checks[name] = {'command': command, 'returncode': proc.returncode, 'stdout_tail': tail_text(proc.stdout), 'stderr_tail': tail_text(proc.stderr)}
            except FileNotFoundError as exc:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                checks[name] = {'command': command, 'returncode': 'missing_executable', 'stdout_tail': '', 'stderr_tail': str(exc)}
            except subprocess.TimeoutExpired as exc:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                checks[name] = {'command': command, 'returncode': 'timeout', 'stdout_tail': tail_text(exc.stdout or ''), 'stderr_tail': tail_text(exc.stderr or '')}
    after = line_count(log_path)
    checks_ok = not missing and all((value.get('returncode') == 0 for value in checks.values()))
    result = {'status': 'ok' if checks_ok and before == after else 'error', 'capability': 'pd_one_ecrn_approval_pilot_check', 'action': 'check', 'test_root': str(root), 'script': str(script_path), 'test_files': [str(path) for path in test_paths], 'missing_prerequisites': missing, 'checks': checks, 'operational_log': {'path': str(log_path), 'lines_before': before, 'lines_after': after, 'unchanged': before == after}, 'safety': {'read_only': True, 'shell_used': False, 'fixed_paths_only': True, 'source_system_writes_forbidden': True, 'external_delivery_forbidden': True}}
    emit(result)
    if result['status'] != 'ok':
        raise SystemExit(1)

def _item_number_workflow_component_paths(component: str) -> tuple[str, Path, Path, Path]:
    normalized = str(component or 'engine').strip().lower()
    if normalized == 'engine':
        return (normalized, ITEM_NUMBER_WORKFLOW_ENGINE(), ITEM_NUMBER_WORKFLOW_SCRIPT(), ITEM_NUMBER_WORKFLOW_TEST())
    if normalized == 'compiler':
        return (normalized, ITEM_NUMBER_COMPILER_ROOT(), ITEM_NUMBER_COMPILER_SCRIPT(), ITEM_NUMBER_COMPILER_TEST())
    if normalized == 'runner':
        return (normalized, ITEM_NUMBER_WORKFLOW_RUNNER(), ITEM_NUMBER_WORKFLOW_RUNNER_SCRIPT(), ITEM_NUMBER_WORKFLOW_RUNNER_TEST())
    fail('component must be one of: compiler, engine, runner')

def cmd_item_number_workflow_check(args: argparse.Namespace) -> None:
    """Compile and run the fixed item-number workflow unit tests for engine or compiler."""
    component, component_root, script_path, unit_test_path = _item_number_workflow_component_paths(getattr(args, 'component', 'engine'))
    if not component_root.is_dir():
        fail(f'item-number workflow {component} directory not found: {component_root}')
    if not script_path.is_file():
        fail(f'item-number workflow {component} script not found: {script_path}')
    if not unit_test_path.is_file():
        fail(f'item-number workflow {component} test not found: {unit_test_path}')
    compile_cmd = [sys.executable, '-m', 'py_compile', str(script_path), str(unit_test_path)]
    compile_proc = subprocess.run(compile_cmd, cwd=str(component_root), env=safe_env(), text=True, capture_output=True, shell=False, timeout=120)
    unit_cmd = [sys.executable, '-m', 'unittest', 'discover', '-s', 'tests', '-v']
    unit_proc = subprocess.run(unit_cmd, cwd=str(component_root), env=safe_env(), text=True, capture_output=True, shell=False, timeout=240)
    result = {'status': 'ok' if compile_proc.returncode == 0 and unit_proc.returncode == 0 else 'error', 'capability': 'pd_one_item_number_workflow_check', 'action': 'check', 'component': component, 'component_root': str(component_root), 'script': str(script_path), 'unit_test': str(unit_test_path), 'unit_test_dir': str(component_root / 'tests'), 'unit_command': unit_cmd, 'unit_mode': 'discover', 'compile': {'command': compile_cmd, 'returncode': compile_proc.returncode, 'stdout_tail': tail_text(compile_proc.stdout), 'stderr_tail': tail_text(compile_proc.stderr)}, 'unit': {'command': unit_cmd, 'returncode': unit_proc.returncode, 'stdout_tail': tail_text(unit_proc.stdout), 'stderr_tail': tail_text(unit_proc.stderr)}, 'safety': {'shell_used': False, 'read_only': True, 'fixed_component_root': True, 'source_system_writes_forbidden': True}}
    emit(result)
    if result['status'] != 'ok':
        raise SystemExit(1)

def convert_office_to_pdf(input_path: Path, out: Path) -> dict[str, Any]:
    source = require_allowed_doc_path(input_path)
    suffix = source.suffix.lower()
    if suffix not in OFFICE_TO_PDF_INPUT_SUFFIXES:
        fail('supported office input types: .csv, .doc, .docx, .odp, .ods, .odt, .ppt, .pptx, .rtf, .xls, .xlsm, .xlsx')
    with source.open('rb') as handle:
        source_header = handle.read(8)
    if suffix == '.xls' and source_header != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        fail('declared .xls input is not an OLE Compound File workbook')
    if suffix in {'.xlsx', '.xlsm', '.ods'} and (not source_header.startswith(b'PK')):
        fail(f'declared {suffix} input is not a ZIP-based Office package')
    out = out.resolve(strict=False)
    if not any((is_within(out, root) for root in DOC_ALLOWED_ROOTS())):
        fail(f'output path is outside approved PD One document roots: {out}')
    if out.suffix.lower() != '.pdf':
        fail('office-convert-pdf output must end with .pdf')
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()
    renderer = find_bin('libreoffice') or find_bin('soffice')
    lo_attempt: dict[str, Any] | None = None
    if renderer:
        with tempfile.TemporaryDirectory(prefix='office-convert-', dir=str(out.parent)) as tmp_dir:
            workdir = Path(tmp_dir)
            source_copy = workdir / f'{out.stem}{source.suffix}'
            shutil.copy2(source, source_copy)
            profile = workdir / 'lo-profile'
            profile.mkdir(exist_ok=True)
            cmd = [renderer, f'-env:UserInstallation=file://{profile.resolve()}', '--headless', '--nologo', '--nofirststartwizard', '--convert-to', 'pdf', '--outdir', str(out.parent), str(source_copy)]
            try:
                proc = run_capture(cmd, timeout=180)
                lo_attempt = {'renderer': renderer, 'returncode': proc.returncode, 'stdout': tail_text(proc.stdout.strip()), 'stderr': tail_text(proc.stderr.strip())}
            except subprocess.TimeoutExpired:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                lo_attempt = {'renderer': renderer, 'returncode': 'timeout', 'stdout': '', 'stderr': 'office conversion timed out after 180s'}
    else:
        lo_attempt = {'renderer': None, 'returncode': 'missing', 'stdout': '', 'stderr': 'LibreOffice/soffice not found on PATH'}
    render_mode = 'libreoffice'
    fallback_info: dict[str, Any] | None = None
    if not out.exists() or out.stat().st_size <= 0:
        detail = ''
        if lo_attempt:
            detail = f'; last LibreOffice attempt: {tail_text(json.dumps(lo_attempt, ensure_ascii=False, sort_keys=True))}'
        if out.exists() and out.stat().st_size <= 0:
            fail(f'office conversion produced an empty pdf: {out}{detail}')
        fail(f'office conversion did not produce expected pdf: {out}{detail}')
    pdf_digest = hashlib.sha256(out.read_bytes()).hexdigest()
    try:
        pdf_validation_raw = extract_pdf(out)
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'pdf validation failed: {exc!r}')
    if pdf_validation_raw.get('error'):
        pdf_validation = {'status': 'skipped', 'reason': pdf_validation_raw['error']}
    else:
        pages = pdf_validation_raw.get('pages') or []
        first_page_preview = ''
        if pages and isinstance(pages[0], dict):
            first_page_preview = str(pages[0].get('text') or '')[:1000]
        pdf_validation = {'status': 'ok', 'page_count': pdf_validation_raw.get('page_count'), 'first_page_text_preview': first_page_preview}
    result: dict[str, Any] = {'status': 'ok', 'capability': 'pd_one_office_convert_to_pdf', 'input': str(source), 'output': str(out), 'source_suffix': suffix, 'renderer': renderer, 'render_mode': render_mode, 'output_size_bytes': out.stat().st_size, 'sha256': pdf_digest, 'conversion': lo_attempt or {}, 'pdf_validation': pdf_validation}
    if fallback_info:
        result['fallback'] = fallback_info
    return result

def cmd_office_convert_pdf(args: argparse.Namespace) -> None:
    result = convert_office_to_pdf(Path(args.input), Path(args.out))
    emit(result)

def _require_str(value: Any, label: str) -> str:
    if not isinstance(value, str):
        fail(f'{label} must be a string')
    return value

def _require_list(value: Any, label: str) -> list[Any]:
    if not isinstance(value, list):
        fail(f'{label} must be a list')
    return value

def _apply_table(table, rows: list[list[Any]], *, header_rows: int=0) -> None:
    for row_idx, row in enumerate(rows):
        for col_idx, cell_value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = '' if cell_value is None else str(cell_value)
            if row_idx < header_rows:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True

def build_html_from_paragraph_spec(spec_path: Path, out: Path) -> dict[str, Any]:
    """Render one bounded JSON ``paragraphs`` spec to a local HTML artifact."""
    source = require_allowed_doc_path(spec_path)
    out = out.resolve(strict=False)
    if source.suffix.lower() != '.json':
        fail('json-html-artifact input must be .json')
    if out.suffix.lower() not in {'.html', '.htm'}:
        fail('json-html-artifact output must be .html or .htm')
    if not any((is_within(out, root) for root in DOC_ALLOWED_ROOTS())):
        fail(f'output path is outside approved roots: {out}')
    if source.stat().st_size > JSON_HTML_ARTIFACT_MAX_BYTES:
        fail(f'JSON spec exceeds {JSON_HTML_ARTIFACT_MAX_BYTES} bytes')
    try:
        spec = json.loads(source.read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'failed to load JSON spec: {exc!r}')
    if not isinstance(spec, dict) or not isinstance(spec.get('paragraphs'), list):
        fail('spec must be an object with a paragraphs list')
    paragraphs = spec['paragraphs']
    if len(paragraphs) > JSON_HTML_ARTIFACT_MAX_PARAGRAPHS:
        fail('paragraphs exceeds safe limit')
    parts = ['<!doctype html><html><head><meta charset="utf-8"><style>', '@page { size: A4; margin: 2cm; } body { font-family: Arial, sans-serif; font-size: 11pt; line-height: 1.45; color:#111; }', 'h1 { text-align:center; font-size:16pt; margin:0 0 6pt; } h2 { text-align:center; font-size:15pt; margin:0 0 6pt; }', '.version { text-align:center; margin:0 0 18pt; } h3 { font-size:12pt; margin:12pt 0 4pt; page-break-after:avoid; }', 'p { text-align:justify; margin:0 0 6pt; } .item { margin-left:18pt; text-indent:-18pt; }', '</style></head><body>']
    for index, paragraph in enumerate(paragraphs):
        if not isinstance(paragraph, str):
            fail(f'paragraphs[{index}] must be a string')
        if len(paragraph) > 20000:
            fail(f'paragraphs[{index}] exceeds 20000 characters')
        escaped = html.escape(paragraph)
        if index == 0:
            parts.append(f'<h1>{escaped}</h1>')
        elif index == 1:
            parts.append(f'<h2>{escaped}</h2>')
        elif index == 2:
            parts.append(f'<p class="version">{escaped}</p>')
        elif paragraph.startswith('Article '):
            parts.append(f'<h3>{escaped}</h3>')
        elif not paragraph:
            parts.append('<p>&nbsp;</p>')
        elif paragraph.startswith(('1. ', '2. ', '3. ', '4. ', '5. ', '6. ', '(1) ', '(2) ', '(3) ', '(4) ')):
            parts.append(f'<p class="item">{escaped}</p>')
        else:
            parts.append(f'<p>{escaped}</p>')
    parts.append('</body></html>')
    rendered = ''.join(parts)
    encoded = rendered.encode('utf-8')
    if len(encoded) > JSON_HTML_ARTIFACT_MAX_BYTES:
        fail(f'rendered HTML exceeds {JSON_HTML_ARTIFACT_MAX_BYTES} bytes')
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f'.{out.name}.', suffix='.tmp', dir=str(out.parent), delete=False) as handle:
        temp_path = Path(handle.name)
        handle.write(encoded)
    try:
        temp_path.replace(out)
    finally:
        temp_path.unlink(missing_ok=True)
    return {'status': 'ok', 'capability': 'pd_one_json_html_artifact_build', 'spec': str(source), 'output': str(out), 'paragraphs': len(paragraphs), 'output_bytes': len(encoded), 'output_sha256': hashlib.sha256(encoded).hexdigest(), 'safety': {'source_read_only': True, 'local_artifact_write_only': True, 'source_system_writes': False, 'external_delivery': False}}

def cmd_json_html_artifact(args: argparse.Namespace) -> None:
    emit(build_html_from_paragraph_spec(Path(args.spec), Path(args.out)))

def _iter_docx_table_paragraphs(table: Any) -> Any:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _iter_docx_table_paragraphs(nested_table)

def _iter_docx_story_paragraphs(doc: Any) -> list[Any]:
    paragraphs: list[Any] = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(_iter_docx_table_paragraphs(table))
    for section in doc.sections:
        for story in (section.header, section.footer):
            paragraphs.extend(story.paragraphs)
            for table in story.tables:
                paragraphs.extend(_iter_docx_table_paragraphs(table))
    deduped: list[Any] = []
    seen: set[int] = set()
    for paragraph in paragraphs:
        marker = id(paragraph._p)
        if marker not in seen:
            deduped.append(paragraph)
            seen.add(marker)
    return deduped

def inspect_or_stage_docx_layout(*, action: str, input_path: Path, target_text: str, output_path: Path | None=None, source_sha256: str | None=None, overwrite: bool=False) -> dict[str, Any]:
    before_hash = _PdUnset
    doc = _PdUnset
    output = _PdUnset
    receipt = _PdUnset
    source = _PdUnset
    target = _PdUnset
    temp_path = _PdUnset

    def _pd_phase_13():
        nonlocal action, before_hash, doc, output, receipt, source, target, temp_path
        'Inspect or stage one exact DOCX section-start layout adjustment.'
        action = str(action or 'inspect').strip().lower()
        if action not in {'inspect', 'stage'}:
            raise ValueError('action must be one of: inspect, stage')
        source = require_allowed_doc_path(Path(input_path))
        if source.suffix.lower() != '.docx':
            raise ValueError('docx-layout-patch only supports .docx files')
        if source.stat().st_size > DOCX_INSPECT_MAX_BYTES:
            raise ValueError(f'DOCX exceeds the {DOCX_INSPECT_MAX_BYTES} byte layout-patch limit')
        target = str(target_text or '')
        if not target or len(target) > DOCX_LAYOUT_TARGET_MAX_CHARS:
            raise ValueError(f'target_text must contain 1-{DOCX_LAYOUT_TARGET_MAX_CHARS} characters')
        before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        if source_sha256 is not None:
            expected_hash = str(source_sha256).strip().lower()
            if not re.fullmatch('[0-9a-f]{64}', expected_hash):
                raise ValueError('source_sha256 must be a 64-character lowercase hexadecimal digest')
            if expected_hash != before_hash:
                raise ValueError('source SHA-256 guard does not match the current DOCX')
        try:
            from docx import Document
        except Exception as exc:
            raise RuntimeError(f'python-docx unavailable: {exc}') from exc
        doc = Document(str(source))
        paragraphs = _iter_docx_story_paragraphs(doc)
        matches = [paragraph for paragraph in paragraphs if paragraph.text == target]
        if len(matches) != 1:
            raise ValueError(f'target_text must match exactly one visible paragraph; matches={len(matches)}')
        paragraph = matches[0]
        before_layout = {'page_break_before': paragraph.paragraph_format.page_break_before, 'keep_with_next': paragraph.paragraph_format.keep_with_next}
        receipt = {'status': 'ok', 'capability': 'pd_one_docx_layout_patch', 'action': action, 'input': str(source), 'source_sha256': before_hash, 'target_text': target, 'match_count': 1, 'before': before_layout, 'requested': {'page_break_before': True, 'keep_with_next': True}, 'safety': {'admin_only': True, 'source_read_only': True, 'separate_staged_output_required': True, 'approved_document_roots_only': True, 'exact_unique_visible_match_required': True, 'office_integrity_gate_required_for_stage': True, 'source_system_writes': False, 'external_delivery': False, 'arbitrary_openxml_forbidden': True}}
        if action == 'inspect':
            receipt['would_change'] = before_layout != {'page_break_before': True, 'keep_with_next': True}
            return _PdReturn(receipt)
        if output_path is None:
            raise ValueError('output_path is required for action=stage')
        output = Path(output_path).resolve(strict=False)
        if output.suffix.lower() != '.docx':
            raise ValueError('output_path must end in .docx')
        if not any((is_within(output, root) for root in DOC_ALLOWED_ROOTS())):
            raise ValueError(f'output_path is outside approved PD One document roots: {output}')
        if output == source:
            raise ValueError('source overwrite is forbidden; choose a separate staged output_path')
        if output.exists() and (not overwrite):
            raise ValueError('output_path already exists; pass overwrite only for the separate staged output')
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.keep_with_next = True
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix=f'.{output.name}.', suffix='.docx', dir=str(output.parent), delete=False) as handle:
            temp_path = Path(handle.name)

    def _pd_phase_14():
        try:
            doc.save(str(temp_path))
            gate_script = Path(resolve_configured_path('~/.hermes/skills/pd-one/pd-one-admin/scripts/office_no_corruption_gate.py'))
            gate_result = {'status': 'unavailable'}
            if gate_script.exists():
                proc = subprocess.run([sys.executable, str(gate_script), str(temp_path)], text=True, capture_output=True, env=safe_env(), shell=False, timeout=180)
                gate_result = {'status': 'ok' if proc.returncode == 0 else 'error', 'returncode': proc.returncode}
                if proc.returncode != 0:
                    raise RuntimeError(f'office no-corruption gate failed: {proc.stderr.strip()[:1200] or proc.stdout.strip()[:1200]}')
            verify_doc = Document(str(temp_path))
            verify_matches = [item for item in _iter_docx_story_paragraphs(verify_doc) if item.text == target]
            if len(verify_matches) != 1:
                raise RuntimeError('staged DOCX target readback was not unique')
            verify_format = verify_matches[0].paragraph_format
            if verify_format.page_break_before is not True or verify_format.keep_with_next is not True:
                raise RuntimeError('staged DOCX layout readback did not preserve the requested properties')
            if hashlib.sha256(source.read_bytes()).hexdigest() != before_hash:
                raise RuntimeError('source DOCX changed during staged layout patch')
            temp_path.replace(output)
        finally:
            temp_path.unlink(missing_ok=True)
        receipt.update({'output': str(output), 'output_bytes': output.stat().st_size, 'output_sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'after': {'page_break_before': True, 'keep_with_next': True}, 'office_integrity_gate': gate_result, 'source_unchanged': hashlib.sha256(source.read_bytes()).hexdigest() == before_hash})
        return _PdReturn(receipt)
    _pd_result = _pd_phase_13()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_14()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def cmd_docx_layout_patch(args: argparse.Namespace) -> None:
    try:
        result = inspect_or_stage_docx_layout(action=args.action, input_path=Path(args.input), target_text=args.target_text, output_path=Path(args.out) if args.out else None, source_sha256=args.source_sha256, overwrite=bool(args.overwrite))
    except (ValueError, RuntimeError, FileNotFoundError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))
    emit(result)

def _replace_one_visible_paragraph_text(paragraph: Any, old_text: str, new_text: str) -> None:
    runs = list(paragraph.runs)
    full_text = ''.join((run.text for run in runs))
    start = full_text.find(old_text)
    if start < 0 or full_text.find(old_text, start + 1) >= 0:
        fail('company announcement replacement must occur exactly once in one paragraph')
    end = start + len(old_text)
    cursor = 0
    inserted = False
    for run in runs:
        original = run.text
        run_start = cursor
        run_end = cursor + len(original)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue
        prefix = original[:max(0, start - run_start)] if run_start <= start < run_end else ''
        suffix = original[max(0, end - run_start):] if run_start < end <= run_end else ''
        if not inserted:
            run.text = prefix + new_text + suffix
            inserted = True
        else:
            run.text = suffix
    if not inserted:
        fail('company announcement replacement could not be mapped to Word runs')

def build_company_announcement_from_spec(spec: dict[str, Any], out: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'python-docx unavailable: {exc!r}')
    replacements = spec.get('replacements')
    if not isinstance(replacements, dict) or not replacements:
        fail('company_announcement template mode requires a non-empty replacements object')
    if len(replacements) > COMPANY_ANNOUNCEMENT_MAX_REPLACEMENTS:
        fail(f'replacements exceeds safe limit of {COMPANY_ANNOUNCEMENT_MAX_REPLACEMENTS}')
    normalized: list[tuple[str, str]] = []
    for old_text, new_text in replacements.items():
        if not isinstance(old_text, str) or not old_text:
            fail('replacement keys must be non-empty strings')
        if not isinstance(new_text, str):
            fail('replacement values must be strings')
        if len(old_text) > COMPANY_ANNOUNCEMENT_MAX_REPLACEMENT_CHARS or len(new_text) > COMPANY_ANNOUNCEMENT_MAX_REPLACEMENT_CHARS:
            fail(f'replacement text exceeds {COMPANY_ANNOUNCEMENT_MAX_REPLACEMENT_CHARS} characters')
        normalized.append((old_text, new_text))
    template = COMPANY_ANNOUNCEMENT_TEMPLATE()
    if not template.is_file():
        fail(f'fixed company announcement template not found: {template}')
    source_bytes = template.read_bytes()
    source_sha256 = hashlib.sha256(source_bytes).hexdigest()
    doc = Document(str(template))
    paragraphs = _iter_docx_story_paragraphs(doc)
    for old_text, new_text in normalized:
        matches = [paragraph for paragraph in paragraphs if old_text in ''.join((run.text for run in paragraph.runs))]
        if len(matches) != 1:
            fail(f'company announcement replacement must match exactly once: {old_text[:120]!r}; matches={len(matches)}')
        _replace_one_visible_paragraph_text(matches[0], old_text, new_text)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f'.{out.name}.', suffix='.docx', dir=str(out.parent), delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        doc.save(str(temp_path))
        gate_script = Path(resolve_configured_path('~/.hermes/skills/pd-one/pd-one-admin/scripts/office_no_corruption_gate.py'))
        if gate_script.exists():
            proc = subprocess.run([sys.executable, str(gate_script), str(temp_path)], text=True, capture_output=True, env=safe_env(), shell=False, timeout=180)
            if proc.returncode != 0:
                fail(f'office no-corruption gate failed: {proc.stderr.strip()[:1200] or proc.stdout.strip()[:1200]}')
        if hashlib.sha256(template.read_bytes()).hexdigest() != source_sha256:
            fail('fixed company announcement template changed during staged build')
        temp_path.replace(out)
    finally:
        temp_path.unlink(missing_ok=True)
    return {'output': str(out), 'sha256': hashlib.sha256(out.read_bytes()).hexdigest(), 'template': 'company_announcement', 'template_path': str(template), 'source_sha256': source_sha256, 'replacement_count': len(normalized), 'paragraphs_scanned': len(paragraphs), 'safety': {'fixed_template': True, 'source_read_only': True, 'local_artifact_write_only': True, 'source_system_writes': False, 'external_delivery': False, 'office_integrity_gate_required': True}}

def build_docx_from_spec(spec: dict[str, Any], out: Path) -> dict[str, Any]:
    doc = _PdUnset
    p = _PdUnset
    sections = _PdUnset

    def _pd_phase_15():
        nonlocal doc, p, sections
        template_mode = spec.get('template')
        if template_mode == 'company_announcement':
            return _PdReturn(build_company_announcement_from_spec(spec, out))
        if template_mode is not None:
            fail(f'unsupported DOCX template mode: {template_mode!r}')
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt
        except Exception as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            fail(f'python-docx unavailable: {exc!r}')
        doc = Document()
        title = spec.get('title')
        if title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_require_str(title, 'title'))
            r.bold = True
            r.font.size = Pt(16)
        subtitle = spec.get('subtitle')
        if subtitle:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_require_str(subtitle, 'subtitle'))
            r.italic = True
        sections = spec.get('sections') or []
        sections = _require_list(sections, 'sections')
        if len(sections) > 200:
            fail('sections exceeds safe limit')

    def _pd_phase_16():
        nonlocal p
        for idx, section in enumerate(sections, start=1):
            if not isinstance(section, dict):
                fail(f'sections[{idx}] must be an object')
            kind = _require_str(section.get('type'), f'sections[{idx}].type').lower()
            if kind == 'heading':
                level = int(section.get('level', 1))
                if level < 1 or level > 3:
                    fail(f'sections[{idx}].level must be between 1 and 3')
                doc.add_heading(_require_str(section.get('text'), f'sections[{idx}].text'), level=level)
            elif kind == 'paragraph':
                p = doc.add_paragraph()
                p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}.get(str(section.get('align', 'left')).lower(), WD_ALIGN_PARAGRAPH.LEFT)
                text = _require_str(section.get('text'), f'sections[{idx}].text')
                run = p.add_run(text)
                if section.get('bold'):
                    run.bold = True
                if section.get('italic'):
                    run.italic = True
            elif kind == 'bullet_list':
                items = _require_list(section.get('items'), f'sections[{idx}].items')
                if len(items) > 100:
                    fail(f'sections[{idx}].items exceeds safe limit')
                for item in items:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(_require_str(item, f'sections[{idx}].items[]'))
            elif kind == 'page_break':
                doc.add_page_break()
            elif kind == 'table':
                rows = _require_list(section.get('rows'), f'sections[{idx}].rows')
                if not rows:
                    fail(f'sections[{idx}].rows cannot be empty')
                normalized_rows: list[list[Any]] = []
                max_cols = 0
                for row_idx, row in enumerate(rows, start=1):
                    row_list = _require_list(row, f'sections[{idx}].rows[{row_idx}]')
                    if len(row_list) > 20:
                        fail(f'sections[{idx}].rows[{row_idx}] exceeds safe column limit')
                    normalized_rows.append(row_list)
                    max_cols = max(max_cols, len(row_list))
                if len(normalized_rows) > 100:
                    fail(f'sections[{idx}].rows exceeds safe row limit')
                table = doc.add_table(rows=len(normalized_rows), cols=max_cols)
                table.style = str(section.get('style') or 'Table Grid')
                header_rows = int(section.get('header_rows', 0))
                if header_rows < 0 or header_rows > len(normalized_rows):
                    fail(f'sections[{idx}].header_rows is out of range')
                _apply_table(table, normalized_rows, header_rows=header_rows)
            else:
                fail(f'sections[{idx}].type unsupported: {kind!r}')
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        gate_script = Path(resolve_configured_path('~/.hermes/skills/pd-one/pd-one-admin/scripts/office_no_corruption_gate.py'))
        if gate_script.exists():
            proc = subprocess.run([sys.executable, str(gate_script), str(out)], text=True, capture_output=True, env=safe_env(), shell=False, timeout=180)
            if proc.returncode != 0:
                fail(f'office no-corruption gate failed: {proc.stderr.strip()[:1200] or proc.stdout.strip()[:1200]}')
        digest = hashlib.sha256(out.read_bytes()).hexdigest()
        return _PdReturn({'output': str(out), 'sha256': digest, 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)})
    _pd_result = _pd_phase_15()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_16()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def _office_visual_qa_file_receipts(out_dir: Path) -> list[dict[str, Any]]:
    receipts: list[dict[str, Any]] = []
    for item in sorted(out_dir.iterdir(), key=lambda value: value.name):
        if not item.is_file():
            continue
        try:
            size_bytes = item.stat().st_size
        except OSError:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            size_bytes = None
        receipts.append({'name': item.name, 'size_bytes': size_bytes})
        if len(receipts) >= OFFICE_VISUAL_QA_MAX_FILES:
            break
    return receipts

def run_office_visual_qa(path: Path, qa_dir: Path) -> dict[str, Any]:
    qa_script = OFFICE_VISUAL_QA_SCRIPT()
    if not qa_script.exists():
        return {'status': 'skipped', 'reason': 'visual-office-qa script unavailable', 'input': str(path), 'out_dir': str(qa_dir), 'safety': {'read_only_source': True, 'source_system_writes': False, 'shell_used': False}}
    qa_dir.mkdir(parents=True, exist_ok=True)
    proc = subprocess.run([sys.executable, str(qa_script), str(path), '--out-dir', str(qa_dir)], text=True, capture_output=True, env=safe_env(), shell=False, timeout=300)
    summary = qa_dir / 'qa-summary.md'
    summary_text = summary.read_text(encoding='utf-8', errors='replace') if summary.exists() else ''
    files = _office_visual_qa_file_receipts(qa_dir)
    result: dict[str, Any] = {'status': 'ok' if proc.returncode == 0 else 'error', 'capability': 'pd_one_office_visual_qa', 'input': str(path), 'out_dir': str(qa_dir), 'summary': str(summary) if summary.exists() else None, 'returncode': proc.returncode, 'page_count': len(list(qa_dir.glob('page-*.png'))), 'files': files, 'safety': {'read_only_source': True, 'source_system_writes': False, 'local_artifact_write': True, 'fixed_helper': True, 'shell_used': False, 'manual_visual_inspection_required': True}}
    if proc.returncode != 0:
        if 'no LibreOffice/soffice renderer found' in summary_text or 'no usable Docker fallback image' in summary_text:
            result.update({'status': 'skipped', 'reason': 'renderer unavailable'})
            return result
        if 'PARTIAL:' in summary_text:
            result.update({'status': 'partial', 'reason': 'renderer produced incomplete visual artifacts'})
            return result
        diagnostic = proc.stderr.strip() or proc.stdout.strip()
        fail(f'visual office QA failed: {diagnostic[:1200] or 'no diagnostic output'}')
    return result

def run_visual_qa(path: Path) -> dict[str, Any]:
    qa_dir = path.parent / f'.{path.stem}-visual-qa'
    return run_office_visual_qa(path, qa_dir)

def _office_docprops_local_name(tag: str) -> str:
    return tag.rsplit('}', 1)[-1]

def _office_docprops_safe_member(name: str) -> str:
    normalized = str(name or '').replace('\\', '/')
    parts = Path(normalized).parts
    if not normalized or normalized.startswith('/') or '..' in parts:
        fail(f'refusing unsafe Office member path: {normalized}')
    return normalized

def _office_docprops_parse_xml(payload: bytes, allowed_fields: tuple[str, ...]) -> dict[str, Any]:
    text = payload.decode('utf-8', errors='replace')
    lowered = text.lower()
    if '<!doctype' in lowered or '<!entity' in lowered:
        fail('Office docProps XML DOCTYPE/entity declarations are not allowed')
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'invalid Office docProps XML: {exc}')
    values: dict[str, Any] = {}
    for element in root.iter():
        local = _office_docprops_local_name(element.tag)
        if local not in allowed_fields or local in values:
            continue
        raw = ''.join(element.itertext()).strip()
        if not raw:
            continue
        clipped = raw[:OFFICE_DOCPROPS_MAX_FIELD_CHARS]
        if local in OFFICE_DOCPROPS_INT_FIELDS:
            try:
                values[local] = int(clipped)
                continue
            except ValueError:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
        values[local] = clipped
    return values

def cmd_office_docprops_inspect(args: argparse.Namespace) -> None:
    """Inspect allowlisted Office Open XML core/app properties without dumping XML."""
    source = require_allowed_doc_path(Path(args.input))
    if source.is_symlink() or any((parent.is_symlink() for parent in source.parents if str(parent).startswith(str(ROOT())))):
        fail('symlinked Office inputs are not allowed')
    suffix = source.suffix.lower()
    if suffix not in OFFICE_DOCPROPS_INPUT_SUFFIXES:
        fail(f'office-docprops-inspect only supports {sorted(OFFICE_DOCPROPS_INPUT_SUFFIXES)} files')
    if not source.is_file():
        fail(f'input must be a regular Office Open XML file: {source}')
    size = source.stat().st_size
    if size <= 0 or size > OFFICE_DOCPROPS_MAX_SOURCE_BYTES:
        fail(f'Office package size must be within 1-{OFFICE_DOCPROPS_MAX_SOURCE_BYTES} bytes: {size}')
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if not zipfile.is_zipfile(source):
        fail('input is not a valid Office Open XML zip package')
    try:
        archive = zipfile.ZipFile(source)
    except zipfile.BadZipFile as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'invalid Office Open XML zip: {exc}')
    with archive:
        namelist = [_office_docprops_safe_member(name) for name in archive.namelist()]
        if len(namelist) > OFFICE_DOCPROPS_MAX_NAMELIST:
            fail(f'Office package exceeds {OFFICE_DOCPROPS_MAX_NAMELIST} member limit')
        infos = {_office_docprops_safe_member(info.filename): info for info in archive.infolist()}
        members: list[dict[str, Any]] = []
        parsed = {'core': {}, 'app': {}}
        omitted_vectors: list[str] = []
        for member in OFFICE_DOCPROPS_MEMBERS:
            info = infos.get(member)
            present = info is not None
            member_size = int(info.file_size) if present else None
            members.append({'name': member, 'present': present, 'size_bytes': member_size})
            if not present:
                continue
            if member_size is None or member_size <= 0 or member_size > OFFICE_DOCPROPS_MAX_MEMBER_BYTES:
                fail(f'{member} size must be within 1-{OFFICE_DOCPROPS_MAX_MEMBER_BYTES} bytes')
            payload = archive.read(member)
            if member.endswith('core.xml'):
                parsed['core'] = _office_docprops_parse_xml(payload, OFFICE_DOCPROPS_CORE_FIELDS)
            else:
                parsed['app'] = _office_docprops_parse_xml(payload, OFFICE_DOCPROPS_APP_FIELDS)
            lowered = payload.decode('utf-8', errors='replace').lower()
            if 'titlesofparts' in lowered:
                omitted_vectors.append('TitlesOfParts')
            if 'headingpairs' in lowered:
                omitted_vectors.append('HeadingPairs')
    emit({'status': 'ok', 'capability': 'pd_one_office_docprops_inspect', 'input': str(source), 'suffix': suffix, 'size_bytes': size, 'sha256': digest, 'members': members, 'core': parsed['core'], 'app': parsed['app'], 'omitted_app_vectors': sorted(set(omitted_vectors)), 'safety': {'admin_only': True, 'read_only': True, 'approved_document_roots_only': True, 'allowlisted_members_only': True, 'raw_xml_omitted': True, 'titles_of_parts_omitted': True, 'source_overwrite_forbidden': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'shell_used': False}})

def cmd_office_visual_qa(args: argparse.Namespace) -> None:
    source = require_allowed_doc_path(Path(args.input))
    if source.suffix.lower() not in OFFICE_VISUAL_QA_INPUT_SUFFIXES:
        fail(f'office-visual-qa only supports {sorted(OFFICE_VISUAL_QA_INPUT_SUFFIXES)} files')
    if source.stat().st_size > OFFICE_VISUAL_QA_MAX_BYTES:
        fail(f'office-visual-qa input exceeds {OFFICE_VISUAL_QA_MAX_BYTES} bytes')
    qa_dir = Path(args.out_dir).expanduser().resolve(strict=False)
    if not any((is_within(qa_dir, root) for root in DOC_ALLOWED_ROOTS())):
        fail(f'output path is outside approved PD One document roots: {qa_dir}')
    emit(run_office_visual_qa(source, qa_dir))

def convert_html_to_docx(input_path: Path, out: Path) -> dict[str, Any]:
    source = require_allowed_doc_path(input_path)
    if source.suffix.lower() not in OFFICE_HTML_TO_DOCX_INPUT_SUFFIXES:
        fail('office-html-to-docx input must be .html or .htm')
    if source.stat().st_size > OFFICE_HTML_TO_DOCX_MAX_BYTES:
        fail(f'office-html-to-docx input exceeds {OFFICE_HTML_TO_DOCX_MAX_BYTES} bytes')
    out = out.resolve(strict=False)
    if not any((is_within(out, root) for root in DOC_ALLOWED_ROOTS())):
        fail(f'output path is outside approved roots: {out}')
    if out.suffix.lower() != '.docx':
        fail('office-html-to-docx output must end with .docx')
    renderer = find_bin('libreoffice') or find_bin('soffice')
    if not renderer:
        fail('LibreOffice/soffice renderer not found')
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()
    with tempfile.TemporaryDirectory(prefix='office-html-to-docx-', dir=str(out.parent)) as tmp_dir:
        workdir = Path(tmp_dir)
        source_copy = workdir / 'source.html'
        shutil.copy2(source, source_copy)
        profile = workdir / 'lo-profile'
        profile.mkdir()
        base_cmd = [renderer, f'-env:UserInstallation=file://{profile.resolve()}', '--headless', '--nologo', '--nofirststartwizard']
        first = run_capture(base_cmd + ['--convert-to', 'odt:writer8', '--outdir', str(workdir), str(source_copy)], timeout=120)
        odt = workdir / 'source.odt'
        if first.returncode != 0 or not odt.is_file() or odt.stat().st_size <= 0:
            fail(f'HTML to ODT conversion failed: {tail_text(first.stderr or first.stdout)}')
        second = run_capture(base_cmd + ['--convert-to', 'docx:Office Open XML Text', '--outdir', str(workdir), str(odt)], timeout=120)
        docx = workdir / 'source.docx'
        if second.returncode != 0 or not docx.is_file() or docx.stat().st_size <= 0:
            fail(f'ODT to DOCX conversion failed: {tail_text(second.stderr or second.stdout)}')
        shutil.copy2(docx, out)
    gate_script = Path(resolve_configured_path('~/.hermes/skills/pd-one/pd-one-admin/scripts/office_no_corruption_gate.py'))
    if gate_script.exists():
        gate = run_capture([sys.executable, str(gate_script), str(out)], timeout=180)
        if gate.returncode != 0:
            fail(f'office no-corruption gate failed: {tail_text(gate.stderr or gate.stdout)}')
    preview_path = out.with_suffix('.pdf')
    pdf_preview = convert_office_to_pdf(out, preview_path)
    return {'status': 'ok', 'capability': 'pd_one_office_html_to_docx', 'input': str(source), 'output': str(out), 'output_bytes': out.stat().st_size, 'output_sha256': hashlib.sha256(out.read_bytes()).hexdigest(), 'conversion': 'html->odt->docx', 'pdf_preview': pdf_preview, 'mattermost_delivery_files': [str(out), str(preview_path)], 'safety': {'shell_used': False, 'source_read_only': True, 'local_artifact_write_only': True, 'no_external_delivery': True}}

def cmd_office_html_to_docx(args: argparse.Namespace) -> None:
    emit(convert_html_to_docx(Path(args.input), Path(args.out)))

def cmd_doc_build(args: argparse.Namespace) -> None:
    spec_path = require_allowed_doc_path(Path(args.spec))
    out = Path(args.out).resolve(strict=False)
    if not any((is_within(out, root) for root in DOC_ALLOWED_ROOTS())):
        fail(f'output path is outside approved roots: {out}')
    if out.suffix.lower() != '.docx':
        fail('doc-build output must end with .docx')
    try:
        spec = json.loads(spec_path.read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'failed to load JSON spec: {exc!r}')
    if not isinstance(spec, dict):
        fail('spec must be a JSON object')
    build = build_docx_from_spec(spec, out)
    result = {'status': 'ok', 'capability': 'pd_one_office_build_artifact', 'spec': str(spec_path), 'output': str(out), 'build': build}
    preview_path = out.with_suffix('.pdf')
    result['pdf_preview'] = convert_office_to_pdf(out, preview_path)
    result['mattermost_delivery_files'] = [str(out), str(preview_path)]
    if args.visual_qa:
        result['visual_qa'] = run_visual_qa(out)
    emit(result)

def _load_ecrn_docx_patch_module() -> Any:
    if not ECRN_DOCX_PATCH_SCRIPT().is_file():
        fail(f'ECRN DOCX patch wrapper not found: {ECRN_DOCX_PATCH_SCRIPT()}')
    spec = importlib.util.spec_from_file_location('pdone_ecrn_docx_ops', ECRN_DOCX_PATCH_SCRIPT())
    if spec is None or spec.loader is None:
        fail(f'could not load ECRN DOCX patch wrapper: {ECRN_DOCX_PATCH_SCRIPT()}')
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module

def cmd_ecrn_docx_unique_text_patch(args: argparse.Namespace) -> None:
    """Inspect or stage one exact visible-text replacement in an ECRN DOCX."""
    try:
        result = _load_ecrn_docx_patch_module().inspect_or_stage(action=args.action, input_path=args.input, old_text=args.old_text, new_text=args.new_text, output_path=args.out, source_sha256=args.source_sha256, overwrite=bool(args.overwrite))
    except (ValueError, RuntimeError, FileNotFoundError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))
    emit(result)
    if result.get('status') != 'ok':
        raise SystemExit(1)
