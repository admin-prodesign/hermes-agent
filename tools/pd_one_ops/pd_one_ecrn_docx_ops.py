#!/usr/bin/env python3
"""Bounded DOCX exact-text inspection and staged revision helpers.

This module intentionally supports only a narrow local-artifact workflow:
inspect an approved DOCX for one exact visible text match, or stage a separate
DOCX with that match replaced. It never edits the source document, sends files,
or claims that an electronic approval/signature remains valid after revision.
"""
from __future__ import annotations

class _PdReturn:
    __slots__ = ('value',)

    def __init__(self, value):
        self.value = value

class _PdHandled:
    pass

class _PdUnset:

    def __bool__(self):
        raise UnboundLocalError('pd one local used before assignment')

    def __repr__(self):
        raise UnboundLocalError('pd one local used before assignment')
_PdUnset = _PdUnset()
import logging
from tools.pd_one_ops.site_config import site_text, resolve_configured_path, external_path, profile_root, text_excludes_configured_hosts, ExternalModule, owner_display_name, repo_cli, package_file, kanban_boards_root
import difflib
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterator

def ROOT():
    return Path(resolve_configured_path('~/.hermes/profiles/pdone'))

def INPUT_ROOTS():
    return (ROOT() / 'cache/documents', ROOT() / 'artifacts/ecrn', ROOT() / 'tmp')

def OUTPUT_ROOTS():
    return (ROOT() / 'artifacts/ecrn', ROOT() / 'tmp')
MAX_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 2000
MAX_REPLACEMENTS = 8

def GATE_SCRIPT():
    return Path(resolve_configured_path('~/.hermes/skills/pd-one/pd-one-admin/scripts/office_no_corruption_gate.py'))

def _is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve(strict=False).relative_to(root.resolve(strict=False))
        return True
    except ValueError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return False

def _require_path(raw: str | Path, roots: tuple[Path, ...], label: str) -> Path:
    path = Path(raw).expanduser().resolve(strict=False)
    if not any((_is_within(path, root) for root in roots)):
        raise ValueError(f'{label} is outside approved ECRN document roots: {path}')
    if not path.is_file():
        raise FileNotFoundError(f'{label} not found: {path}')
    if path.stat().st_size > MAX_BYTES:
        raise ValueError(f'{label} exceeds {MAX_BYTES} bytes: {path}')
    return path

def _require_text(value: str, label: str) -> str:
    if not isinstance(value, str) or not value:
        raise ValueError(f'{label} must be a non-empty string')
    if len(value) > MAX_TEXT_CHARS:
        raise ValueError(f'{label} exceeds {MAX_TEXT_CHARS} characters')
    return value

def _require_sha256(value: str | None) -> str | None:
    if value is None or not str(value).strip():
        return None
    digest = str(value).strip().lower()
    if len(digest) != 64 or any((ch not in '0123456789abcdef' for ch in digest)):
        raise ValueError('source_sha256 must be a 64-character hexadecimal SHA-256 digest')
    return digest

def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()

def _load_document(path: Path) -> Any:
    if path.suffix.lower() != '.docx':
        raise ValueError(f'ECRN DOCX patch only supports .docx input: {path}')
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f'python-docx unavailable: {exc}') from exc
    return Document(str(path))

def _iter_paragraphs(doc: Any) -> Iterator[tuple[Any, dict[str, Any]]]:
    seen: list[Any] = []

    def already_seen(paragraph: Any) -> bool:
        element = paragraph._p
        if any((element is previous for previous in seen)):
            return True
        seen.append(element)
        return False
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        if already_seen(paragraph):
            continue
        yield (paragraph, {'scope': 'body', 'paragraph_index': paragraph_index})
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    if already_seen(paragraph):
                        continue
                    yield (paragraph, {'scope': 'table_cell', 'table_index': table_index, 'row_index': row_index, 'cell_index': cell_index, 'paragraph_index': paragraph_index})

def _find_run_span(paragraph: Any, old_text: str) -> dict[str, int] | None:
    run_texts = [run.text or '' for run in paragraph.runs]
    full_text = ''.join(run_texts)
    if full_text.count(old_text) != 1:
        return None
    start = full_text.index(old_text)
    end = start + len(old_text)
    cursor = 0
    start_index = end_index = None
    start_offset = end_offset = 0
    for index, text in enumerate(run_texts):
        next_cursor = cursor + len(text)
        if start_index is None and cursor <= start < next_cursor:
            start_index = index
            start_offset = start - cursor
        if cursor < end <= next_cursor:
            end_index = index
            end_offset = end - cursor
            break
        cursor = next_cursor
    if start_index is None or end_index is None:
        return None
    return {'start_index': start_index, 'end_index': end_index, 'start_offset': start_offset, 'end_offset': end_offset, 'global_start': start, 'global_end': end}

def _find_matches(doc: Any, old_text: str) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    for paragraph, location in _iter_paragraphs(doc):
        count = paragraph.text.count(old_text)
        if not count:
            continue
        run_span = _find_run_span(paragraph, old_text) if count == 1 else None
        run_indexes = []
        if run_span is not None:
            run_indexes = list(range(run_span['start_index'], run_span['end_index'] + 1))
        matches.append({'location': location, 'occurrences': count, 'run_indexes': run_indexes, 'run_span': run_span, 'patchable_without_rebuilding_runs': run_span is not None and count == 1})
    return matches

def _replace_global_span(paragraph: Any, start: int, end: int, replacement: str) -> None:
    run_texts = [run.text or '' for run in paragraph.runs]
    if not run_texts:
        raise RuntimeError('target paragraph has no editable runs')

    def locate_start(position: int) -> tuple[int, int]:
        cursor = 0
        for index, text in enumerate(run_texts):
            next_cursor = cursor + len(text)
            if position < next_cursor or (position == cursor and bool(text)):
                return (index, position - cursor)
            cursor = next_cursor
        if position == cursor:
            return (len(run_texts) - 1, len(run_texts[-1]))
        raise RuntimeError('replacement start is outside paragraph text')

    def locate_end(position: int) -> tuple[int, int]:
        cursor = 0
        for index, text in enumerate(run_texts):
            next_cursor = cursor + len(text)
            if cursor < position <= next_cursor:
                return (index, position - cursor)
            cursor = next_cursor
        raise RuntimeError('replacement end is outside paragraph text')
    if start < 0 or end < start:
        raise RuntimeError('invalid replacement span')
    start_index, start_offset = locate_start(start)
    if start == end:
        paragraph.runs[start_index].text = run_texts[start_index][:start_offset] + replacement + run_texts[start_index][start_offset:]
        return
    end_index, end_offset = locate_end(end)
    if end_index < start_index:
        raise RuntimeError('replacement span run order is invalid')
    if start_index == end_index:
        paragraph.runs[start_index].text = run_texts[start_index][:start_offset] + replacement + run_texts[start_index][end_offset:]
        return
    paragraph.runs[start_index].text = run_texts[start_index][:start_offset] + replacement
    for index in range(start_index + 1, end_index):
        paragraph.runs[index].text = ''
    paragraph.runs[end_index].text = run_texts[end_index][end_offset:]

def _apply_minimal_run_patch(paragraph: Any, old_text: str, new_text: str, run_span: dict[str, int]) -> None:
    global_start = int(run_span['global_start'])
    matcher = difflib.SequenceMatcher(a=old_text, b=new_text, autojunk=False)
    for tag, old_start, old_end, new_start, new_end in reversed(matcher.get_opcodes()):
        if tag == 'equal':
            continue
        _replace_global_span(paragraph, global_start + old_start, global_start + old_end, new_text[new_start:new_end])

def _base_receipt(*, action: str, source: Path, source_sha256: str, old_text: str, new_text: str, matches: list[dict[str, Any]]) -> dict[str, Any]:
    occurrence_count = sum((int(match['occurrences']) for match in matches))
    patchable = len(matches) == 1 and occurrence_count == 1 and bool(matches[0].get('patchable_without_rebuilding_runs'))
    receipt: dict[str, Any] = {'status': 'ok' if patchable else 'blocked', 'capability': 'pd_one_ecrn_docx_unique_text_patch', 'action': action, 'input': str(source), 'source_sha256': source_sha256, 'requested_old_text_sha256': hashlib.sha256(old_text.encode('utf-8')).hexdigest(), 'requested_new_text_sha256': hashlib.sha256(new_text.encode('utf-8')).hexdigest(), 'match_count': len(matches), 'occurrence_count': occurrence_count, 'matches': matches, 'safety': {'admin_only': True, 'read_only': action == 'inspect', 'source_never_overwritten': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'arbitrary_shell_forbidden': True, 'unique_visible_match_required': True, 'run_formatting_rebuild_forbidden': True, 'source_approval_or_signature_preservation_not_claimed': True, 'reapproval_required_before_delivery': True}}
    if not matches:
        receipt.update({'blocker': 'unique_match_not_found', 'next_step': 'Re-read the current source through the wrapper and confirm the exact visible old text before staging a revision.'})
    elif occurrence_count != 1:
        receipt.update({'blocker': 'unique_match_required', 'next_step': 'Do not stage automatically; narrow the request or resolve the duplicate visible matches manually.'})
    elif not patchable:
        receipt.update({'blocker': 'match_spans_multiple_runs', 'next_step': "Use an approved Office editing lane that preserves the document's run formatting; this wrapper will not rebuild a rich paragraph."})
    return receipt

def inspect_or_stage(*, action: str, input_path: str | Path, old_text: str, new_text: str, output_path: str | Path | None=None, source_sha256: str | None=None, overwrite: bool=False) -> dict[str, Any]:
    actual_source_sha256 = _PdUnset
    output = _PdUnset
    receipt = _PdUnset
    source = _PdUnset
    temp_path = _PdUnset

    def _pd_phase_1():
        nonlocal action, actual_source_sha256, new_text, old_text, output, receipt, source, temp_path
        'Inspect or stage a unique visible-text replacement in a separate DOCX.'
        action = str(action or 'inspect').strip().lower()
        if action not in {'inspect', 'stage'}:
            raise ValueError('action must be one of: inspect, stage')
        old_text = _require_text(old_text, 'old_text')
        new_text = _require_text(new_text, 'new_text')
        if old_text == new_text:
            raise ValueError('old_text and new_text must differ')
        source = _require_path(input_path, INPUT_ROOTS(), 'input')
        if source.suffix.lower() != '.docx':
            raise ValueError(f'input must end with .docx: {source}')
        actual_source_sha256 = _sha256(source)
        expected_sha256 = _require_sha256(source_sha256)
        if expected_sha256 and expected_sha256 != actual_source_sha256:
            return _PdReturn({'status': 'blocked', 'capability': 'pd_one_ecrn_docx_unique_text_patch', 'action': action, 'input': str(source), 'source_sha256': actual_source_sha256, 'expected_source_sha256': expected_sha256, 'blocker': 'source_sha256_mismatch', 'next_step': 'Re-read the current source and provide its current SHA-256 before staging.', 'safety': {'admin_only': True, 'read_only': True, 'source_never_overwritten': True, 'external_delivery': False}})
        doc = _load_document(source)
        matches = _find_matches(doc, old_text)
        receipt = _base_receipt(action=action, source=source, source_sha256=actual_source_sha256, old_text=old_text, new_text=new_text, matches=matches)
        if action == 'inspect' or receipt['status'] != 'ok':
            return _PdReturn(receipt)
        if output_path is None:
            raise ValueError('output_path is required for action=stage')
        output = Path(output_path).expanduser().resolve(strict=False)
        if not any((_is_within(output, root) for root in OUTPUT_ROOTS())):
            raise ValueError(f'output is outside approved ECRN artifact roots: {output}')
        if output.suffix.lower() != '.docx':
            raise ValueError('output must end with .docx')
        if output == source:
            raise ValueError('output must be different from input; source overwrite is forbidden')
        if output.exists() and (not overwrite):
            return _PdReturn({**receipt, 'status': 'blocked', 'blocker': 'output_exists', 'output': str(output), 'next_step': 'Choose a new staged output path or explicitly use overwrite in the admin-only staged-artifact lane.'})
        output.parent.mkdir(parents=True, exist_ok=True)
        temp_path = None

    def _pd_phase_2():
        nonlocal temp_path
        try:
            fd, temp_name = tempfile.mkstemp(prefix=f'.{output.stem}-', suffix='.docx', dir=str(output.parent))
            os.close(fd)
            temp_path = Path(temp_name)
            shutil.copy2(source, temp_path)
            staged_doc = _load_document(temp_path)
            staged_matches = _find_matches(staged_doc, old_text)
            if len(staged_matches) != 1 or sum((int(match['occurrences']) for match in staged_matches)) != 1:
                raise RuntimeError('source changed between inspection and staging; unique match no longer holds')
            match = staged_matches[0]
            run_span = match.get('run_span') or {}
            try:
                int(run_span['global_start'])
                int(run_span['global_end'])
            except (KeyError, TypeError, ValueError) as exc:
                raise RuntimeError('staged match did not resolve to a safe run span') from exc
            target_paragraph = None
            for paragraph, location in _iter_paragraphs(staged_doc):
                if location == match['location']:
                    target_paragraph = paragraph
                    break
            if target_paragraph is None:
                raise RuntimeError('could not resolve the inspected paragraph in the staged copy')
            _apply_minimal_run_patch(target_paragraph, old_text, new_text, run_span)
            staged_doc.save(str(temp_path))
            gate_result = _run_no_corruption_gate(temp_path)
            if gate_result['status'] != 'ok':
                raise RuntimeError(f'Office no-corruption gate failed: {gate_result.get('detail', 'unknown error')}')
            os.replace(temp_path, output)
            temp_path = None
            readback_doc = _load_document(output)
            readback_matches_old = _find_matches(readback_doc, old_text)
            readback_matches_new = _find_matches(readback_doc, new_text)
            output_sha256 = _sha256(output)
            if sum((int(match['occurrences']) for match in readback_matches_old)) != 0:
                raise RuntimeError('readback still contains old text')
            if sum((int(match['occurrences']) for match in readback_matches_new)) != 1:
                raise RuntimeError('readback does not contain exactly one replacement text')
            receipt.update({'status': 'ok', 'output': str(output), 'output_sha256': output_sha256, 'changed': True, 'replacement': {'old_text_sha256': hashlib.sha256(old_text.encode('utf-8')).hexdigest(), 'new_text_sha256': hashlib.sha256(new_text.encode('utf-8')).hexdigest()}, 'readback': {'old_occurrence_count': 0, 'new_occurrence_count': 1, 'source_sha256_unchanged': _sha256(source) == actual_source_sha256}, 'office_no_corruption_gate': gate_result, 'approval_note': 'This is a staged revised artifact only; any electronic approval/signature is not carried forward as valid. Re-approval is required before any delivery or posting.'})
            return _PdReturn(receipt)
        finally:
            if temp_path is not None:
                try:
                    temp_path.unlink(missing_ok=True)
                except Exception:
                    logging.getLogger('pd_one_ops').exception('pd one command failed')
    _pd_result = _pd_phase_1()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_2()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def inspect_or_stage_many(*, action: str, input_path: str | Path, replacements: list[list[str]] | list[tuple[str, str]], output_path: str | Path | None=None, source_sha256: str | None=None, overwrite: bool=False) -> dict[str, Any]:
    actual_source_sha256 = _PdUnset
    match = _PdUnset
    new_text = _PdUnset
    old_text = _PdUnset
    output = _PdUnset
    pairs = _PdUnset
    receipt = _PdUnset
    source = _PdUnset
    temp_path = _PdUnset

    def _pd_phase_3():
        nonlocal action, actual_source_sha256, new_text, old_text, output, pairs, receipt, source, temp_path
        'Inspect or atomically stage 1-8 exact visible-text replacements.\n\n    This generic company-document route deliberately stages a separate local\n    artifact. It never overwrites the source and never performs delivery.\n    '
        action = str(action or 'inspect').strip().lower()
        if action not in {'inspect', 'stage'}:
            raise ValueError('action must be one of: inspect, stage')
        if not isinstance(replacements, list) or not 1 <= len(replacements) <= MAX_REPLACEMENTS:
            raise ValueError(f'replacements must contain 1-{MAX_REPLACEMENTS} old/new pairs')
        pairs = []
        for index, pair in enumerate(replacements):
            if not isinstance(pair, (list, tuple)) or len(pair) != 2:
                raise ValueError(f'replacement {index} must be an old/new pair')
            old_text = _require_text(pair[0], f'replacements[{index}].old_text')
            new_text = _require_text(pair[1], f'replacements[{index}].new_text')
            if old_text == new_text:
                raise ValueError(f'replacement {index} old_text and new_text must differ')
            pairs.append((old_text, new_text))
        old_texts = [pair[0] for pair in pairs]
        new_texts = [pair[1] for pair in pairs]
        if len(set(old_texts)) != len(old_texts) or len(set(new_texts)) != len(new_texts):
            raise ValueError('replacement old_text and new_text values must each be unique')
        for index, (old_text, new_text) in enumerate(pairs):
            for other_index, (other_old, other_new) in enumerate(pairs):
                if index == other_index:
                    continue
                if old_text in other_old or old_text in other_new or new_text in other_old:
                    raise ValueError('overlapping or chained replacement values are forbidden')
        source = _require_path(input_path, INPUT_ROOTS(), 'input')
        actual_source_sha256 = _sha256(source)
        expected_sha256 = _require_sha256(source_sha256)
        if expected_sha256 and expected_sha256 != actual_source_sha256:
            return _PdReturn({'status': 'blocked', 'capability': 'pd_one_docx_exact_text_patch', 'action': action, 'input': str(source), 'source_sha256': actual_source_sha256, 'expected_source_sha256': expected_sha256, 'blocker': 'source_sha256_mismatch', 'safety': {'admin_only': True, 'read_only': True, 'source_never_overwritten': True, 'external_delivery': False}})
        doc = _load_document(source)
        checks: list[dict[str, Any]] = []
        all_patchable = True
        for old_text, new_text in pairs:
            matches = _find_matches(doc, old_text)
            occurrence_count = sum((int(match['occurrences']) for match in matches))
            patchable = len(matches) == 1 and occurrence_count == 1 and bool(matches[0].get('patchable_without_rebuilding_runs'))
            all_patchable = all_patchable and patchable
            checks.append({'old_text_sha256': hashlib.sha256(old_text.encode('utf-8')).hexdigest(), 'new_text_sha256': hashlib.sha256(new_text.encode('utf-8')).hexdigest(), 'match_count': len(matches), 'occurrence_count': occurrence_count, 'patchable': patchable, 'locations': [match['location'] for match in matches[:2]]})
        receipt = {'status': 'ok' if all_patchable else 'blocked', 'capability': 'pd_one_docx_exact_text_patch', 'action': action, 'input': str(source), 'source_sha256': actual_source_sha256, 'replacement_count': len(pairs), 'checks': checks, 'safety': {'admin_only': True, 'read_only': action == 'inspect', 'source_never_overwritten': True, 'separate_staged_output_required': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'arbitrary_shell_forbidden': True, 'unique_visible_match_required_per_replacement': True, 'office_integrity_gate_required_for_stage': True}}
        if not all_patchable:
            receipt.update({'blocker': 'all_replacements_must_have_one_patchable_visible_match', 'next_step': 'Correct the exact visible replacement pairs and inspect again; do not run arbitrary OpenXML mutation.'})
            return _PdReturn(receipt)
        if action == 'inspect':
            return _PdReturn(receipt)
        if output_path is None:
            raise ValueError('output_path is required for action=stage')
        output = Path(output_path).expanduser().resolve(strict=False)
        if not any((_is_within(output, root) for root in OUTPUT_ROOTS())):
            raise ValueError(f'output is outside approved staged artifact roots: {output}')
        if output.suffix.lower() != '.docx' or output == source:
            raise ValueError('output must be a separate .docx under an approved staged artifact root')
        if output.exists() and (not overwrite):
            return _PdReturn({**receipt, 'status': 'blocked', 'blocker': 'output_exists', 'output': str(output)})
        output.parent.mkdir(parents=True, exist_ok=True)
        temp_path = None

    def _pd_phase_4():
        nonlocal match, new_text, old_text, temp_path
        try:
            fd, temp_name = tempfile.mkstemp(prefix=f'.{output.stem}-', suffix='.docx', dir=str(output.parent))
            os.close(fd)
            temp_path = Path(temp_name)
            shutil.copy2(source, temp_path)
            staged_doc = _load_document(temp_path)
            for old_text, new_text in pairs:
                current_matches = _find_matches(staged_doc, old_text)
                if len(current_matches) != 1 or sum((int(match['occurrences']) for match in current_matches)) != 1:
                    raise RuntimeError('replacement set changed during staging; unique match no longer holds')
                match = current_matches[0]
                run_span = match.get('run_span')
                if not run_span:
                    raise RuntimeError('staged match did not resolve to a safe run span')
                target_paragraph = next((paragraph for paragraph, location in _iter_paragraphs(staged_doc) if location == match['location']), None)
                if target_paragraph is None:
                    raise RuntimeError('could not resolve an inspected paragraph in the staged copy')
                _apply_minimal_run_patch(target_paragraph, old_text, new_text, run_span)
            staged_doc.save(str(temp_path))
            gate_result = _run_no_corruption_gate(temp_path)
            if gate_result['status'] != 'ok':
                raise RuntimeError(f'Office no-corruption gate failed: {gate_result.get('detail', 'unknown error')}')
            os.replace(temp_path, output)
            temp_path = None
            readback_doc = _load_document(output)
            readback = []
            for old_text, new_text in pairs:
                old_count = sum((int(match['occurrences']) for match in _find_matches(readback_doc, old_text)))
                new_count = sum((int(match['occurrences']) for match in _find_matches(readback_doc, new_text)))
                if old_count != 0 or new_count != 1:
                    raise RuntimeError('staged artifact replacement readback failed')
                readback.append({'old_occurrence_count': old_count, 'new_occurrence_count': new_count})
            receipt.update({'output': str(output), 'output_sha256': _sha256(output), 'changed': True, 'readback': readback, 'source_sha256_unchanged': _sha256(source) == actual_source_sha256, 'office_no_corruption_gate': gate_result})
            return _PdReturn(receipt)
        finally:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
    _pd_result = _pd_phase_3()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_4()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def _run_no_corruption_gate(path: Path) -> dict[str, Any]:
    if not GATE_SCRIPT().is_file():
        return {'status': 'blocked', 'detail': f'gate script not found: {GATE_SCRIPT()}'}
    proc = subprocess.run([sys.executable, str(GATE_SCRIPT()), str(path)], text=True, capture_output=True, shell=False, timeout=180)
    detail = (proc.stderr.strip() or proc.stdout.strip())[-1600:]
    return {'status': 'ok' if proc.returncode == 0 else 'error', 'returncode': proc.returncode, 'detail': detail}
__all__ = ['inspect_or_stage', 'inspect_or_stage_many']
