from __future__ import annotations
__all__ = ['cmd_approval_ticket_cancel', '_bounded_json_object', 'cmd_admin_approval_kanban_status', 'load_hr_policy_lookup_module', 'cmd_hr_policy_lookup', 'post_file_entries', 'match_post_terms', 'cmd_mattermost_member_search', 'channel_policy_suggestion', 'cmd_mattermost_channel_setup', 'load_json', 'is_within', 'require_allowed_doc_path', 'ledger_ticket', 'require_approved_ticket', 'stable_dump', 'ensure_user', 'validate_mm_user_id', 'cmd_policy_change_dry_run', 'cmd_policy_change_apply', 'extract_docx', 'inspect_docx_paragraph', 'inspect_docx_body_xml', 'inspect_docx_table_structure', 'inspect_docx_table_matches', 'extract_legacy_doc', 'PPTX_EXTRACT_MAX_ARCHIVE_BYTES', 'PPTX_EXTRACT_MAX_SLIDE_XML_BYTES', 'PPTX_EXTRACT_MAX_SLIDES', 'PPTX_EXTRACT_MAX_TEXTS_PER_SLIDE', 'PPTX_EXTRACT_MAX_TEXT', '_pptx_reject_unsafe_zip_name', '_pptx_package_stats', '_extract_pptx_stdlib', 'extract_pptx', '_normalize_pptx_ooxml_member', '_pptx_ooxml_drawingml_texts', 'inspect_pptx_ooxml_text', '_pptx_trim_text', '_pptx_shape_type_name', '_pptx_is_group', '_pptx_collect_text_frame', 'inspect_pptx', 'extract_pdf', '_legacy_xls_column_name', '_legacy_xls_cell_address', '_legacy_xls_visibility_label', '_legacy_xls_cell_type_label', '_legacy_xls_trim_text', '_legacy_xls_jsonable', '_legacy_xls_sheet_attrs', '_legacy_xls_decode_cell_value', '_legacy_xls_encryption_probe', 'inspect_legacy_xls', 'extracted_plain_text', 'ANNOUNCEMENT_SIGNALS', 'REWARD_DISCIPLINE_SIGNALS', 'classify_company_announcement', 'extract_by_suffix', 'cmd_doc_extract', 'cmd_artifact_base64_encode', 'cmd_local_artifact_copy', 'cmd_video_frame_sample', 'cmd_svg_artifact_inspect', '_parse_3mf_slicer_settings', '_safe_3mf_zip_info', '_validate_artifact_identity_glob', 'cmd_artifact_identity_inventory']
import logging
from tools.pd_one_ops.site_config import site_text, resolve_configured_path, external_path, profile_root, text_excludes_configured_hosts, ExternalModule, owner_display_name, repo_cli, package_file, kanban_boards_root
import argparse
import base64
import calendar
import contextlib
import collections
import csv
import difflib
import fnmatch
import importlib.util
import hashlib
import html
import math
import io
import json
import os
import re
import shlex
import shutil
import socket
import subprocess
import types
import sys
import tempfile
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any
from urllib.parse import urlparse
from zoneinfo import ZoneInfo
from tools.pd_one_ops.pd_one_safe_ops_part0 import *
from tools.pd_one_ops.pd_one_safe_ops_part1 import *

def cmd_approval_ticket_cancel(args: argparse.Namespace) -> None:
    """Cancel one queued/in_progress/approved PDAPP ticket without executing it.

    Owner/admin abandonment of a named ticket is a ledger terminal write, not a
    decision-post approve/reject. This wrapper is the typed path: exact ticket
    id, required reason, dry-run default, live confirmation phrase, and
    status readback. It does not enqueue, claim, complete, retry, expand
    scope, or start the executor/Kanban plane.
    """
    ticket_id = (args.ticket_id or '').strip()
    if not re.match('^PDAPP-\\d{8}-[A-F0-9]{8}$', ticket_id):
        fail('--ticket-id must look like PDAPP-YYYYMMDD-XXXXXXXX')
    reason = (args.reason or '').strip()
    if len(reason) < 8:
        fail('--reason must be at least 8 characters')
    dry_run = bool(getattr(args, 'dry_run', True))
    tickets_mod, _ = load_admin_approval_modules()
    ledger = tickets_mod.ApprovalLedger(str(APPROVAL_LEDGER()))
    try:
        current = ledger.get(ticket_id)
    except KeyError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        emit({'status': 'not_found', 'capability': 'pd_one_approval_ticket_cancel', 'ticket_id': ticket_id, 'safety': {'read_only': dry_run, 'no_privileged_action_executed': True}})
        return
    ticket_status = current.get('status')
    common = {'capability': 'pd_one_approval_ticket_cancel', 'ticket_id': ticket_id, 'ticket_status': ticket_status, 'reason': reason, 'safety': {'dry_run_default': True, 'live_write_confirmation_required': ADMIN_APPROVAL_CANCEL_CONFIRMATION, 'no_executor_start': True, 'no_scope_expansion': True, 'no_retry': True, 'decision_post_not_used': True}}
    if ticket_status == 'cancelled':
        emit({**common, 'status': 'already_cancelled', 'ticket_status_after': 'cancelled', 'would_apply': False})
        return
    if ticket_status not in CANCELLABLE_TICKET_STATUSES:
        emit({**common, 'status': 'blocked', 'blocked_reason': 'status_not_cancellable', 'would_apply': False, 'ticket_status_after': ticket_status})
        return
    if dry_run:
        emit({**common, 'status': 'dry_run', 'would_apply': True, 'ticket_status_after': 'cancelled'})
        return
    if args.confirm_live_write != ADMIN_APPROVAL_CANCEL_CONFIRMATION:
        fail(f'live ticket cancel requires --confirm-live-write {ADMIN_APPROVAL_CANCEL_CONFIRMATION!r}')
    ticket_after = ledger.cancel(ticket_id, reason=reason)
    emit({**common, 'status': 'ok', 'would_apply': False, 'ticket_status_after': ticket_after.get('status'), 'cancel_reason': ticket_after.get('cancel_reason'), 'updated_at_utc': ticket_after.get('updated_at_utc'), 'safety': {**common['safety'], 'readback_verification_required': True, 'no_privileged_action_executed': False}})

def _bounded_json_object(path: Path) -> tuple[bool, dict[str, Any] | None, str | None]:
    if not path.is_file():
        return (False, None, None)
    try:
        size = path.stat().st_size
    except OSError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return (True, None, f'stat_failed:{exc}')
    if size > ADMIN_APPROVAL_KANBAN_JSON_MAX_BYTES:
        return (True, None, 'too_large')
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return (True, None, f'invalid_json:{type(exc).__name__}')
    if not isinstance(data, dict):
        return (True, None, 'not_object')
    return (True, data, None)

def cmd_admin_approval_kanban_status(args: argparse.Namespace) -> None:
    """Read bounded PDAPP Kanban mapping/result/workspace existence without listing files."""
    ticket_id = (args.ticket_id or '').strip()
    if not PDAPP_TICKET_ID_RE.match(ticket_id):
        fail('--ticket-id must look like PDAPP-YYYYMMDD-XXXXXXXX')
    mapping_file = ADMIN_APPROVAL_KANBAN_STATE() / 'mappings' / f'{ticket_id}.json'
    result_file = ADMIN_APPROVAL_KANBAN_STATE() / 'results' / f'{ticket_id}.json'
    mapping_exists, mapping, mapping_error = _bounded_json_object(mapping_file)
    result_exists, result, result_error = _bounded_json_object(result_file)
    mapping_receipt: dict[str, Any] = {'exists': mapping_exists, 'valid': False, 'path_kind': 'admin-approval-kanban/mappings/<ticket-id>.json'}
    workspace_exists = None
    if mapping_error:
        mapping_receipt['error'] = mapping_error
    elif mapping:
        schema_ok = mapping.get('schema') == ADMIN_APPROVAL_KANBAN_MAPPING_SCHEMA
        ticket_ok = mapping.get('ticket_id') == ticket_id
        mapping_receipt['valid'] = bool(schema_ok and ticket_ok)
        if mapping_receipt['valid']:
            task_id = str(mapping.get('task_id') or '')
            board = str(mapping.get('board') or ADMIN_APPROVAL_KANBAN_BOARD)
            mapping_receipt.update({'schema': mapping.get('schema'), 'task_id': task_id or None, 'board': board, 'ledger_claimed': bool(mapping.get('ledger_claimed')), 'unblocked': bool(mapping.get('unblocked')), 'created_at_utc': mapping.get('created_at_utc'), 'execution_attempt_id': mapping.get('execution_attempt_id')})
            if board == ADMIN_APPROVAL_KANBAN_BOARD and KANBAN_TASK_ID_RE.match(task_id):
                workspace = KANBAN_BOARDS_ROOT() / board / 'workspaces' / task_id
                workspace_exists = workspace.is_dir()
    result_receipt: dict[str, Any] = {'exists': result_exists, 'valid': False, 'path_kind': 'admin-approval-kanban/results/<ticket-id>.json'}
    if result_error:
        result_receipt['error'] = result_error
    elif result:
        schema_ok = result.get('schema') == ADMIN_APPROVAL_EXECUTION_RESULT_SCHEMA
        ticket_ok = result.get('ticket_id') == ticket_id
        result_receipt['valid'] = bool(schema_ok and ticket_ok)
        if result_receipt['valid']:
            remaining = result.get('remaining_gates')
            result_receipt.update({'schema': result.get('schema'), 'outcome': result.get('outcome'), 'execution_attempt_id': result.get('execution_attempt_id'), 'remaining_gates_count': len(remaining) if isinstance(remaining, list) else None})
    emit({'status': 'ok', 'capability': 'pd_one_admin_approval_kanban_status', 'ticket_id': ticket_id, 'mapping': mapping_receipt, 'result': result_receipt, 'workspace_exists': workspace_exists, 'safety': {'read_only': True, 'admin_only': True, 'bounded_ticket_lookup': True, 'workspace_listing_forbidden': True, 'mcp_source_read_forbidden': True, 'systemd_runtime_probe_forbidden': True, 'summaries_not_returned': True, 'no_privileged_action_executed': True, 'no_shell_glue': True}})

def load_hr_policy_lookup_module():
    module_path = ROOT() / 'tools/pd_one_hr_policy_lookup.py'
    spec = importlib.util.spec_from_file_location('pd_one_hr_policy_lookup_runtime', module_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f'cannot load HR policy lookup module: {module_path}')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def cmd_hr_policy_lookup(args: argparse.Namespace) -> None:
    """Search only the approved current HR policy documents for an exact active requester."""
    requester_user_id = validate_mattermost_id(args.requester_user_id, 'requester_user_id')
    query = (args.query or '').strip()
    if not query:
        fail('--query is required')
    document_keys = split_csv_lines(args.document_keys)
    try:
        result = load_hr_policy_lookup_module().lookup_hr_policy(requester_user_id, query, document_keys=document_keys or None, max_results=args.max_results)
    except (PermissionError, ValueError, RuntimeError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(str(exc))
    emit(result)

def post_file_entries(post: dict[str, Any]) -> list[dict[str, str]]:
    """Return attachment id/name pairs from Mattermost post metadata (no extra API)."""
    entries: list[dict[str, str]] = []
    seen: set[str] = set()
    metadata = post.get('metadata') if isinstance(post.get('metadata'), dict) else {}
    files = metadata.get('files') if isinstance(metadata, dict) else None
    if not isinstance(files, list):
        files = []
    for item in files:
        if not isinstance(item, dict):
            continue
        fid = str(item.get('id') or '').strip()
        name = str(item.get('name') or '').strip()
        key = fid or name
        if not key or key in seen:
            continue
        seen.add(key)
        entries.append({'id': fid, 'name': name})
    return entries

def match_post_terms(post: dict[str, Any], terms: list[str]) -> tuple[str, str]:
    """Return (matched_term, matched_on) using message body and attachment filenames."""
    message_l = (post.get('message') or '').lower()
    names_l = ' '.join((entry.get('name') or '' for entry in post_file_entries(post))).lower()
    msg_hit = ''
    file_hit = ''
    for term in terms[:12]:
        needle = (term or '').lower().strip()
        if not needle:
            continue
        if not msg_hit and needle in message_l:
            msg_hit = term
        if not file_hit and needle in names_l:
            file_hit = term
        if msg_hit and file_hit:
            break
    if msg_hit and file_hit:
        return (msg_hit, 'both')
    if msg_hit:
        return (msg_hit, 'message')
    if file_hit:
        return (file_hit, 'filename')
    return ('', '')

def cmd_mattermost_member_search(args: argparse.Namespace) -> None:
    base = _PdUnset
    ch = _PdUnset
    cid = _PdUnset
    headers = _PdUnset
    hits = _PdUnset
    matched_on = _PdUnset
    matched_term = _PdUnset
    max_fallback_posts = _PdUnset
    max_results = _PdUnset
    me = _PdUnset
    post = _PdUnset
    query = _PdUnset
    requested_channels = _PdUnset
    search_channels = _PdUnset
    search_errors = _PdUnset
    sender_id = _PdUnset
    since_ms = _PdUnset
    terms = _PdUnset

    def _pd_phase_5():
        nonlocal base, cid, headers, hits, matched_on, matched_term, max_fallback_posts, max_results, me, post, query, requested_channels, search_channels, search_errors, sender_id, since_ms, terms
        'Search only channels where the requester is a verified Mattermost member.'
        sender_id = validate_mattermost_id(args.sender_id, 'sender_id')
        query = (args.query or '').strip()
        if not query:
            fail('--query is required')
        max_results = max(1, min(int(args.max_results), 50))
        max_channels = max(1, min(int(args.max_channels), 200))
        max_fallback_posts = max(0, min(int(args.max_fallback_posts), 3000))
        since_ms = parse_since_ms(args.since)
        requested_channels = [validate_mattermost_id(c.strip(), 'channel_id') for c in args.channel_id or [] if c.strip()]
        users_policy = load_json(POLICY_USERS()) if POLICY_USERS().exists() else {}
        sender_profile = (users_policy.get('users') or {}).get(sender_id) or {}
        if sender_profile.get('status') != 'active':
            emit({'status': 'blocked', 'reason': 'sender_not_active_profiled_user', 'senderId': sender_id})
            return _PdReturn(None)
        base, headers, me = mattermost_context()
        membership = mattermost_user_membership(base, headers, sender_id)
        if requested_channels:
            missing = [cid for cid in requested_channels if cid not in membership]
            if missing:
                emit({'status': 'blocked', 'reason': 'requester_not_member_of_requested_channel', 'senderId': sender_id, 'blockedChannelIds': missing, 'safety': {'botAdminVisibilityIgnored': True, 'authorizationBasis': 'Mattermost requester active channel membership'}})
                return _PdReturn(None)
            search_channels = {cid: membership[cid] for cid in requested_channels}
        else:
            search_channels = dict(list(membership.items())[:max_channels])
        terms = [t.strip() for t in re.split('[,，\\n]+', query) if t.strip()]
        if not terms:
            terms = [query]
        team_ids = sorted({str(ch.get('team_id') or '') for ch in search_channels.values() if ch.get('team_id')})
        hits = {}
        search_errors = []
        for team_id in team_ids:
            for term in terms[:8]:
                try:
                    result = mm_request('POST', base, {**headers, 'Content-Type': 'application/json'}, f'/api/v4/teams/{team_id}/posts/search', json={'terms': term, 'is_or_search': False, 'include_deleted_channels': False}, timeout=35)
                except SystemExit as _exc:
                    logging.getLogger('pd_one_ops').exception('pd one command failed')
                    search_errors.append(f'team={team_id} term={term!r} failed')
                    continue
                for post in (result.get('posts') or {}).values():
                    cid = post.get('channel_id')
                    if cid not in search_channels or post.get('delete_at'):
                        continue
                    if since_ms and int(post.get('create_at') or 0) < since_ms:
                        continue
                    matched_term, matched_on = match_post_terms(post, [term])
                    hits[post.get('id') or ''] = {'post': post, 'matchedTerm': matched_term or term, 'matchedOn': matched_on or 'message'}
                    if len(hits) >= max_results:
                        break
                if len(hits) >= max_results:
                    break
            if len(hits) >= max_results:
                break

    def _pd_phase_6():
        nonlocal ch, cid, matched_on, matched_term, post
        if requested_channels and max_fallback_posts and (len(hits) < max_results):
            for cid in requested_channels:
                page = 0
                scanned = 0
                while scanned < max_fallback_posts and len(hits) < max_results:
                    payload = mm_request('GET', base, headers, f'/api/v4/channels/{cid}/posts', params={'page': page, 'per_page': 200}, timeout=35)
                    order = payload.get('order') or []
                    if not order:
                        break
                    for pid in order:
                        post = (payload.get('posts') or {}).get(pid) or {}
                        scanned += 1
                        if post.get('delete_at') or (since_ms and int(post.get('create_at') or 0) < since_ms):
                            continue
                        matched_term, matched_on = match_post_terms(post, terms[:12])
                        if matched_term:
                            hits[post.get('id') or ''] = {'post': post, 'matchedTerm': matched_term or query, 'matchedOn': matched_on}
                        if scanned >= max_fallback_posts or len(hits) >= max_results:
                            break
                    if len(order) < 200:
                        break
                    page += 1
        rows = []
        for item in sorted(hits.values(), key=lambda x: int((x.get('post') or {}).get('create_at') or 0), reverse=True)[:max_results]:
            post = item.get('post') or {}
            cid = post.get('channel_id') or ''
            ch = search_channels.get(cid) or {}
            post_id = post.get('id', '')
            original_post_id = post.get('root_id') or post_id
            files = post_file_entries(post)
            rows.append({'postId': post_id, 'originalPostId': original_post_id, 'channelId': cid, 'channelName': ch.get('display_name') or ch.get('name') or cid, 'createAt': post.get('create_at', 0), 'userId': post.get('user_id', ''), 'matchedTerm': item.get('matchedTerm', ''), 'matchedOn': item.get('matchedOn') or 'message', 'files': files, 'link': post_link(base, post_id), 'original_post_url': post_link(base, original_post_id), 'excerpt': compact_excerpt(post.get('message', ''), item.get('matchedTerm', ''))})
        emit({'status': 'ok', 'capability': 'pd_one_mattermost_member_search', 'senderId': sender_id, 'query': query, 'since': args.since or '', 'hitCapReached': len(hits) >= max_results, 'searchedChannelCount': len(search_channels), 'searchedChannels': [{'channelId': cid, 'name': ch.get('display_name') or ch.get('name') or cid} for cid, ch in list(search_channels.items())[:20]], 'sources_searched': [{'source': 'mattermost_member_visible_channels', 'mode': 'team_posts_search_plus_optional_channel_fallback', 'searchedChannelCount': len(search_channels), 'channelSample': [{'channelId': cid, 'name': ch.get('display_name') or ch.get('name') or cid} for cid, ch in list(search_channels.items())[:20]], 'since': args.since or '', 'query': query}], 'resultCount': len(rows), 'results': rows, 'errors': search_errors[:5], 'safety': {'readOnly': True, 'authorizationBasis': 'exact sender id + live Mattermost active channel membership', 'botAdminVisibilityIgnored': True, 'rawTranscriptDumps': 'denied; bounded excerpts and post links only', 'sourceSystemWrites': 'denied', 'externalSends': 'denied'}, 'replyGuidance': 'List Sources checked / 已查來源 from sources_searched in the bilingual user reply.', 'botUserId': me.get('id', '')})
    _pd_result = _pd_phase_5()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_6()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def channel_policy_suggestion(channel: dict[str, Any], members: list[dict[str, Any]], users: dict[str, dict[str, Any]], summary: dict[str, Any]) -> dict[str, Any]:
    policy = load_json(POLICY_USERS()) if POLICY_USERS().exists() else {}
    user_records = policy.get('users') if isinstance(policy, dict) else {}
    if isinstance(user_records, list):
        user_records = {u.get('mattermostUserId') or u.get('id'): u for u in user_records if isinstance(u, dict)}
    member_rows = []
    role_counts: collections.Counter[str] = collections.Counter()
    for m in members:
        uid = m.get('user_id') or ''
        u = users.get(uid, {})
        rec = user_records.get(uid, {}) if isinstance(user_records, dict) else {}
        roles = rec.get('roles') or rec.get('assignedRoles') or []
        for role in roles:
            role_counts[str(role)] += 1
        member_rows.append({'user_id': uid, 'username': u.get('username') or uid, 'policy_found': bool(rec), 'roles': roles, 'delete_at': u.get('delete_at'), 'is_bot': u.get('is_bot')})
    suggested_roles = [r for r, _ in role_counts.most_common(8) if r != 'admin']
    if 'key_employee_operations' not in suggested_roles:
        suggested_roles.insert(0, 'key_employee_operations')
    return {'channel_id': channel.get('id'), 'current_display_name': channel.get('display_name'), 'suggested_roles': suggested_roles[:8], 'member_count': len(member_rows), 'unprofiled_members': [m for m in member_rows if not m.get('policy_found')], 'top_role_counts': role_counts.most_common(12), 'suggested_policy_note': 'Review before editing policies/mattermost-channels.md; this wrapper does not grant permissions or expand channel toolsets automatically.', 'member_sample': member_rows[:30], 'evidence_summary': summary}

def cmd_mattermost_channel_setup(args: argparse.Namespace) -> None:
    channel_id = validate_mattermost_id(args.channel_id, 'channel_id')
    action = args.action
    max_posts = max(0, min(int(args.max_posts), 2000))
    if action == 'apply_metadata' and args.confirm_live_write != MATTERMOST_CHANNEL_SETUP_CONFIRMATION:
        fail(f'apply_metadata requires --confirm-live-write {MATTERMOST_CHANNEL_SETUP_CONFIRMATION!r}')
    base, headers, me = mattermost_context()
    channel = mm_request('GET', base, headers, f'/api/v4/channels/{channel_id}')
    posts = fetch_channel_posts(base, headers, channel_id, max_posts)
    post_user_ids = [str(p.get('user_id') or '') for p in posts]
    members = fetch_channel_members(base, headers, channel_id) if args.include_members else []
    member_user_ids = [str(m.get('user_id') or '') for m in members]
    users = mm_users_by_id(base, headers, post_user_ids + member_user_ids)
    summary = summarize_channel(posts, users)
    target_display = args.target_display_name or propose_channel_display_name(str(channel.get('display_name') or ''))
    target_purpose = args.target_purpose or propose_channel_purpose(summary)
    plan = {'status': 'ok', 'capability': 'pd_one_mattermost_channel_setup', 'action': action, 'safety': {'shell_used': False, 'token_printed': False, 'live_write_requested': action == 'apply_metadata', 'live_confirmation_required': MATTERMOST_CHANNEL_SETUP_CONFIRMATION if action == 'apply_metadata' else None, 'policy_permissions_changed': False}, 'bot': {'id': me.get('id'), 'username': me.get('username'), 'roles': me.get('roles')}, 'channel_before': {k: channel.get(k) for k in ['id', 'name', 'display_name', 'purpose', 'header', 'type', 'team_id', 'total_msg_count']}, 'history_summary': summary, 'metadata_plan': {'endpoint': f'/api/v4/channels/{channel_id}/patch', 'display_name': target_display, 'purpose': target_purpose, 'lengths': {'display_name': len(target_display or ''), 'purpose': len(target_purpose or '')}}, 'policy_suggestion': channel_policy_suggestion(channel, members, users, summary) if args.include_members else None}
    if action in {'inspect', 'dry_run_update'}:
        emit(plan)
        return
    payload = {'display_name': target_display, 'purpose': target_purpose}
    applied = mm_request('PUT', base, {**headers, 'Content-Type': 'application/json'}, f'/api/v4/channels/{channel_id}/patch', json=payload, timeout=35)
    readback = mm_request('GET', base, headers, f'/api/v4/channels/{channel_id}')
    plan['apply_status'] = 'applied'
    plan['apply_response'] = {k: applied.get(k) for k in ['id', 'display_name', 'purpose', 'update_at']}
    plan['readback'] = {'display_name': readback.get('display_name'), 'purpose': readback.get('purpose'), 'display_name_match': readback.get('display_name') == target_display, 'purpose_match': readback.get('purpose') == target_purpose}
    emit(plan)

def load_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except FileNotFoundError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'file not found: {path}')
    except json.JSONDecodeError as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'invalid JSON in {path}: {exc}')

def is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve(strict=False).relative_to(root.resolve(strict=False))
        return True
    except ValueError:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return False

def require_allowed_doc_path(path: Path) -> Path:
    p = path.resolve(strict=False)
    allowed_files = {candidate.resolve(strict=False) for candidate in DOC_ALLOWED_FILES()}
    if p not in allowed_files and (not any((is_within(p, root) for root in DOC_ALLOWED_ROOTS()))):
        fail(f'path is outside approved PD One document roots: {p}')
    if not p.exists():
        fail(f'input file not found: {p}')
    return p

def ledger_ticket(ticket_id: str) -> dict[str, Any]:
    if not re.match('^PDAPP-\\d{8}-[A-F0-9]{8}$', ticket_id or ''):
        fail('ticket_id must look like PDAPP-YYYYMMDD-XXXXXXXX')
    if not APPROVAL_LEDGER().exists():
        fail('approval ledger not found')
    state: dict[str, Any] | None = None
    for line in APPROVAL_LEDGER().read_text(encoding='utf-8', errors='replace').splitlines():
        if not line.strip():
            continue
        try:
            ev = json.loads(line)
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            continue
        if ev.get('ticket_id') != ticket_id:
            continue
        status = ev.get('status') or ev.get('to_status') or ev.get('decision')
        if state is None:
            state = {'ticket_id': ticket_id, 'events': []}
        state['events'].append(ev.get('action') or ev.get('event') or ev.get('type') or 'event')
        if status:
            state['status'] = status
        if ev.get('andy_decision'):
            state['andy_decision'] = ev.get('andy_decision')
        if ev.get('approved_scope'):
            state['approved_scope'] = ev.get('approved_scope')
    if not state:
        fail(f'ticket not found: {ticket_id}')
    return state

def require_approved_ticket(ticket_id: str, purpose: str) -> dict[str, Any]:
    ticket = ledger_ticket(ticket_id)
    status = str(ticket.get('status') or '').lower()
    if status not in {'approved', 'queued', 'in_progress'}:
        fail(f'ticket {ticket_id} is not in an approved executable state: {status}')
    scope_text = json.dumps(ticket, ensure_ascii=False).lower()
    if purpose.lower() not in scope_text and 'permission' not in scope_text and ('policy' not in scope_text):
        fail(f'ticket {ticket_id} approved scope does not obviously cover {purpose}; refusing apply')
    return ticket

def stable_dump(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + '\n'

def ensure_user(data: dict[str, Any], user_id: str) -> dict[str, Any]:
    users = data.get('users')
    if not isinstance(users, dict) or user_id not in users:
        fail(f'Mattermost user id not found in users policy: {user_id}')
    if not isinstance(users[user_id], dict):
        fail(f'user entry is not an object: {user_id}')
    return users[user_id]

def validate_mm_user_id(user_id: str) -> str:
    u = (user_id or '').strip()
    if not re.match('^[a-z0-9]{20,32}$', u):
        fail('Mattermost user id must be exact lowercase id, not display name')
    return u

def cmd_policy_change_dry_run(args: argparse.Namespace) -> None:
    user_id = validate_mm_user_id(args.user_id)
    data = load_json(POLICY_USERS())
    before = stable_dump(data).splitlines(keepends=True)
    user = ensure_user(data, user_id)
    changes: list[str] = []
    if args.add_role:
        roles = user.setdefault('roles', [])
        if not isinstance(roles, list):
            fail('user roles field is not a list')
        if args.add_role not in roles:
            roles.append(args.add_role)
            changes.append(f'add_role:{args.add_role}')
    if args.add_dm_scope:
        dm = user.setdefault('dm', {})
        if not isinstance(dm, dict):
            fail('user dm field is not an object')
        scopes = dm.setdefault('allowedScopes', [])
        if not isinstance(scopes, list):
            fail('user dm.allowedScopes field is not a list')
        if args.add_dm_scope not in scopes:
            scopes.append(args.add_dm_scope)
            changes.append(f'add_dm_scope:{args.add_dm_scope}')
    if args.set_active is not None:
        active = args.set_active == 'true'
        if user.get('active') != active:
            user['active'] = active
            changes.append(f'set_active:{active}')
    after_text = stable_dump(data)
    load_json(POLICY_ROLES())
    diff = ''.join(difflib.unified_diff(before, after_text.splitlines(keepends=True), fromfile=str(POLICY_USERS()), tofile=str(POLICY_USERS()) + '.proposed'))
    receipt = {'status': 'ok', 'capability': 'pd_one_policy_change_dry_run', 'mode': 'dry-run', 'user_id': user_id, 'changes': changes, 'changed': bool(changes), 'diff': diff, 'proposed_data': data if args.include_proposed_data else None, 'target': str(POLICY_USERS()), 'apply_note': 'Use policy-change-apply only with an approved ticket and exact same requested changes.'}
    if not args.include_proposed_data:
        receipt.pop('proposed_data', None)
    if args.receipt:
        p = Path(args.receipt)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(stable_dump(receipt), encoding='utf-8')
    emit(receipt)

def cmd_policy_change_apply(args: argparse.Namespace) -> None:
    if not args.ticket_id:
        fail('--ticket-id is required')
    _ticket = require_approved_ticket(args.ticket_id, 'policy')
    dry_args = argparse.Namespace(user_id=args.user_id, add_role=args.add_role, add_dm_scope=args.add_dm_scope, set_active=args.set_active, receipt=None, include_proposed_data=True)
    user_id = validate_mm_user_id(dry_args.user_id)
    data = load_json(POLICY_USERS())
    before_text = stable_dump(data)
    user = ensure_user(data, user_id)
    changes: list[str] = []
    if dry_args.add_role:
        roles = user.setdefault('roles', [])
        if dry_args.add_role not in roles:
            roles.append(dry_args.add_role)
            changes.append(f'add_role:{dry_args.add_role}')
    if dry_args.add_dm_scope:
        dm = user.setdefault('dm', {})
        scopes = dm.setdefault('allowedScopes', [])
        if dry_args.add_dm_scope not in scopes:
            scopes.append(dry_args.add_dm_scope)
            changes.append(f'add_dm_scope:{dry_args.add_dm_scope}')
    if dry_args.set_active is not None:
        active = dry_args.set_active == 'true'
        if user.get('active') != active:
            user['active'] = active
            changes.append(f'set_active:{active}')
    if not changes:
        emit({'status': 'ok', 'capability': 'pd_one_policy_change_apply_approved', 'mode': 'live', 'changed': False, 'user_id': user_id, 'ticket_id': args.ticket_id})
        return
    backup = POLICY_USERS().with_suffix(POLICY_USERS().suffix + f'.bak-{args.ticket_id}-{int(time.time())}')
    shutil.copy2(POLICY_USERS(), backup)
    POLICY_USERS().write_text(stable_dump(data), encoding='utf-8')
    readback = load_json(POLICY_USERS())
    if stable_dump(readback) != stable_dump(data):
        POLICY_USERS().write_text(before_text, encoding='utf-8')
        fail('policy apply readback mismatch; restored previous content')
    emit({'status': 'ok', 'capability': 'pd_one_policy_change_apply_approved', 'mode': 'live', 'changed': True, 'user_id': user_id, 'changes': changes, 'ticket_id': args.ticket_id, 'backup': str(backup), 'readback_ok': True})

def extract_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        ns = {'w': namespace}
        with zipfile.ZipFile(path) as archive:
            try:
                info = archive.getinfo('word/document.xml')
            except KeyError as exc:
                raise ValueError('DOCX is missing word/document.xml') from exc
            if info.file_size > 20 * 1024 * 1024:
                raise ValueError('DOCX word/document.xml exceeds the 20 MiB extraction limit')
            document_xml = archive.read(info)
        if b'<!DOCTYPE' in document_xml.upper() or b'<!ENTITY' in document_xml.upper():
            raise ValueError('DOCX document XML contains a forbidden declaration')
        root = ET.fromstring(document_xml)
        body_wide_paragraphs = root.findall('.//w:p', ns)
        if len(body_wide_paragraphs) > 5000:
            raise ValueError('DOCX exceeds the 5000 paragraph extraction limit')
        indexed_paragraphs = []
        for index, paragraph in enumerate(body_wide_paragraphs):
            text = ''.join((node.text or '' for node in paragraph.findall('.//w:t', ns)))
            if text.strip():
                indexed_paragraphs.append({'index': index, 'text': text[:4000]})
        paragraphs = [item['text'] for item in indexed_paragraphs]
        return {'paragraph_count': len(paragraphs), 'paragraphs': paragraphs[:200], 'body_paragraphs': indexed_paragraphs[:200], 'body_paragraphs_truncated': len(indexed_paragraphs) > 200, 'body_wide_paragraph_count': len(body_wide_paragraphs), 'tables': [], 'extraction_backend': 'stdlib_wordprocessingml'}
    doc = Document(str(path))
    indexed_paragraphs = [{'index': index, 'text': paragraph.text} for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip()]
    paragraphs = [item['text'] for item in indexed_paragraphs]
    tables = []
    for table in doc.tables[:20]:
        tables.append([[cell.text for cell in row.cells] for row in table.rows[:50]])
    return {'paragraph_count': len(paragraphs), 'paragraphs': paragraphs[:200], 'body_paragraphs': indexed_paragraphs[:200], 'body_paragraphs_truncated': len(indexed_paragraphs) > 200, 'tables': tables, 'extraction_backend': 'python_docx'}

def inspect_docx_paragraph(path: Path, *, table_index: int, row_index: int, cell_index: int, paragraph_index: int=0, include_runs: bool=True) -> dict[str, Any]:
    """Read one bounded DOCX table-cell paragraph and its runs.

    This replaces recurring ad hoc ``python-docx`` probes that walk a caller-
    supplied table/row/cell/paragraph path.  It is deliberately targeted: the
    caller must name one paragraph, and the receipt is capped so a large or
    sensitive document cannot be dumped accidentally through a generic helper.
    """
    if path.suffix.lower() != '.docx':
        raise ValueError('docx-inspect only supports .docx files')
    if path.stat().st_size > DOCX_INSPECT_MAX_BYTES:
        raise ValueError(f'DOCX exceeds the {DOCX_INSPECT_MAX_BYTES} byte inspection limit')
    indices = {'table_index': table_index, 'row_index': row_index, 'cell_index': cell_index, 'paragraph_index': paragraph_index}
    for name, value in indices.items():
        if isinstance(value, bool) or not isinstance(value, int) or value < 0 or (value > 500):
            raise ValueError(f'{name} must be an integer from 0 through 500')
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f'python-docx unavailable: {exc}') from exc
    doc = Document(str(path))
    if table_index >= len(doc.tables):
        raise ValueError(f'table_index {table_index} is outside document table count {len(doc.tables)}')
    table = doc.tables[table_index]
    if row_index >= len(table.rows):
        raise ValueError(f'row_index {row_index} is outside table row count {len(table.rows)}')
    row = table.rows[row_index]
    if cell_index >= len(row.cells):
        raise ValueError(f'cell_index {cell_index} is outside row cell count {len(row.cells)}')
    cell = row.cells[cell_index]
    if paragraph_index >= len(cell.paragraphs):
        raise ValueError(f'paragraph_index {paragraph_index} is outside cell paragraph count {len(cell.paragraphs)}')
    paragraph = cell.paragraphs[paragraph_index]

    def bounded_text(value: str) -> str:
        text = value or ''
        return text[:DOCX_INSPECT_MAX_TEXT_CHARS] + ('…' if len(text) > DOCX_INSPECT_MAX_TEXT_CHARS else '')
    runs = None
    runs_truncated = False
    if include_runs:
        runs = [bounded_text(run.text) for run in paragraph.runs[:DOCX_INSPECT_MAX_RUNS]]
        runs_truncated = len(paragraph.runs) > DOCX_INSPECT_MAX_RUNS
    return {'table_count': len(doc.tables), 'table_row_count': len(table.rows), 'row_cell_count': len(row.cells), 'cell_paragraph_count': len(cell.paragraphs), 'target': {'table_index': table_index, 'row_index': row_index, 'cell_index': cell_index, 'paragraph_index': paragraph_index}, 'paragraph': {'text': bounded_text(paragraph.text), 'run_count': len(paragraph.runs), 'runs': runs, 'runs_truncated': runs_truncated}}

def inspect_docx_body_xml(path: Path, paragraph_indices: list[int]) -> dict[str, Any]:
    """Return bounded WordprocessingML for selected body-wide paragraph indices."""
    if path.suffix.lower() != '.docx':
        raise ValueError('docx-body-xml-inspect only supports .docx files')
    if path.stat().st_size > DOCX_INSPECT_MAX_BYTES:
        raise ValueError(f'DOCX exceeds the {DOCX_INSPECT_MAX_BYTES} byte inspection limit')
    if not paragraph_indices or len(paragraph_indices) > 20:
        raise ValueError('docx-body-xml-inspect requires 1-20 paragraph indices')
    if any((isinstance(index, bool) or not isinstance(index, int) or index < 0 or (index > 5000) for index in paragraph_indices)):
        raise ValueError('paragraph indices must be integers from 0 through 5000')
    if len(set(paragraph_indices)) != len(paragraph_indices):
        raise ValueError('paragraph indices must be unique')
    namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    ns = {'w': namespace}
    with zipfile.ZipFile(path) as archive:
        try:
            info = archive.getinfo('word/document.xml')
        except KeyError as exc:
            raise ValueError('DOCX is missing word/document.xml') from exc
        if info.file_size > 20 * 1024 * 1024:
            raise ValueError('DOCX word/document.xml exceeds the 20 MiB inspection limit')
        document_xml = archive.read(info)
    if b'<!DOCTYPE' in document_xml.upper() or b'<!ENTITY' in document_xml.upper():
        raise ValueError('DOCX document XML contains a forbidden declaration')
    root = ET.fromstring(document_xml)
    paragraphs = root.findall('.//w:p', ns)
    max_index = max(paragraph_indices)
    if max_index >= len(paragraphs):
        raise ValueError(f'paragraph_index {max_index} is outside body-wide paragraph count {len(paragraphs)}')
    receipts: list[dict[str, Any]] = []
    for index in paragraph_indices:
        paragraph = paragraphs[index]
        text_nodes = paragraph.findall('.//w:t', ns)
        visible_text = ''.join((node.text or '' for node in text_nodes))
        run_texts = [''.join((node.text or '' for node in run.findall('.//w:t', ns))) for run in paragraph.findall('.//w:r', ns)]
        xml = ET.tostring(paragraph, encoding='unicode')
        receipts.append({'index': index, 'text': visible_text[:4000] + ('…' if len(visible_text) > 4000 else ''), 'run_count': len(run_texts), 'runs': run_texts[:100], 'runs_truncated': len(run_texts) > 100, 'xml': xml[:5000] + ('…' if len(xml) > 5000 else ''), 'xml_truncated': len(xml) > 5000})
    return {'body_wide_paragraph_count': len(paragraphs), 'paragraphs': receipts}

def inspect_docx_table_structure(path: Path) -> dict[str, Any]:
    """Read bounded WordprocessingML table geometry and visible cell text."""
    if path.suffix.lower() != '.docx':
        raise ValueError('docx-table-structure only supports .docx files')
    if path.stat().st_size > DOCX_INSPECT_MAX_BYTES:
        raise ValueError(f'DOCX exceeds the {DOCX_INSPECT_MAX_BYTES} byte inspection limit')
    namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    q = lambda name: f'{{{namespace}}}{name}'
    max_tables, max_rows, max_cells, max_text_chars = (20, 100, 2000, 1000)
    with zipfile.ZipFile(path) as archive:
        try:
            document_xml = archive.read('word/document.xml')
        except KeyError as exc:
            raise ValueError('DOCX is missing word/document.xml') from exc
    if b'<!DOCTYPE' in document_xml.upper() or b'<!ENTITY' in document_xml.upper():
        raise ValueError('DOCX document XML contains a forbidden declaration')
    root = ET.fromstring(document_xml)
    all_tables = root.findall(f'.//{q('tbl')}')
    tables: list[dict[str, Any]] = []
    cells_scanned = 0
    rows_truncated = False
    cells_truncated = False
    for table_index, table in enumerate(all_tables[:max_tables]):
        grid = table.find(q('tblGrid'))
        grid_columns = list(grid.findall(q('gridCol'))) if grid is not None else []
        widths = [column.get(q('w')) or column.get('w:w') for column in grid_columns]
        table_rows = list(table.findall(q('tr')))
        rows: list[dict[str, Any]] = []
        if len(table_rows) > max_rows:
            rows_truncated = True
        for row_index, row in enumerate(table_rows[:max_rows]):
            row_cells = list(row.findall(q('tc')))
            cell_receipts: list[dict[str, Any]] = []
            for cell_index, cell in enumerate(row_cells):
                if cells_scanned >= max_cells:
                    cells_truncated = True
                    break
                text = ''.join((node.text or '' for node in cell.findall(f'.//{q('t')}')))
                truncated = len(text) > max_text_chars
                cell_receipts.append({'cell_index': cell_index, 'text': text[:max_text_chars] + ('…' if truncated else ''), 'text_truncated': truncated})
                cells_scanned += 1
            rows.append({'row_index': row_index, 'cell_count': len(row_cells), 'cells': cell_receipts, 'cells_truncated': len(cell_receipts) < len(row_cells)})
            if cells_scanned >= max_cells:
                break
        tables.append({'table_index': table_index, 'grid_column_count': len(grid_columns), 'grid_widths_twips': widths, 'row_count': len(table_rows), 'rows': rows, 'rows_truncated': len(rows) < len(table_rows)})
        if cells_scanned >= max_cells:
            break
    return {'status': 'ok', 'capability': 'pd_one_docx_table_structure', 'input': str(path), 'table_count': len(all_tables), 'tables_scanned': len(tables), 'tables_truncated': len(tables) < len(all_tables), 'rows_truncated': rows_truncated, 'cells_scanned': cells_scanned, 'cells_truncated': cells_truncated, 'tables': tables, 'safety': {'read_only': True, 'admin_only': True, 'approved_document_root_required': True, 'bounded_receipt': True, 'shell_used': False, 'external_delivery': False}}

def inspect_docx_table_matches(path: Path, terms: list[str]) -> dict[str, Any]:
    """Return bounded matching table-cell structure from one approved DOCX."""
    if path.suffix.lower() != '.docx':
        raise ValueError('docx-table-search only supports .docx files')
    if path.stat().st_size > DOCX_INSPECT_MAX_BYTES:
        raise ValueError(f'DOCX exceeds the {DOCX_INSPECT_MAX_BYTES} byte inspection limit')
    normalized_terms = [str(term).strip() for term in terms if str(term).strip()]
    if not normalized_terms:
        raise ValueError('docx-table-search requires at least one non-empty term')
    if len(normalized_terms) > 8:
        raise ValueError('docx-table-search accepts at most 8 terms')
    if any((len(term) > 128 for term in normalized_terms)):
        raise ValueError('docx-table-search terms must be at most 128 characters')
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f'python-docx unavailable: {exc}') from exc
    doc = Document(str(path))
    terms_folded = [(term, term.casefold()) for term in normalized_terms]
    matches: list[dict[str, Any]] = []
    seen_cells: set[Any] = set()
    tables_scanned = 0
    rows_scanned = 0
    cells_scanned = 0
    duplicate_cells_skipped = 0
    max_tables = 20
    max_rows = 100
    max_cells = 2000
    max_matches = 30
    max_text_chars = min(DOCX_INSPECT_MAX_TEXT_CHARS, 4000)
    for table_index, table in enumerate(doc.tables[:max_tables]):
        tables_scanned += 1
        for row_index, row in enumerate(table.rows[:max_rows]):
            rows_scanned += 1
            for cell_index, cell in enumerate(row.cells):
                cell_key = cell._tc
                if cell_key in seen_cells:
                    duplicate_cells_skipped += 1
                    continue
                seen_cells.add(cell_key)
                if cells_scanned >= max_cells:
                    break
                cells_scanned += 1
                cell_text = cell.text or ''
                matched_terms = [original for original, folded in terms_folded if folded in cell_text.casefold()]
                if not matched_terms:
                    continue
                matches.append({'table': table_index, 'row': row_index, 'cell': cell_index, 'matched_terms': matched_terms, 'text': cell_text[:max_text_chars], 'direct_paragraphs': [(paragraph.text or '')[:max_text_chars] for paragraph in cell.paragraphs[:50]], 'nested_tables': len(cell.tables)})
                if len(matches) >= max_matches:
                    break
            if cells_scanned >= max_cells or len(matches) >= max_matches:
                break
        if cells_scanned >= max_cells or len(matches) >= max_matches:
            break
    return {'status': 'ok', 'capability': 'pd_one_docx_table_search', 'input': str(path), 'terms': normalized_terms, 'matches': matches, 'limits': {'max_tables': max_tables, 'max_rows_per_table': max_rows, 'max_cells': max_cells, 'max_matches': max_matches, 'max_text_chars': max_text_chars}, 'scan': {'tables_scanned': tables_scanned, 'rows_scanned': rows_scanned, 'cells_scanned': cells_scanned, 'duplicate_cells_skipped': duplicate_cells_skipped}, 'safety': {'read_only': True, 'source_written': False, 'external_calls': False}}

def extract_legacy_doc(path: Path) -> dict[str, Any]:
    """Best-effort text extraction for legacy Word .doc files.

    This is intentionally read-only and dependency-light. It is good enough for
    document-type routing/classification and short memo translation follow-ups;
    complex conversion should still use a real Office renderer when available.
    """
    try:
        import olefile
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'error': 'olefile unavailable'}
    try:
        ole = olefile.OleFileIO(str(path))
        data = ole.openstream('WordDocument').read()
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'error': f'legacy .doc extraction failed: {exc}'}
    chunks: list[str] = []
    for enc in ('utf-16le', 'cp950'):
        text = data.decode(enc, errors='ignore')
        for chunk in re.findall('[^\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f]{4,}', text):
            cleaned = re.sub('[\\uf000-\\uffff]', '', chunk).strip()
            if cleaned and any((ch.isalnum() or '一' <= ch <= '鿿' for ch in cleaned)):
                chunks.append(cleaned)
    seen: set[str] = set()
    out: list[str] = []
    for chunk in chunks:
        key = re.sub('\\s+', ' ', chunk)
        if key in seen:
            continue
        seen.add(key)
        out.append(chunk[:8000])
        if len(out) >= 80:
            break
    return {'paragraph_count': len(out), 'paragraphs': out}
PPTX_EXTRACT_MAX_ARCHIVE_BYTES = 80 * 1024 * 1024
PPTX_EXTRACT_MAX_SLIDE_XML_BYTES = 8 * 1024 * 1024
PPTX_EXTRACT_MAX_SLIDES = 80
PPTX_EXTRACT_MAX_TEXTS_PER_SLIDE = 40
PPTX_EXTRACT_MAX_TEXT = 4000

def _pptx_reject_unsafe_zip_name(name: str) -> None:
    candidate = Path(name)
    if candidate.is_absolute() or name.startswith('/') or name.startswith('\\') or ('..' in candidate.parts):
        raise ValueError(f'PPTX contains an unsafe zip member path: {name}')

def _pptx_package_stats(path: Path) -> dict[str, Any]:
    size = path.stat().st_size
    if size > PPTX_EXTRACT_MAX_ARCHIVE_BYTES:
        raise ValueError(f'PPTX exceeds the {PPTX_EXTRACT_MAX_ARCHIVE_BYTES} byte extraction limit')
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        for name in names:
            _pptx_reject_unsafe_zip_name(name)
        slide_names = []
        for name in names:
            match = re.fullmatch('ppt/slides/slide(\\d+)\\.xml', name)
            if match:
                slide_names.append((int(match.group(1)), name))
        slide_names.sort(key=lambda item: item[0])
        uncompressed = sum((info.file_size for info in archive.infolist()))
    return {'slide_xml_count': len(slide_names), 'total_files': len(names), 'size_bytes': size, 'uncompressed_size_bytes': uncompressed, 'slide_xml_names': [name for _, name in slide_names[:PPTX_EXTRACT_MAX_SLIDES]], 'slides_truncated': len(slide_names) > PPTX_EXTRACT_MAX_SLIDES}

def _extract_pptx_stdlib(path: Path, stats: dict[str, Any]) -> dict[str, Any]:
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    slides: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        for index, name in enumerate(stats['slide_xml_names'], 1):
            info = archive.getinfo(name)
            if info.file_size > PPTX_EXTRACT_MAX_SLIDE_XML_BYTES:
                raise ValueError(f'PPTX {name} exceeds the {PPTX_EXTRACT_MAX_SLIDE_XML_BYTES} byte XML limit')
            document_xml = archive.read(info)
            if b'<!DOCTYPE' in document_xml.upper() or b'<!ENTITY' in document_xml.upper():
                raise ValueError('PPTX slide XML contains a forbidden declaration')
            root = ET.fromstring(document_xml)
            texts: list[str] = []
            for node in root.findall('.//a:t', ns):
                text = (node.text or '').strip()
                if not text:
                    continue
                texts.append(text[:PPTX_EXTRACT_MAX_TEXT])
                if len(texts) >= PPTX_EXTRACT_MAX_TEXTS_PER_SLIDE:
                    break
            slides.append({'slide': index, 'texts': texts, 'xml_name': name})
    return {'slide_count': stats['slide_xml_count'], 'slides': slides, 'package': {'slide_xml_count': stats['slide_xml_count'], 'total_files': stats['total_files'], 'size_bytes': stats['size_bytes'], 'uncompressed_size_bytes': stats['uncompressed_size_bytes']}, 'slides_truncated': bool(stats['slides_truncated']), 'extraction_backend': 'stdlib_presentationml'}

def extract_pptx(path: Path) -> dict[str, Any]:
    """Read bounded PPTX package stats and slide text.

    This is the canonical replacement for ad hoc ZipFile/ElementTree
    PresentationML probes. Prefer python-pptx when it can open the file;
    otherwise fall back to stdlib slide XML text extraction.
    """
    stats = _pptx_package_stats(path)
    package = {'slide_xml_count': stats['slide_xml_count'], 'total_files': stats['total_files'], 'size_bytes': stats['size_bytes'], 'uncompressed_size_bytes': stats['uncompressed_size_bytes']}
    try:
        from pptx import Presentation
        presentation = Presentation(str(path))
        slides = []
        for idx, slide in enumerate(presentation.slides, 1):
            if idx > PPTX_EXTRACT_MAX_SLIDES:
                break
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text and shape.text.strip():
                    texts.append(shape.text[:PPTX_EXTRACT_MAX_TEXT])
                if len(texts) >= PPTX_EXTRACT_MAX_TEXTS_PER_SLIDE:
                    break
            slides.append({'slide': idx, 'texts': texts})
        return {'slide_count': len(list(presentation.slides)), 'slides': slides, 'package': package, 'slides_truncated': len(list(presentation.slides)) > PPTX_EXTRACT_MAX_SLIDES, 'extraction_backend': 'python_pptx'}
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return _extract_pptx_stdlib(path, stats)

def _normalize_pptx_ooxml_member(raw: str) -> str:
    member = str(raw or '').strip().replace('\\', '/')
    if not member or member.startswith('/') or '..' in member.split('/'):
        raise ValueError(f'invalid PPTX member path: {raw}')
    if PPTX_OOXML_TEXT_INSPECT_MEMBER_RE.fullmatch(member) is None:
        raise ValueError(f'PPTX member is outside the allowlist: {member}')
    return member

def _pptx_ooxml_drawingml_texts(xml_bytes: bytes) -> list[str]:
    if b'<!DOCTYPE' in xml_bytes.upper() or b'<!ENTITY' in xml_bytes.upper():
        raise ValueError('PPTX member XML contains a forbidden declaration')
    root = ET.fromstring(xml_bytes)
    texts: list[str] = []
    for element in root.iter():
        tag = element.tag or ''
        local = tag.rsplit('}', 1)[-1]
        if local == 't':
            value = element.text or ''
            if value:
                texts.append(value)
    return texts

def inspect_pptx_ooxml_text(path: Path, members: list[str]) -> dict[str, Any]:
    """Read bounded DrawingML a:t text from allowlisted PPTX package members."""
    if path.suffix.lower() != '.pptx':
        raise ValueError('pptx-ooxml-text-inspect only supports .pptx files')
    if path.stat().st_size > PPTX_OOXML_TEXT_INSPECT_MAX_BYTES:
        raise ValueError(f'PPTX exceeds the {PPTX_OOXML_TEXT_INSPECT_MAX_BYTES} byte inspection limit')
    if not members or len(members) > PPTX_OOXML_TEXT_INSPECT_MAX_MEMBERS:
        raise ValueError(f'pptx-ooxml-text-inspect requires 1-{PPTX_OOXML_TEXT_INSPECT_MAX_MEMBERS} members')
    normalized = [_normalize_pptx_ooxml_member(member) for member in members]
    if len(set(normalized)) != len(normalized):
        raise ValueError('PPTX members must be unique')
    receipts: list[dict[str, Any]] = []
    missing = 0
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for name in names:
            _pptx_reject_unsafe_zip_name(name)
        for member in normalized:
            if member not in names:
                missing += 1
                receipts.append({'member': member, 'exists': False, 'text_count': 0, 'texts': [], 'truncated': False})
                continue
            info = archive.getinfo(member)
            if info.file_size > PPTX_OOXML_TEXT_INSPECT_MAX_MEMBER_BYTES:
                raise ValueError(f'PPTX member exceeds the {PPTX_OOXML_TEXT_INSPECT_MAX_MEMBER_BYTES} byte limit: {member}')
            texts = _pptx_ooxml_drawingml_texts(archive.read(info))
            truncated = len(texts) > PPTX_OOXML_TEXT_INSPECT_MAX_TEXTS
            bounded = []
            for text in texts[:PPTX_OOXML_TEXT_INSPECT_MAX_TEXTS]:
                clipped = text[:PPTX_OOXML_TEXT_INSPECT_MAX_TEXT_CHARS]
                if len(text) > PPTX_OOXML_TEXT_INSPECT_MAX_TEXT_CHARS:
                    clipped += '…'
                    truncated = True
                bounded.append(clipped)
            receipts.append({'member': member, 'exists': True, 'text_count': len(texts), 'texts': bounded, 'truncated': truncated})
    status = 'ok' if missing == 0 else 'error' if missing == len(normalized) else 'partial'
    return {'status': status, 'capability': 'pd_one_pptx_ooxml_text_inspect', 'input': str(path), 'type': '.pptx', 'member_count': len(normalized), 'members_missing': missing, 'members': receipts, 'safety': {'read_only': True, 'admin_only': True, 'approved_document_root': True, 'bounded_members': True, 'raw_xml_not_returned': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'external_delivery': False}}

def _pptx_trim_text(value: Any) -> tuple[str, bool]:
    text = '' if value is None else str(value)
    if len(text) > PPTX_INSPECT_MAX_TEXT:
        return (text[:PPTX_INSPECT_MAX_TEXT] + '…', True)
    return (text, False)

def _pptx_shape_type_name(shape: Any) -> str:
    value = getattr(shape, 'shape_type', None)
    name = getattr(value, 'name', None)
    if name:
        return str(name)
    if value is None:
        return 'unknown'
    return str(value).split('.')[-1]

def _pptx_is_group(shape: Any) -> bool:
    return _pptx_shape_type_name(shape).upper() == 'GROUP' and hasattr(shape, 'shapes')

def _pptx_collect_text_frame(shape: Any) -> tuple[list[dict[str, Any]], bool]:
    if not getattr(shape, 'has_text_frame', False):
        return ([], False)
    texts: list[dict[str, Any]] = []
    truncated = False
    for pi, paragraph in enumerate(shape.text_frame.paragraphs):
        if pi >= PPTX_INSPECT_MAX_PARAS:
            truncated = True
            break
        runs: list[dict[str, Any]] = []
        for ri, run in enumerate(paragraph.runs):
            if ri >= PPTX_INSPECT_MAX_RUNS:
                truncated = True
                break
            run_text = getattr(run, 'text', '') or ''
            if not run_text:
                continue
            trimmed, run_truncated = _pptx_trim_text(run_text)
            truncated = truncated or run_truncated
            runs.append({'ri': ri, 'text': trimmed, 'truncated': run_truncated})
        para_text = getattr(paragraph, 'text', '') or ''
        if not para_text and (not runs):
            continue
        trimmed, para_truncated = _pptx_trim_text(para_text)
        truncated = truncated or para_truncated
        texts.append({'pi': pi, 'text': trimmed, 'truncated': para_truncated, 'runs': runs})
    return (texts, truncated)

def _pptx_table_cells(shape: Any) -> tuple[list[dict[str, Any]], bool]:
    table_cells: list[dict[str, Any]] = []
    truncated = False
    try:
        table = shape.table
        for row_i, row in enumerate(table.rows):
            for col_i, cell in enumerate(row.cells):
                cell_text = getattr(cell, 'text', '') or ''
                if not cell_text.strip():
                    continue
                trimmed, cell_truncated = _pptx_trim_text(cell_text)
                truncated = truncated or cell_truncated
                table_cells.append({'row': row_i, 'col': col_i, 'text': trimmed, 'truncated': cell_truncated})
                if len(table_cells) >= PPTX_INSPECT_MAX_PARAS:
                    return table_cells, True
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return [], truncated
    return table_cells, truncated

def inspect_pptx(path: Path, *, include_text: bool=True) -> dict[str, Any]:
    """Read-only recursive PPTX shape/text inspect under approved document roots."""
    path = require_allowed_doc_path(Path(path))
    if path.suffix.lower() != '.pptx':
        fail('pptx-inspect only supports .pptx files')
    size_bytes = path.stat().st_size
    if size_bytes > PPTX_INSPECT_MAX_BYTES:
        fail(f'presentation exceeds {PPTX_INSPECT_MAX_BYTES} bytes: {path}')
    try:
        from pptx import Presentation
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'python-pptx unavailable: {exc}')
    data = path.read_bytes()
    prs = Presentation(str(path))
    slide_count = len(prs.slides)
    width = int(prs.slide_width)
    height = int(prs.slide_height)
    shapes_out: list[dict[str, Any]] = []
    shapes_seen = 0
    shapes_truncated = False
    text_truncated = False

    def walk(container: Any, prefix: str, depth: int) -> None:
        nonlocal shapes_seen, shapes_truncated, text_truncated
        if shapes_truncated or depth > PPTX_INSPECT_MAX_DEPTH:
            if depth > PPTX_INSPECT_MAX_DEPTH:
                shapes_truncated = True
            return
        try:
            children = list(container)
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            return
        for index, shape in enumerate(children):
            if shapes_seen >= PPTX_INSPECT_MAX_SHAPES:
                shapes_truncated = True
                return
            shapes_seen += 1
            path_label = f'{prefix}/sh{index}'
            record: dict[str, Any] = {'path': path_label, 'name': str(getattr(shape, 'name', '') or ''), 'shape_type': _pptx_shape_type_name(shape)}
            if include_text:
                texts, truncated = _pptx_collect_text_frame(shape)
                text_truncated = text_truncated or truncated
                if texts:
                    record['texts'] = texts
            if getattr(shape, 'has_table', False):
                record_cells, cells_truncated = _pptx_table_cells(shape)
                text_truncated = text_truncated or cells_truncated
                if record_cells:
                    record['table_cells'] = record_cells
            shapes_out.append(record)
            if _pptx_is_group(shape):
                walk(shape.shapes, path_label, depth + 1)
    for slide_index, slide in enumerate(prs.slides, 1):
        if slide_index > PPTX_INSPECT_MAX_SLIDES:
            shapes_truncated = True
            break
        walk(slide.shapes, f's{slide_index}', 1)
    return {'status': 'ok', 'capability': 'pd_one_pptx_inspect', 'input': str(path), 'name': path.name, 'type': '.pptx', 'size_bytes': size_bytes, 'sha256': hashlib.sha256(data).hexdigest(), 'slide_count': slide_count, 'slides_inspected': min(slide_count, PPTX_INSPECT_MAX_SLIDES), 'slide_width_emu': width, 'slide_height_emu': height, 'slide_width': width, 'slide_height': height, 'include_text': bool(include_text), 'shape_count': len(shapes_out), 'shapes_seen': shapes_seen, 'shapes_truncated': shapes_truncated, 'text_truncated': text_truncated, 'shapes': shapes_out, 'safety': {'read_only': True, 'admin_only': True, 'approved_document_root': True, 'bounded_target': True, 'bounded_output': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'external_delivery': False}}

def extract_pdf(path: Path) -> dict[str, Any]:
    try:
        import fitz
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'error': 'pymupdf unavailable'}
    doc = fitz.open(str(path))
    pages = []
    for i, page in enumerate(doc):
        if i >= 50:
            break
        pages.append({'page': i + 1, 'text': page.get_text('text')[:8000]})
    return {'page_count': doc.page_count, 'pages': pages}

def _legacy_xls_column_name(index_1_based: int) -> str:
    n = max(1, int(index_1_based))
    letters = ''
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters

def _legacy_xls_cell_address(row_index_0: int, col_index_0: int) -> str:
    return f'{_legacy_xls_column_name(col_index_0 + 1)}{row_index_0 + 1}'

def _legacy_xls_visibility_label(value: Any) -> str:
    labels = {0: 'visible', 1: 'hidden', 2: 'very_hidden'}
    try:
        return labels.get(int(value), f'unknown({value})')
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return f'unknown({value})'

def _legacy_xls_cell_type_label(ctype: Any) -> str:
    labels = {0: 'empty', 1: 'text', 2: 'number', 3: 'date', 4: 'boolean', 5: 'error', 6: 'blank'}
    try:
        return labels.get(int(ctype), f'ctype_{ctype}')
    except Exception:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return f'ctype_{ctype}'

def _legacy_xls_trim_text(value: Any, limit: int=240) -> Any:
    if isinstance(value, str) and len(value) > limit:
        return value[:limit] + '…'
    return value

def _legacy_xls_jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        return {str(key): _legacy_xls_jsonable(val) for key, val in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_legacy_xls_jsonable(item) for item in value]
    return str(value)

def _legacy_xls_sheet_attrs(sheet: Any) -> dict[str, Any]:
    attrs: dict[str, Any] = {}
    for name in dir(sheet):
        if name.startswith('_'):
            continue
        lower = name.lower()
        if not any((token in lower for token in ('print', 'page', 'zoom', 'fit', 'margin', 'protect', 'pane', 'window'))):
            continue
        try:
            value = getattr(sheet, name)
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            continue
        if callable(value):
            continue
        attrs[name] = _legacy_xls_jsonable(value)
    if hasattr(sheet, 'visibility'):
        try:
            attrs['visibility'] = _legacy_xls_jsonable(getattr(sheet, 'visibility'))
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
    return dict(sorted(attrs.items()))

def _legacy_xls_decode_cell_value(xlrd_module: Any, book: Any, cell: Any) -> Any:
    ctype = getattr(cell, 'ctype', None)
    value = getattr(cell, 'value', None)
    if ctype in {getattr(xlrd_module, 'XL_CELL_EMPTY', 0), getattr(xlrd_module, 'XL_CELL_BLANK', 6)}:
        return None
    if ctype == getattr(xlrd_module, 'XL_CELL_DATE', 3):
        try:
            decoded = xlrd_module.xldate_as_datetime(value, getattr(book, 'datemode', 0))
            return decoded.isoformat() if hasattr(decoded, 'isoformat') else str(decoded)
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            return value
    if ctype == getattr(xlrd_module, 'XL_CELL_BOOLEAN', 4):
        return bool(value)
    if isinstance(value, str):
        return _legacy_xls_trim_text(value)
    return value

def _legacy_xls_encryption_probe(path: Path) -> dict[str, Any]:
    try:
        import msoffcrypto
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'available': False, 'error': str(exc)}
    try:
        with path.open('rb') as fh:
            office = msoffcrypto.OfficeFile(fh)
            return {'available': True, 'encrypted': bool(office.is_encrypted())}
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'available': True, 'error': str(exc)}

def inspect_legacy_xls(path: Path, *, sheet: str | None=None, max_sheets: int=10, preview_rows: int=25, preview_cols: int=25, max_nonempty_cells: int=400, max_merged_cells: int=20) -> dict[str, Any]:
    try:
        import olefile
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'status': 'blocked', 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'error': f'olefile unavailable: {exc}', 'blocker': 'dependency_missing', 'next_step': 'Install olefile or route the workbook through another approved Office reader.', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    try:
        import xlrd
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'status': 'blocked', 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'error': f'xlrd unavailable: {exc}', 'blocker': 'dependency_missing', 'next_step': 'Install xlrd or use another approved workbook reader.', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    size_bytes = path.stat().st_size
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    encryption = _legacy_xls_encryption_probe(path)
    file_meta: dict[str, Any] = {'size_bytes': size_bytes, 'sha256': digest, 'suffix': path.suffix.lower()}
    if isinstance(encryption, dict):
        file_meta['encryption_probe'] = encryption
    if encryption.get('available') and encryption.get('encrypted') is True:
        return {'status': 'blocked', 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'file': file_meta, 'blocker': 'encrypted_workbook', 'message': 'Workbook is encrypted; provide the open password or an unencrypted copy.', 'next_step': 'Ask for the workbook open password or an unencrypted copy, then retry the read-only inspection.', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    ole_info: dict[str, Any]
    try:
        ole = olefile.OleFileIO(str(path))
        try:
            streams = ['/'.join(parts) for parts in ole.listdir()]
            ole_info = {'available': True, 'stream_count': len(streams), 'has_workbook_stream': bool(ole.exists('Workbook') or ole.exists('Book')), 'streams_preview': streams[:20], 'streams_truncated': len(streams) > 20}
        finally:
            try:
                ole.close()
            except Exception:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        ole_info = {'available': True, 'error': str(exc)}
    try:
        book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    except Exception as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        return {'status': 'blocked', 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'file': file_meta, 'ole': ole_info, 'error': f'xlrd open failed: {exc}', 'blocker': 'workbook_read_failed', 'next_step': 'If the workbook is password-protected, provide the open password or an unencrypted copy.', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
    try:
        sheet_names = list(book.sheet_names())
        if sheet:
            target_sheet = sheet.strip()
            if target_sheet not in sheet_names:
                return {'status': 'blocked', 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'file': file_meta, 'ole': ole_info, 'sheet_filter': target_sheet, 'sheet_names': sheet_names[:50], 'sheet_names_truncated': len(sheet_names) > 50, 'blocker': 'sheet_not_found', 'next_step': 'Use one of the reported sheet names or omit --sheet to inspect the first bounded sheets.', 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True}}
            selected_indexes = [sheet_names.index(target_sheet)]
        else:
            selected_indexes = list(range(min(len(sheet_names), max(1, int(max_sheets)))))
        reports: list[dict[str, Any]] = []
        partial = False
        for idx in selected_indexes:
            name = sheet_names[idx]
            try:
                sh = book.sheet_by_index(idx)
                merged_cells = list(getattr(sh, 'merged_cells', []) or [])
                row_cap = min(max(1, int(preview_rows)), int(getattr(sh, 'nrows', 0)))
                col_cap = min(max(1, int(preview_cols)), int(getattr(sh, 'ncols', 0)))
                preview_rows_out: list[list[Any]] = []
                for row_idx in range(row_cap):
                    preview_rows_out.append([_legacy_xls_decode_cell_value(xlrd, book, sh.cell(row_idx, col_idx)) for col_idx in range(col_cap)])
                sample_cells: list[dict[str, Any]] = []
                sample_truncated = False
                nonempty_seen = 0
                empty_ctypes = {getattr(xlrd, 'XL_CELL_EMPTY', 0), getattr(xlrd, 'XL_CELL_BLANK', 6)}
                for row_idx in range(int(getattr(sh, 'nrows', 0))):
                    for col_idx in range(int(getattr(sh, 'ncols', 0))):
                        cell = sh.cell(row_idx, col_idx)
                        if getattr(cell, 'ctype', None) in empty_ctypes:
                            continue
                        nonempty_seen += 1
                        sample_cells.append({'address': _legacy_xls_cell_address(row_idx, col_idx), 'row': row_idx + 1, 'col': col_idx + 1, 'ctype': _legacy_xls_cell_type_label(getattr(cell, 'ctype', None)), 'value': _legacy_xls_decode_cell_value(xlrd, book, cell)})
                        if len(sample_cells) >= max(1, int(max_nonempty_cells)):
                            sample_truncated = True
                            break
                    if sample_truncated:
                        break
                report = {'status': 'ok', 'name': name, 'index': idx, 'visibility': _legacy_xls_visibility_label(getattr(sh, 'visibility', None)), 'attrs': _legacy_xls_sheet_attrs(sh), 'nrows': int(getattr(sh, 'nrows', 0)), 'ncols': int(getattr(sh, 'ncols', 0)), 'preview_rows': preview_rows_out, 'preview_rows_requested': int(preview_rows), 'preview_cols_requested': int(preview_cols), 'preview_truncated': bool(int(getattr(sh, 'nrows', 0)) > row_cap or int(getattr(sh, 'ncols', 0)) > col_cap), 'sample_cells': sample_cells, 'sampled_nonempty_cells': nonempty_seen, 'sample_cells_truncated': sample_truncated, 'merged_cells_count': len(merged_cells), 'merged_cells': [list(map(int, merged)) for merged in merged_cells[:max(0, int(max_merged_cells))]], 'merged_cells_truncated': len(merged_cells) > max(0, int(max_merged_cells))}
                reports.append(report)
            except Exception as exc:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                partial = True
                reports.append({'status': 'error', 'name': name, 'index': idx, 'error': str(exc)})
        status = 'partial' if partial or any((item.get('status') != 'ok' for item in reports)) else 'ok'
        return {'status': status, 'capability': 'pd_one_legacy_xls_inspect', 'input': str(path), 'file': file_meta, 'ole': ole_info, 'workbook': {'datemode': int(getattr(book, 'datemode', 0)), 'sheet_count': len(sheet_names), 'sheet_names': sheet_names[:50], 'sheet_names_truncated': len(sheet_names) > 50, 'sheet_filter': sheet.strip() if sheet else None, 'max_sheets': int(max_sheets), 'inspected_sheet_count': len(reports), 'inspected_sheets_truncated': bool(sheet is None and len(selected_indexes) < len(sheet_names)), 'max_nonempty_cells': int(max_nonempty_cells), 'max_merged_cells': int(max_merged_cells)}, 'sheets': reports, 'safety': {'read_only': True, 'shell_used': False, 'source_system_writes_forbidden': True, 'legacy_xls_only': True, 'bounded_sheet_preview': True}}
    finally:
        try:
            book.release_resources()
        except Exception:
            logging.getLogger('pd_one_ops').exception('pd one command failed')

def extracted_plain_text(extracted: dict[str, Any]) -> str:
    parts: list[str] = []
    for p in extracted.get('paragraphs') or []:
        if isinstance(p, str):
            parts.append(p)
    for table in extracted.get('tables') or []:
        if isinstance(table, list):
            for row in table:
                if isinstance(row, list):
                    parts.extend((str(cell) for cell in row))
    for slide in extracted.get('slides') or []:
        if isinstance(slide, dict):
            parts.extend((str(t) for t in slide.get('texts') or []))
    for page in extracted.get('pages') or []:
        if isinstance(page, dict):
            parts.append(str(page.get('text') or ''))
    return '\n'.join((p for p in parts if p)).strip()
ANNOUNCEMENT_SIGNALS = {'公告': 3, 'announcement': 3, '文號': 2, '發文單位': 2, '受文單位': 2, '知會單位': 2, '發文日期': 2, '主旨': 2, '核准': 1, '審核': 1, '承辦': 1, 'pro-design co': 2}
REWARD_DISCIPLINE_SIGNALS = {'行政懲處': 5, '懲處': 4, '獎懲': 4, '員工獎懲辦法': 5, '申誡': 4, '小過': 4, '大過': 4, '嘉獎': 4, '記功': 4, 'disciplinary': 4, 'discipline': 4, 'demerit': 4, 'reprimand': 4, 'reward': 3}

def classify_company_announcement(path: Path, extracted: dict[str, Any]) -> dict[str, Any]:
    text = (path.name + '\n' + extracted_plain_text(extracted)).casefold()
    hits: list[str] = []
    announcement_score = 0
    reward_discipline_score = 0
    for term, weight in ANNOUNCEMENT_SIGNALS.items():
        if term.casefold() in text:
            hits.append(term)
            announcement_score += weight
    for term, weight in REWARD_DISCIPLINE_SIGNALS.items():
        if term.casefold() in text:
            hits.append(term)
            reward_discipline_score += weight
    should_use = announcement_score >= 5 or reward_discipline_score >= 4 or (announcement_score >= 3 and reward_discipline_score >= 3)
    variant = 'reward-discipline' if reward_discipline_score >= 4 else 'company-announcement' if should_use else 'none'
    return {'document_type': variant, 'use_company_announcement_template': should_use, 'requires_visual_qa_for_docx': should_use, 'confidence': 'high' if announcement_score + reward_discipline_score >= 8 else 'medium' if should_use else 'low', 'announcement_score': announcement_score, 'reward_discipline_score': reward_discipline_score, 'matched_signals': sorted(set(hits)), 'recommended_skill': 'company-announcement-writer' if should_use else 'none', 'recommended_reference': 'references/reward-discipline-announcements.md' if variant == 'reward-discipline' else 'references/announcement-template.md' if should_use else 'none', 'required_handling': 'For translate/polish/format requests, produce the employee-facing output as a company-announcement-template artifact (normally DOCX) and run visual-office-qa before delivery; a plain-text translation alone is insufficient. Publishing/sending beyond replying with the requested draft/export remains approval-gated.' if should_use else 'No company-announcement template required by deterministic classifier.'}

def extract_by_suffix(path: Path) -> tuple[str, dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == '.docx':
        return (suffix, extract_docx(path))
    if suffix == '.doc':
        return (suffix, extract_legacy_doc(path))
    if suffix == '.pptx':
        return (suffix, extract_pptx(path))
    if suffix == '.pdf':
        return (suffix, extract_pdf(path))
    fail('supported document types: .doc, .docx, .pptx, .pdf')

def cmd_doc_extract(args: argparse.Namespace) -> None:
    path = require_allowed_doc_path(Path(args.input))
    suffix, extracted = extract_by_suffix(path)
    result = {'status': 'ok', 'capability': 'pd_one_doc_extract', 'input': str(path), 'type': suffix, 'extracted': extracted}
    if args.out:
        out = Path(args.out).resolve(strict=False)
        if not any((is_within(out, root) for root in DOC_ALLOWED_ROOTS())):
            fail(f'output path is outside approved roots: {out}')
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(stable_dump(result), encoding='utf-8')
        result['out'] = str(out)
    emit(result)

def cmd_artifact_base64_encode(args: argparse.Namespace) -> None:
    """Base64-encode one bounded local media artifact from approved roots."""
    path = Path(args.input).resolve(strict=False)
    if not any((is_within(path, root) for root in ARTIFACT_BASE64_ALLOWED_ROOTS())):
        fail(f'path is outside approved PD One artifact roots: {path}')
    if not path.is_file():
        fail(f'artifact is not a regular file: {path}')
    suffix = path.suffix.lower()
    if suffix not in ARTIFACT_BASE64_ALLOWED_SUFFIXES:
        fail(f'supported artifact types: {', '.join(sorted(ARTIFACT_BASE64_ALLOWED_SUFFIXES))}')
    size = path.stat().st_size
    if size > ARTIFACT_BASE64_MAX_BYTES:
        fail(f'artifact exceeds {ARTIFACT_BASE64_MAX_BYTES} byte limit: {size}')
    data = path.read_bytes()
    emit({'status': 'ok', 'capability': 'pd_one_artifact_base64_encode', 'input': str(path), 'mime_type': ARTIFACT_BASE64_MIME_TYPES[suffix], 'size_bytes': size, 'sha256': hashlib.sha256(data).hexdigest(), 'base64': base64.b64encode(data).decode('ascii'), 'safety': {'read_only': True, 'admin_only': True, 'approved_artifact_roots_only': True, 'max_bytes': ARTIFACT_BASE64_MAX_BYTES, 'source_system_writes_forbidden': True, 'external_delivery': False, 'shell_used': False}})

def cmd_local_artifact_copy(args: argparse.Namespace) -> None:
    """Copy one legacy employee-workspace artifact to the canonical artifact tree."""
    source = Path(args.input).resolve(strict=False)
    output = Path(args.out).resolve(strict=False)
    if not is_within(source, LOCAL_ARTIFACT_COPY_SOURCE_ROOT()):
        fail(f'input is outside the approved legacy artifact root: {source}')
    if not is_within(output, LOCAL_ARTIFACT_COPY_DEST_ROOT()):
        fail(f'output is outside the canonical PD One artifact root: {output}')
    if not source.is_file() or source.is_symlink():
        fail('input must be a regular, non-symlink file')
    relative = source.relative_to(LOCAL_ARTIFACT_COPY_SOURCE_ROOT().resolve())
    expected_output = (LOCAL_ARTIFACT_COPY_DEST_ROOT() / relative).resolve(strict=False)
    if output != expected_output:
        fail(f'output must preserve the source-relative artifact path: {expected_output}')
    suffix = source.suffix.lower()
    if suffix not in LOCAL_ARTIFACT_COPY_ALLOWED_SUFFIXES or output.suffix.lower() != suffix:
        fail('input/output suffix is not an approved artifact type or does not match')
    source_size = source.stat().st_size
    if source_size <= 0 or source_size > LOCAL_ARTIFACT_COPY_MAX_BYTES:
        fail(f'input size must be 1-{LOCAL_ARTIFACT_COPY_MAX_BYTES} bytes')
    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    if args.source_sha256 and args.source_sha256.lower() != source_sha256:
        fail('source SHA-256 guard mismatch')
    safety = {'source_read_only': True, 'overwrite_forbidden': True, 'relative_path_preserved': True, 'local_artifact_write_only': True, 'external_delivery': False}
    if output.exists():
        if not output.is_file() or output.is_symlink():
            fail('existing output is not a regular, non-symlink file')
        if output.stat().st_size != source_size or hashlib.sha256(output.read_bytes()).hexdigest() != source_sha256:
            fail('output already exists with different content; refusing overwrite')
        emit({'status': 'ok', 'capability': 'pd_one_local_artifact_copy', 'action': 'already_present', 'input': str(source), 'output': str(output), 'size_bytes': source_size, 'sha256': source_sha256, 'safety': safety})
        return
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f'.{output.name}.', dir=str(output.parent), delete=False) as tmp:
        temporary = Path(tmp.name)
    try:
        shutil.copy2(source, temporary)
        if temporary.stat().st_size != source_size or hashlib.sha256(temporary.read_bytes()).hexdigest() != source_sha256:
            fail('staged artifact readback did not match source')
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    emit({'status': 'ok', 'capability': 'pd_one_local_artifact_copy', 'action': 'copied', 'input': str(source), 'output': str(output), 'size_bytes': output.stat().st_size, 'sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'safety': safety})

def cmd_video_frame_sample(args: argparse.Namespace) -> None:
    """Probe one cached video and stage three fixed representative JPEG frames."""
    source = Path(args.input).resolve(strict=False)
    if not is_within(source, VIDEO_FRAME_INPUT_ROOT()):
        fail(f'video is outside the approved cache root: {source}')
    if not source.is_file() or source.suffix.lower() not in VIDEO_FRAME_ALLOWED_SUFFIXES:
        fail(f'input must be a regular cached video ({', '.join(sorted(VIDEO_FRAME_ALLOWED_SUFFIXES))}): {source}')
    source_size = source.stat().st_size
    if source_size <= 0 or source_size > VIDEO_FRAME_MAX_SOURCE_BYTES:
        fail(f'video size must be 1-{VIDEO_FRAME_MAX_SOURCE_BYTES} bytes: {source_size}')
    fractions = VIDEO_FRAME_DEFAULT_FRACTIONS
    ffprobe = shutil.which('ffprobe')
    ffmpeg = shutil.which('ffmpeg')
    if not ffprobe or not ffmpeg:
        fail('ffprobe and ffmpeg are required')
    probe_cmd = [ffprobe, '-v', 'error', '-show_entries', 'format=duration,size:stream=codec_type,codec_name,width,height,r_frame_rate', '-of', 'json', str(source)]
    try:
        probe_proc = subprocess.run(probe_cmd, text=True, capture_output=True, shell=False, timeout=30)
    except subprocess.TimeoutExpired:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail('ffprobe timed out after 30 seconds')
    if probe_proc.returncode != 0:
        fail(f'ffprobe failed with exit {probe_proc.returncode}: {probe_proc.stderr[-500:]}')
    try:
        probe = json.loads(probe_proc.stdout)
        duration = float(probe['format']['duration'])
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        logging.getLogger('pd_one_ops').exception('pd one command failed')
        fail(f'ffprobe returned an invalid duration receipt: {exc}')
    if not math.isfinite(duration) or duration <= 0 or duration > VIDEO_FRAME_MAX_DURATION_SECONDS:
        fail(f'video duration must be within 0-{VIDEO_FRAME_MAX_DURATION_SECONDS} seconds: {duration}')
    video_stream = next((s for s in probe.get('streams', []) if s.get('codec_type') == 'video'), None)
    if not video_stream:
        fail('ffprobe found no video stream')
    source_key = hashlib.sha256(str(source).encode('utf-8')).hexdigest()[:12]
    safe_stem = re.sub('[^A-Za-z0-9._-]+', '_', source.stem)[:48] or 'video'
    output_dir = (VIDEO_FRAME_OUTPUT_ROOT() / f'{safe_stem}-{source_key}').resolve(strict=False)
    if not is_within(output_dir, VIDEO_FRAME_OUTPUT_ROOT()):
        fail('resolved output directory escaped the approved video-frame root')
    output_dir.mkdir(parents=True, exist_ok=True)
    frames: list[dict[str, Any]] = []
    for index, fraction in enumerate(fractions, 1):
        timestamp = duration * fraction
        final_path = output_dir / f'frame_{index}.jpg'
        temp_path = output_dir / f'.frame_{index}.{os.getpid()}.tmp.jpg'
        command = [ffmpeg, '-nostdin', '-hide_banner', '-loglevel', 'error', '-y', '-ss', f'{timestamp:.6f}', '-i', str(source), '-frames:v', '1', '-q:v', '2', str(temp_path)]
        try:
            proc = subprocess.run(command, text=True, capture_output=True, shell=False, timeout=60)
        except subprocess.TimeoutExpired:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            temp_path.unlink(missing_ok=True)
            fail(f'ffmpeg timed out extracting frame {index}')
        if proc.returncode != 0 or not temp_path.is_file() or temp_path.stat().st_size <= 0:
            temp_path.unlink(missing_ok=True)
            fail(f'ffmpeg failed extracting frame {index} with exit {proc.returncode}: {proc.stderr[-500:]}')
        os.replace(temp_path, final_path)
        frame_data = final_path.read_bytes()
        frames.append({'index': index, 'position_fraction': fraction, 'timestamp_seconds': round(timestamp, 6), 'path': str(final_path), 'exists': True, 'size_bytes': len(frame_data), 'sha256': hashlib.sha256(frame_data).hexdigest()})
    emit({'status': 'ok', 'capability': 'pd_one_video_frame_sample', 'input': str(source), 'source_size_bytes': source_size, 'probe': {'duration_seconds': duration, 'codec_name': video_stream.get('codec_name'), 'width': video_stream.get('width'), 'height': video_stream.get('height'), 'r_frame_rate': video_stream.get('r_frame_rate')}, 'output_dir': str(output_dir), 'frames': frames, 'safety': {'admin_only': True, 'source_read_only': True, 'local_staged_artifact_write_only': True, 'approved_roots_only': True, 'shell_used': False, 'fixed_sample_positions': list(VIDEO_FRAME_DEFAULT_FRACTIONS), 'source_system_writes_forbidden': True, 'external_delivery': False}})

def cmd_svg_artifact_inspect(args: argparse.Namespace) -> None:
    data = _PdUnset
    elements = _PdUnset
    embedded_png_count = _PdUnset
    failures = _PdUnset
    image_receipts = _PdUnset
    images = _PdUnset
    invalid_text_y_count = _PdUnset
    path = _PdUnset
    required = _PdUnset
    root = _PdUnset
    size = _PdUnset
    text_y_values = _PdUnset

    def _pd_phase_7():
        nonlocal data, elements, embedded_png_count, failures, image_receipts, images, invalid_text_y_count, path, required, root, size, text_y_values
        'Validate bounded structural/text invariants for one local SVG artifact.'
        path = Path(args.input).resolve(strict=False)
        if not any((is_within(path, root) for root in ARTIFACT_BASE64_ALLOWED_ROOTS())):
            fail(f'path is outside approved PD One artifact roots: {path}')
        if not path.is_file() or path.suffix.lower() != '.svg':
            fail(f'input must be a regular .svg file: {path}')
        size = path.stat().st_size
        if size > SVG_INSPECT_MAX_BYTES:
            fail(f'SVG exceeds {SVG_INSPECT_MAX_BYTES} byte limit: {size}')
        data = path.read_bytes()
        try:
            source = data.decode('utf-8')
        except UnicodeDecodeError:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            fail('SVG must be valid UTF-8')
        lowered = source.lower()
        if '<!doctype' in lowered or '<!entity' in lowered:
            fail('SVG DOCTYPE/entity declarations are not allowed')
        try:
            root = ET.fromstring(source)
        except ET.ParseError as exc:
            logging.getLogger('pd_one_ops').exception('pd one command failed')
            fail(f'invalid SVG XML: {exc}')
        svg_ns = 'http://www.w3.org/2000/svg'
        if root.tag != f'{{{svg_ns}}}svg':
            fail('root element must be an SVG element in the SVG namespace')
        elements = list(root.iter())
        if len(elements) > SVG_INSPECT_MAX_ELEMENTS:
            fail(f'SVG exceeds {SVG_INSPECT_MAX_ELEMENTS} element limit')
        required = list(args.required_text or [])
        if len(required) > SVG_INSPECT_MAX_REQUIRED_TEXTS:
            fail(f'at most {SVG_INSPECT_MAX_REQUIRED_TEXTS} required-text values are allowed')
        if any((not term or len(term) > SVG_INSPECT_MAX_REQUIRED_TEXT_CHARS for term in required)):
            fail(f'required-text values must be 1-{SVG_INSPECT_MAX_REQUIRED_TEXT_CHARS} characters')
        visible_text = '\n'.join((part.strip() for part in root.itertext() if part.strip()))
        missing = [term for term in required if term not in visible_text]
        images = [element for element in elements if element.tag == f'{{{svg_ns}}}image']
        if len(images) > SVG_INSPECT_MAX_IMAGES:
            fail(f'SVG exceeds {SVG_INSPECT_MAX_IMAGES} image element limit')
        xlink_href = '{http://www.w3.org/1999/xlink}href'
        image_receipts = []
        embedded_png_count = 0
        for index, image in enumerate(images):
            href = image.attrib.get('href', image.attrib.get(xlink_href, ''))
            embedded_png = href.startswith('data:image/png;base64,')
            embedded_png_count += int(embedded_png)
            image_receipts.append({'index': index, 'embedded_png': embedded_png, 'box': {key: image.attrib.get(key, '')[:SVG_INSPECT_MAX_ATTR_CHARS] for key in ('x', 'y', 'width', 'height')}})
        text_y_values = []
        invalid_text_y_count = 0
        for element in elements:
            if element.tag != f'{{{svg_ns}}}text' or 'y' not in element.attrib:
                continue
            raw_y = element.attrib['y'].strip()
            try:
                text_y_values.append(float(raw_y))
            except ValueError:
                logging.getLogger('pd_one_ops').exception('pd one command failed')
                invalid_text_y_count += 1
        failures = []
        if missing:
            failures.append(f'missing required text: {missing}')
        if args.expected_embedded_png_count is not None:
            if not 0 <= args.expected_embedded_png_count <= SVG_INSPECT_MAX_IMAGES:
                fail(f'expected embedded PNG count must be 0-{SVG_INSPECT_MAX_IMAGES}')
            if embedded_png_count != args.expected_embedded_png_count:
                failures.append(f'expected {args.expected_embedded_png_count} embedded PNG image(s), found {embedded_png_count}')

    def _pd_phase_8():
        if args.max_text_y_exclusive is not None:
            if not math.isfinite(args.max_text_y_exclusive):
                fail('--max-text-y-exclusive must be finite')
            if invalid_text_y_count:
                failures.append(f'found {invalid_text_y_count} non-numeric text y value(s)')
            if not text_y_values:
                failures.append('no numeric text y values found')
            elif max(text_y_values) >= args.max_text_y_exclusive:
                failures.append(f'maximum text y {max(text_y_values)} is not below {args.max_text_y_exclusive}')
        if args.expected_image_box:
            expected_box_values = args.expected_image_box.split(',')
            if len(expected_box_values) != 4 or any((not value for value in expected_box_values)):
                fail('--expected-image-box must be x,y,width,height')
            expected_box = dict(zip(('x', 'y', 'width', 'height'), expected_box_values))
            embedded_boxes = [item['box'] for item in image_receipts if item['embedded_png']]
            if expected_box not in embedded_boxes:
                failures.append(f'expected embedded PNG image box not found: {expected_box}')
        if failures:
            fail('; '.join(failures))
        emit({'status': 'ok', 'capability': 'pd_one_svg_artifact_inspect', 'input': str(path), 'size_bytes': size, 'sha256': hashlib.sha256(data).hexdigest(), 'xml_valid': True, 'canvas': {'width': root.attrib.get('width', '')[:SVG_INSPECT_MAX_ATTR_CHARS], 'height': root.attrib.get('height', '')[:SVG_INSPECT_MAX_ATTR_CHARS], 'viewBox': root.attrib.get('viewBox', '')[:SVG_INSPECT_MAX_ATTR_CHARS]}, 'element_count': len(elements), 'required_text_count': len(required), 'required_text_preserved': True, 'image_count': len(images), 'embedded_png_count': embedded_png_count, 'images': image_receipts, 'text_y': {'numeric_count': len(text_y_values), 'invalid_count': invalid_text_y_count, 'maximum': max(text_y_values) if text_y_values else None, 'required_max_exclusive': args.max_text_y_exclusive}, 'safety': {'read_only': True, 'admin_only': True, 'approved_artifact_roots_only': True, 'bounded_output': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'shell_used': False}})
    _pd_result = _pd_phase_7()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value
    _pd_result = _pd_phase_8()
    if isinstance(_pd_result, _PdReturn):
        return _pd_result.value

def _parse_3mf_slicer_settings(text: str) -> list[dict[str, str]]:
    found: list[dict[str, str]] = []
    seen: set[str] = set()
    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith(';') or line.startswith('#'):
            line = line[1:].strip()
        if not line or '=' not in line:
            continue
        key, value = line.split('=', 1)
        key = key.strip()
        value = value.strip().strip('"')
        if not key or key in seen:
            continue
        if key in THREEMF_SLICER_SETTING_KEYS:
            seen.add(key)
            found.append({'key': key[:80], 'value': value[:THREEMF_SLICER_MAX_VALUE_CHARS]})
            if len(found) >= THREEMF_SLICER_MAX_SETTINGS:
                break
    return found

def _safe_3mf_zip_info(info: zipfile.ZipInfo) -> str:
    name = info.filename.replace('\\', '/')
    parts = Path(name).parts
    if name.startswith('/') or '..' in parts:
        fail(f'refusing unsafe 3MF member path: {name}')
    return name

def _validate_artifact_identity_glob(pattern: str) -> str:
    text = str(pattern or '').strip()
    if not text or len(text) > ARTIFACT_IDENTITY_MAX_GLOB_CHARS:
        fail('glob pattern is missing or too long')
    if any((marker in text for marker in ('/', '\\', '..', '**'))):
        fail('glob pattern must stay in the selected directory')
    return text

def cmd_artifact_identity_inventory(args: argparse.Namespace) -> None:
    """List bounded artifact identities and optionally peek one file magic."""
    root = Path(args.root).resolve(strict=False)
    if root.is_symlink() or any((parent.is_symlink() for parent in root.parents if str(parent).startswith(str(ROOT())))):
        fail('artifact root must not be a symlink')
    if not any((is_within(root, allowed) for allowed in ARTIFACT_IDENTITY_INPUT_ROOTS())):
        fail(f'path is outside approved PD One artifact roots: {root}')
    if not root.is_dir() or root.is_symlink():
        fail(f'root must be a regular directory: {root}')
    raw_globs = [str(item).strip() for item in getattr(args, 'glob', None) or [] if str(item).strip()]
    peek_name = str(getattr(args, 'peek', '') or '').strip()
    if not raw_globs and (not peek_name):
        fail('provide at least one --glob or a --peek file')
    if len(raw_globs) > ARTIFACT_IDENTITY_MAX_GLOBS:
        fail(f'at most {ARTIFACT_IDENTITY_MAX_GLOBS} glob patterns are allowed')
    patterns = [_validate_artifact_identity_glob(item) for item in raw_globs]
    peek_bytes = int(getattr(args, 'peek_bytes', ARTIFACT_IDENTITY_DEFAULT_PEEK_BYTES) or ARTIFACT_IDENTITY_DEFAULT_PEEK_BYTES)
    if peek_bytes < 1 or peek_bytes > ARTIFACT_IDENTITY_MAX_PEEK_BYTES:
        fail(f'peek bytes must be 1-{ARTIFACT_IDENTITY_MAX_PEEK_BYTES}')
    matches: list[dict[str, Any]] = []
    seen: set[str] = set()
    skipped_unallowed = 0
    truncated = False
    for pattern in patterns:
        for candidate in sorted(root.glob(pattern), key=lambda item: item.name):
            if candidate.is_symlink() or not candidate.is_file():
                continue
            resolved = candidate.resolve(strict=False)
            if not is_within(resolved, root):
                fail(f'glob match escaped inventory root: {resolved}')
            suffix = resolved.suffix.lower()
            if suffix not in ARTIFACT_IDENTITY_ALLOWED_SUFFIXES:
                skipped_unallowed += 1
                continue
            key = str(resolved)
            if key in seen:
                continue
            seen.add(key)
            if len(matches) >= ARTIFACT_IDENTITY_MAX_MATCHES:
                truncated = True
                break
            matches.append({'name': resolved.name, 'relative_path': str(resolved.relative_to(root)), 'suffix': suffix, 'size_bytes': resolved.stat().st_size, 'exists': True})
        if truncated:
            break
    peek_receipt: dict[str, Any] | None = None
    if peek_name:
        peek_path = Path(peek_name)
        if not peek_path.is_absolute():
            peek_path = root / peek_path
        peek_path = peek_path.resolve(strict=False)
        if peek_path.is_symlink() or not is_within(peek_path, root):
            fail('peek path must be a non-symlink file under the inventory root')
        suffix = peek_path.suffix.lower()
        if suffix not in ARTIFACT_IDENTITY_ALLOWED_SUFFIXES:
            fail(f'peek suffix is not allowed: {suffix}')
        exists = peek_path.is_file() and (not peek_path.is_symlink())
        size = peek_path.stat().st_size if exists else None
        magic_hex = None
        if exists:
            magic_hex = peek_path.read_bytes()[:peek_bytes].hex()
        peek_receipt = {'name': peek_path.name, 'relative_path': str(peek_path.relative_to(root)), 'suffix': suffix, 'exists': exists, 'size_bytes': size, 'magic_hex': magic_hex, 'peek_bytes': peek_bytes}
    emit({'status': 'ok', 'capability': 'pd_one_artifact_identity_inventory', 'root': str(root), 'globs': patterns, 'match_count': len(matches), 'matches': matches, 'matches_truncated': truncated, 'skipped_unallowed_suffix': skipped_unallowed, 'peek': peek_receipt, 'safety': {'admin_only': True, 'read_only': True, 'approved_roots_only': True, 'single_directory_globs_only': True, 'bounded_matches': True, 'bounded_magic_peek': True, 'source_overwrite_forbidden': True, 'source_system_writes_forbidden': True, 'external_delivery': False, 'shell_used': False, 'zip_used': False}})
